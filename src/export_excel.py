"""Exportación a Excel (openpyxl) de resultados del ciclo Kalina KSC-11.

``generar_excel`` empaqueta el resultado de ``resolver_ciclo`` +
``calcular_exergia`` más los parámetros de entrada en un libro ``.xlsx`` en
memoria (``bytes``), listo para ``st.download_button(data=...)``. NO escribe
a disco. Hojas (Kalina_ksc_11_tercero.md §10, con unidades y metadatos):

- ``Entradas``: una fila por entrada de ``parametros`` (nombre, valor,
  unidad) + una fila con el ``backend_nombre``.
- ``Estados``: una fila por estado e1..e10 con T, P, h, s, x, m, q, fase y
  exergía física [kJ/kg] (de ``resultado_exergia``).
- ``Componentes``: energía externa relevante por componente. Convención: solo
  HRVG (Qi), turbina (Wt), regenerador (Qreg), condensador (Qout) y bomba
  (Wp) reportan energía; separador, válvula y absorbedor son adiabáticos/
  isoentálpicos y su celda de energía queda en "—".
- ``Exergía``: Sgen/Ed por los 8 componentes + Ed_total + Ex_Qi + Ex_Qout +
  el residual del balance global (Ex_Qi − Wnet − Ex_Qout − Ed_total).
- ``Resumen``: Qi, Qout, Wt, Wp, Wnet, eta, eta_exergetico (= Wnet/Ex_Qi) y
  metadatos (fecha ``datetime.now()``, versión fija, backend usado).

Convención de ``parametros``: dict ``{nombre: (valor, unidad)}``; también se
acepta un valor simple (la unidad queda entonces en "—").
"""

from __future__ import annotations

from datetime import datetime
from io import BytesIO

from openpyxl import Workbook

__all__ = ["generar_excel", "VERSION_PROYECTO"]

VERSION_PROYECTO = "ciclo_kalina_tercero v0.1"

_COMPONENTES = ("hrvg", "separador", "turbina", "regenerador", "valvula",
                "absorbedor", "condensador", "bomba")

_NOMBRES = {"hrvg": "HRVG", "separador": "Separador", "turbina": "Turbina",
            "regenerador": "Regenerador", "valvula": "Válvula",
            "absorbedor": "Absorbedor", "condensador": "Condensador",
            "bomba": "Bomba"}

_ENERGIA = {"hrvg": ("Qi", "calor suministrado (fuente)"),
            "turbina": ("Wt", "trabajo de salida"),
            "regenerador": ("Qreg", "intercambio interno al ciclo"),
            "condensador": ("Qout", "calor rechazado (sumidero)"),
            "bomba": ("Wp", "trabajo de entrada")}

_SIN_ENERGIA = {"separador": "separador ideal, adiabático: sin energía externa",
                "valvula": "isoentálpica: sin energía externa",
                "absorbedor": "mezcla adiabática: sin energía externa"}


def _tachado(valor):
    """``valor`` o "—" si es None (q, fase y exergía pueden no estar definidos)."""
    return valor if valor is not None else "—"


def _hoja_entradas(wb, parametros, backend_nombre):
    ws = wb.active
    ws.title = "Entradas"
    ws.append(["Variable", "Valor", "Unidad"])
    for nombre, entrada in parametros.items():
        if isinstance(entrada, tuple):
            valor, unidad = entrada
        else:
            valor, unidad = entrada, "—"
        ws.append([nombre, valor, unidad])
    ws.append(["backend", backend_nombre, "—"])


def _hoja_estados(wb, resultado_ciclo, resultado_exergia):
    ws = wb.create_sheet("Estados")
    ws.append(["Estado", "T [K]", "P [kPa]", "h [kJ/kg]", "s [kJ/kg·K]",
               "x [-]", "m [kg/s]", "q [-]", "Fase", "ex_fisica [kJ/kg]"])
    estados = resultado_ciclo["estados"]
    ex = resultado_exergia["estados"]
    for clave in sorted(estados, key=lambda k: int(k[1:])):
        e = estados[clave]
        ex_fisica = ex[clave].exergia_fisica if clave in ex else None
        ws.append([e.etiqueta or clave, e.T, e.P, e.h, e.s, e.x, e.m,
                   _tachado(e.q), _tachado(e.fase), _tachado(ex_fisica)])


def _hoja_componentes(wb, resultado_ciclo):
    ws = wb.create_sheet("Componentes")
    ws.append(["Componente", "Energía [kW]", "Tipo de energía"])
    for comp in _COMPONENTES:
        if comp in _ENERGIA:
            energia, tipo = _ENERGIA[comp]
            ws.append([_NOMBRES[comp], resultado_ciclo[energia], tipo])
        else:
            ws.append([_NOMBRES[comp], "—", _SIN_ENERGIA[comp]])


def _hoja_exergia(wb, resultado_ciclo, resultado_exergia):
    ws = wb.create_sheet("Exergía")
    ws.append(["Componente", "Sgen [kW/K]", "Ed [kW]"])
    sgen, ed = resultado_exergia["sgen"], resultado_exergia["ed"]
    for comp in _COMPONENTES:
        ws.append([_NOMBRES[comp], sgen[comp], ed[comp]])
    ed_total = resultado_exergia["ed_total"]
    ex_qi = resultado_exergia["ex_qi"]
    ex_qout = resultado_exergia["ex_qout"]
    ws.append(["Ed_total (suma de componentes)", "—", ed_total])
    ws.append(["Ex_Qi (exergía de la fuente)", "—", ex_qi])
    ws.append(["Ex_Qout (exergía del calor rechazado)", "—", ex_qout])
    residual = ex_qi - resultado_ciclo["Wnet"] - ex_qout - ed_total
    ws.append(["Residual: Ex_Qi - Wnet - Ex_Qout - Ed_total", "—", residual])


def _hoja_resumen(wb, resultado_ciclo, resultado_exergia, backend_nombre):
    ws = wb.create_sheet("Resumen")
    ws.append(["Magnitud", "Valor", "Unidad"])
    for nombre in ("Qi", "Qout", "Wt", "Wp", "Wnet"):
        ws.append([nombre, resultado_ciclo[nombre], "kW"])
    ws.append(["eta", resultado_ciclo["eta"], "-"])
    ex_qi = resultado_exergia["ex_qi"]
    eta_ex = resultado_ciclo["Wnet"] / ex_qi if ex_qi else "—"
    ws.append(["eta_exergetico (Wnet/Ex_Qi)", eta_ex, "-"])
    ws.append(["Fecha de generación",
               datetime.now().isoformat(timespec="seconds"), "—"])
    ws.append(["Versión del proyecto", VERSION_PROYECTO, "—"])
    ws.append(["Backend", backend_nombre, "—"])


def generar_excel(resultado_ciclo, resultado_exergia, parametros,
                  backend_nombre) -> bytes:
    """Libro ``.xlsx`` completo (bytes) de una corrida del ciclo.

    Argumentos: ``resultado_ciclo`` (dict de ``resolver_ciclo``),
    ``resultado_exergia`` (dict de ``calcular_exergia``), ``parametros``
    (dict ``{nombre: (valor, unidad)}`` de las entradas de la corrida, ver
    docstring del módulo) y ``backend_nombre`` (str, añadido a las hojas
    ``Entradas`` y ``Resumen``).

    Devuelve ``bytes`` — en memoria, sin escribir a disco.
    """
    wb = Workbook()
    _hoja_entradas(wb, parametros, backend_nombre)
    _hoja_estados(wb, resultado_ciclo, resultado_exergia)
    _hoja_componentes(wb, resultado_ciclo)
    _hoja_exergia(wb, resultado_ciclo, resultado_exergia)
    _hoja_resumen(wb, resultado_ciclo, resultado_exergia, backend_nombre)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()