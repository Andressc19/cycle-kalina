"""Marca visual en Excel de los valores fuera del rango confiable del motor.

Gris `BFBFBF` + comentario con el detalle (parámetro, valor, límite, motor). El
gris se aplica solo a las celdas afectadas, no a la fila.
"""

from __future__ import annotations

from openpyxl.comments import Comment
from openpyxl.styles import PatternFill

__all__ = ["GRIS_RANGO", "LEYENDA_RANGO", "marcar_tabla", "marcar_estados"]

GRIS_RANGO = "BFBFBF"
LEYENDA_RANGO = ("Gris: valor calculado fuera del rango confiable del motor "
                 "(IAPWS G4-01 §6, Tillner-Roth & Friend 1998); el comentario de la "
                 "celda dice qué parámetro, qué límite y qué motor. Tomar con pinzas.")
_AUTOR = "rangos_motor"


def _gris(cell, texto: str) -> None:
    cell.fill = PatternFill("solid", fgColor=GRIS_RANGO)
    cell.comment = Comment(texto, _AUTOR, width=400, height=120)


def marcar_tabla(ws, cab: list[str], fila_cab: int = 1) -> int:
    """Pinta las celdas que nombra `columnas_rango` en cada fila de datos.

    `cab` son los nombres limpios de las columnas, en orden. Devuelve cuántas
    celdas marcó; sin las columnas de rango no hace nada.
    """
    if "columnas_rango" not in cab or "detalle_rango" not in cab:
        return 0
    idx = {c: j for j, c in enumerate(cab, 1)}
    j_cols, j_det = idx["columnas_rango"], idx["detalle_rango"]
    n = 0
    for r in range(fila_cab + 1, ws.max_row + 1):
        cols = ws.cell(r, j_cols).value
        if not cols:
            continue
        detalle = str(ws.cell(r, j_det).value or "")
        for c in str(cols).split(";"):
            if c in idx:
                _gris(ws.cell(r, idx[c]), detalle)
                n += 1
    return n


def marcar_estados(ws, violaciones) -> int:
    """Hoja `Estados` de `export_excel`: gris en la celda T o P del estado."""
    col = {"T": 2, "P": 3, "fase": 9}
    fila = {}
    for r in range(2, ws.max_row + 1):
        fila[str(ws.cell(r, 1).value)] = r
    n = 0
    for v in violaciones:
        r = fila.get(v.estado) or fila.get(f"e{v.estado}")
        if r is not None and v.parametro in col:
            _gris(ws.cell(r, col[v.parametro]), v.texto())
            n += 1
    return n
