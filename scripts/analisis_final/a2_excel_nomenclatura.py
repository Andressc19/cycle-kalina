# -*- coding: utf-8 -*-
"""A2 · Hoja Nomenclatura (2026-10-12): cómo se numera todo y la lista de estudios y partes."""
from openpyxl.styles import Alignment, Font, PatternFill

import codigos
from a2_excel_hojas import BORDE, GRIS, NEGRITA, TITULO

TEXTO = [
    "Cómo se numera todo (como los capítulos de un informe)",
    "",
    "• PREGUNTA (1 a 6): qué queremos saber. El anexo A es una curiosidad que no entra en las conclusiones.",
    "• ESTUDIO (por ejemplo 2.4): un barrido que responde esa pregunta. El primer número es la pregunta.",
    "• PARTE (por ejemplo 2.4a): cuando un estudio tiene partes con datos fijos distintos, cada una lleva "
    "una letra. REGLA: solo se comparan puntos de la misma parte.",
    "• Tabla 4.1-1, Figura 3.2-2: tablas y figuras van numeradas por estudio (4.1-1 = tabla 1 del estudio "
    "4.1). Las de síntesis de toda la campaña llevan S (Tabla S-4, Figura S-1).",
    "• En cada estudio, la Tabla x.y-1 son las condiciones (qué quedó fijo y qué se barrió), la x.y-2 el "
    "conteo de corridas por clase y las siguientes, los resultados.",
    "• El 'código de origen' (M05, KF-M05-3, B09-3…) es el nombre interno de los archivos de cálculo; solo "
    "sirve para rastrear un dato hasta su CSV.",
]


def hoja_nomenclatura(wb, posicion=1):
    ws = wb.create_sheet("Nomenclatura", posicion)
    for i, t in enumerate(TEXTO, 1):
        c = ws.cell(i, 1, t)
        c.font = TITULO if i == 1 else Font(size=10)
    r = len(TEXTO) + 2
    cab = ["Pregunta", "Estudio", "Nombre del estudio", "Partes (letra = datos fijos distintos)",
           "Código de origen"]
    for j, t in enumerate(cab, 1):
        c = ws.cell(r, j, t); c.font = NEGRITA; c.fill = GRIS; c.border = BORDE
    origen = {}
    for caso, (e, _) in codigos.ESTUDIO.items():
        origen.setdefault(e, []).append(caso.replace("KF-", ""))
    partes = {}
    for sub, (p, n) in codigos._P.items():
        e = codigos.estudio(next(k for k in sorted(codigos.ESTUDIO, key=len, reverse=True)
                                 if sub == k or sub.startswith(k + "-")))
        if p != e or len(partes.get(e, [])) or n:
            partes.setdefault(e, []).append(f"{p}: {n}" if p != e else n)
    partes["4.3"] = ["rondas 0 a 3; cada conjunto se nombra con sus variables (p. ej. 4.3 x_b+P_alta)"]
    vistos = []
    for e in sorted(set(v[0] for v in codigos.ESTUDIO.values()),
                    key=lambda x: (x[0] == "A", [float(t) if t.isdigit() else 0 for t in x.split(".")])):
        if e in vistos:
            continue
        vistos.append(e)
        r += 1
        nombre = next(n for k, (x, n) in codigos.ESTUDIO.items() if x == e)
        fila = [e.split(".")[0] + " · " + codigos.PREGUNTA[e.split(".")[0]], e, nombre,
                " | ".join(dict.fromkeys(partes.get(e, []))), ", ".join(origen[e])]
        for j, v in enumerate(fila, 1):
            c = ws.cell(r, j, v); c.border = BORDE
            c.alignment = Alignment(wrap_text=True, vertical="top")
        if e.startswith("A"):
            for j in range(1, 6):
                ws.cell(r, j).fill = PatternFill("solid", fgColor="F8B4B4")
    for col, w in zip("ABCDE", (46, 9, 40, 90, 20)):
        ws.column_dimensions[col].width = w
    return ws.title
