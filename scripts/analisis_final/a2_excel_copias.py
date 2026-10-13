# -*- coding: utf-8 -*-
"""A2 · Copias de los gráficos junto a la tabla de la que salen sus datos (Claude, 2026-10-12).

Pedido del usuario: los gráficos de las hojas «N Gráficos» no se mueven; cada uno, además, va en
la hoja de la tabla de donde se extraen sus datos (columna «Fuente» de su figura):
- figuras que salen de todos_los_puntos.csv → hoja «N.M-R Resultados por barrido», a la derecha
  del bloque de su parte o de su variable (o al inicio si son de todo el estudio);
- figuras que salen de una tabla resumen → debajo de esa tabla.
Las copias leen las mismas celdas de «N Datos de gráficos»: no hay datos duplicados.
"""
import copy
import math
import re

from openpyxl.styles import Font
from openpyxl.utils import get_column_letter as L

from a2_excel_resultados import ANCLAS
from a2_graf_base import CM_COL, CM_FILA, COPIAS
from pies import SIMBOLO

COL_MIN = 12                                  # primera columna posible de los gráficos en las hojas -R
VAR_DE = {v: k for k, v in SIMBOLO.items()}   # símbolo (η_t) → nombre de la variable (eta_t)

# figura → ("R", estudio, parte | None, "parte" | "var" | "top")  o  ("T", id de la tabla)
MAPA = {"2.1-1": ("R", "2.1", "2.1", "var"),
        "2.2-1": ("R", "2.2", "2.2a", "parte"), "2.2-2": ("R", "2.2", "2.2b", "parte"),
        "2.3-1": ("R", "2.3", "2.3a", "var"), "2.3-2": ("R", "2.3", "2.3b", "parte"),
        "2.3-3": ("R", "2.3", "2.3c", "parte"), "2.3-4": ("R", "2.3", "2.3d", "parte"),
        "2.4-1": ("R", "2.4", None, "top"),
        "3.1-1": ("R", "3.1", None, "top"), "3.1-3": ("R", "3.1", "3.1", "var"),
        "3.1-2": ("T", "3.1-3"), "3.1-4": ("T", "3.1-3"),
        "3.2-1": ("R", "3.2", None, "top"), "3.2-2": ("R", "3.2", None, "top"),
        "3.2-4": ("R", "3.2", None, "top"), "3.2-3": ("T", "3.2-4"),
        "3.3-1": ("R", "3.3", None, "top"),
        "4.1-1": ("R", "4.1", "4.1a", "parte"), "4.1-2": ("R", "4.1", "4.1b", "parte"),
        "4.1-3": ("R", "4.1", "4.1c", "parte"), "4.1-4": ("R", "4.1", "4.1d", "parte"),
        "4.1-5": ("R", "4.1", "4.1e", "parte"),
        "4.2-1": ("R", "4.2", None, "top"), "4.3-1": ("T", "4.3-3"),
        "5.1-1": ("R", "5.1", None, "top"), "6.1-1": [("T", "6.1-3"), ("T", "4.3-4")],   # m6_comparacion.csv y m9_motor_real.csv
        "6.2-1": ("R", "6.2", "6.2c", "parte"), "6.2-2": ("R", "6.2", "6.2b", "var"),
        "6.2d": ("R", "6.2", "6.2d", "parte"), "S-1": ("T", "S-1")}


def _col_libre(ws, ancho):
    """Primera columna sin tabla ni texto desbordado (filas ≥ 6) en una hoja -R, más una de aire."""
    def pt(c):
        return (26 if c == 1 else 17 if c <= 12 else 8.43) * 5.25
    inicio = [0.0]
    for c in range(1, 80):
        inicio.append(inicio[-1] + pt(c))              # inicio[c-1] = borde izquierdo de la columna c
    col = max(COL_MIN, ancho + 2)
    for fila in ws.iter_rows(min_row=6):
        for c in fila:
            if isinstance(c.value, str) and len(c.value) > 12:
                fin = inicio[c.column - 1] + len(c.value) * (c.font.sz or 11) * 0.52
                k = next((j for j in range(c.column, 79) if inicio[j] >= fin), 79)
                col = max(col, k + 2)
    return col


def _titulo(ch):
    try:
        return "".join(r.t for p in ch.title.tx.rich.p for r in (p.r or []))
    except AttributeError:
        return ""


def _hoja_tabla(wb, tid):
    for ws in wb.worksheets:
        if ws.title.startswith(f"{tid} "):
            return ws
    raise KeyError(tid)


def _puntos(ch):
    """Puntos que dibuja un gráfico de dispersión (suma de sus series); None si no es de puntos."""
    if type(ch).__name__ != "ScatterChart":
        return None
    total = 0
    for s in ch.series:
        m = re.match(r"'(.+)'!\$([A-Z]+)\$(\d+)(?::\$[A-Z]+\$(\d+))?", s.yVal.numRef.f)
        total += int(m.group(4) or m.group(3)) - int(m.group(3)) + 1
    return total


AUDITORIA = []                                # (figura, panel, puntos del gráfico, filas de la tabla)


def colocar_copias(wb):
    """Coloca las copias; devuelve {hoja: número de gráficos añadidos}."""
    cursor, puestos, sin_destino, _cols_r = {}, {}, [], {}
    regs = []                                # una figura con datos de varias tablas va en cada una
    for reg in COPIAS:
        dests = MAPA.get(reg["fig"])
        if dests is None:
            sin_destino.append(reg["fig"] or reg["subtitulo"])
            continue
        for i, d in enumerate(dests if isinstance(dests, list) else [dests]):
            r = dict(reg, dest=d)
            if i:
                r["charts"] = [copy.deepcopy(c) for c in reg["charts"]]
            regs.append(r)
    for reg in regs:
        dest = reg["dest"]
        if dest[0] == "T":
            ws, col, modo, est, parte = _hoja_tabla(wb, dest[1]), 1, "top", None, None
            cursor.setdefault(ws.title, ws.max_row + 3)
            deseada = cursor[ws.title]
        else:
            _, est, parte, modo = dest
            anc = ANCLAS[est]
            ws = anc["ws"]
            col = _cols_r.setdefault(ws.title, _col_libre(ws, anc["ancho"]))
            cursor.setdefault(ws.title, anc["top"])
            deseada = anc["partes"].get(parte, anc["top"]) if modo == "parte" else anc["top"]
        if modo == "var":                    # el rótulo va justo encima del primer panel
            deseada = min([ANCLAS[est]["var"].get((parte, VAR_DE.get(_titulo(c), _titulo(c))), 10**6)
                           for c in reg["charts"]] + [10**6])
            deseada = anc["top"] if deseada == 10**6 else deseada - 1
        fila = max(cursor[ws.title], deseada)
        ws.cell(fila, col, reg["subtitulo"]).font = Font(bold=True, size=11)
        fila += 1
        filas_g = math.ceil(reg["alto"] / CM_FILA) + 1
        cols_g = math.ceil(reg["ancho"] / CM_COL) + 1
        for ch in reg["charts"]:
            v = VAR_DE.get(_titulo(ch), _titulo(ch))
            if est:
                a = ANCLAS[est]
                n_tab = (a["n_var"].get((parte, v)) if modo == "var" else
                         a["n_parte"].get(parte) if modo == "parte" else a["n_est"])
                AUDITORIA.append((reg["fig"], _titulo(ch)[:40], _puntos(ch), n_tab))
        if modo == "var":                    # un panel por variable, junto a su bloque de tabla
            for ch in reg["charts"]:
                v = VAR_DE.get(_titulo(ch), _titulo(ch))
                fila = max(fila, ANCLAS[est]["var"].get((parte, v), fila))
                ws.add_chart(ch, f"{L(col)}{fila}")
                fila += filas_g
        else:                                # rejilla de `por_fila` gráficos por fila
            for i, ch in enumerate(reg["charts"]):
                ws.add_chart(ch, f"{L(col + (i % reg['por_fila']) * cols_g)}"
                                 f"{fila + (i // reg['por_fila']) * filas_g}")
            fila += math.ceil(len(reg["charts"]) / reg["por_fila"]) * filas_g
        for linea in reg["pie"][1:] if reg["pie"][:1] and reg["pie"][0].startswith("Figura") else reg["pie"]:
            ws.cell(fila, col, linea).font = Font(size=9, italic=True)
            fila += 1
        cursor[ws.title] = fila + 2
        puestos[ws.title] = puestos.get(ws.title, 0) + len(reg["charts"])
    if sin_destino:
        raise RuntimeError(f"figuras sin hoja de destino: {sin_destino}")
    return puestos
