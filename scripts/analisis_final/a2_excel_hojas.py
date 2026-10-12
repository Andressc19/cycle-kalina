# -*- coding: utf-8 -*-
"""A2 · Hojas del libro Excel_kalina_final_final_final.xlsx (todas menos el Tablero).

Escrito por el director (Claude) el 2026-10-10 (OpenCode sin modelo disponible).
Formato §6.3: encabezado de condiciones (todas sus líneas, bloque por bloque); tabla después de una fila en blanco, con filtro;
notas al final; colores por clase; todo en español.
"""
import glob
import os

import numpy as np
import pandas as pd
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from a2_columnas import tipo_por_unidad
from a2_comun import COLOR_CLASE, RAIZ, TABLAS, fmt_excel
from codigos_texto import traducir
from nombres import GLOSARIO

FIG = os.path.join(RAIZ, "analisis_final", "03_figuras")
NEGRITA, TITULO = Font(bold=True), Font(bold=True, size=13)
GRIS = PatternFill("solid", fgColor="DDEBF7")
_FINO = Side(style="thin", color="A6A6A6")
BORDE = Border(left=_FINO, right=_FINO, top=_FINO, bottom=_FINO)
CLASES_TXT = {"KALINA": "KALINA", "CORREGIBLE": "CORREGIBLE", "Corregible": "CORREGIBLE",
              "INVIABLE": "INVIABLE", "Inviable": "INVIABLE", "NO_CONVERGIO": "NO_CONVERGIO",
              "No convergió": "NO_CONVERGIO", "N2": "N2"}
FIGURAS_DE = {"M02-N1": ["FIG-M02-01"], "M02-N2": ["FIG-M02-02"], "M02R": ["FIG-M02R"],
              "M02S": ["FIG-M02S"], "M03": ["FIG-M03"], "M04": ["FIG-M04-"], "M04c": ["FIG-M04c"],
              "M05": ["FIG-M05-S"], "M05c": ["FIG-M05c"], "M06": ["FIG-M06"], "M07": ["FIG-M07"],
              "M09": ["FIG-M09"], "M10": ["FIG-M10-"], "M10b": ["FIG-M10-01"], "G": ["FIG-G"],
              "ref": ["Figura_"]}


def cuadricula(ws, r1, c1, r2, c2):
    """Bordes finos en todas las celdas del rango (pedido del usuario, 2026-10-11)."""
    for fila in ws.iter_rows(min_row=r1, max_row=r2, min_col=c1, max_col=c2):
        for c in fila:
            c.border = BORDE


def nombre_hoja(txt):
    for c in '[]:*?/\\':
        txt = txt.replace(c, "_")
    return txt[:31]


def _anchos(ws, n_col, fila_ini, fila_fin):
    for j in range(1, n_col + 1):
        largo = max((len(str(ws.cell(i, j).value or "")) for i in range(fila_ini, fila_fin + 1)),
                    default=8)
        ws.column_dimensions[get_column_letter(j)].width = min(max(10, largo + 2), 48)


def hoja_tabla(wb, meta, condiciones):
    """Una tabla TAB-* con encabezado (todas las líneas de condiciones), cuerpo y notas."""
    carpeta = os.path.join(TABLAS, str(meta["carpeta_nueva"]))       # nomenclatura nueva
    df = pd.read_csv(os.path.join(carpeta, meta["archivo"] + ".csv"))
    ws = wb.create_sheet(nombre_hoja(meta["hoja"]))
    ws["A1"] = f"Tabla {meta['id_nuevo']} · {meta['titulo_nuevo']}"; ws["A1"].font = TITULO
    ws["A2"] = meta["cabecera"]
    ws["A3"] = f"Punto base: {meta['punto_base_nuevo']}" if meta["carpeta_nueva"] != "S" else ""
    lineas = condiciones.splitlines() if condiciones and meta["carpeta_nueva"] != "S" else         ["Síntesis: cada fila es un estudio o una parte con sus propios datos fijos."]
    # Todas las líneas: con varios bloques, cortar a 2 mostraba solo el primero.
    for k, linea in enumerate(lineas):
        ws.cell(4 + k, 1, linea).font = NEGRITA if linea.startswith("Parte ") else Font()
    r_filtro = 4 + len(lineas)
    ws.cell(r_filtro, 1, f"Fuente: {traducir(meta['fuente'])}")
    for r in range(1, r_filtro + 1):
        ws.cell(r, 1).alignment = Alignment(wrap_text=False)
    h = r_filtro + 2                                  # fila de encabezados de la tabla
    for j, col in enumerate(df.columns, 1):
        c = ws.cell(h, j, col); c.font = NEGRITA; c.fill = GRIS
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for i, fila in enumerate(df.itertuples(index=False), h + 1):
        for j, (col, v) in enumerate(zip(df.columns, fila), 1):
            v = None if (isinstance(v, float) and pd.isna(v)) else v
            if isinstance(v, (bool, np.bool_)):
                v = "sí" if v else "no"              # con formato numérico salía "#########"
            if isinstance(v, float) and not float(v).is_integer():
                v = float(f"{v:.5g}")              # 10.87115255 → 10.871
            cel = ws.cell(i, j, v)
            if isinstance(v, str) and v[:1] in "=+-@" and v[:1]:
                cel.data_type = "s"               # texto, no fórmula (evita #¿NOMBRE?)
            fmt = fmt_excel(tipo_por_unidad(col), v) if isinstance(v, (int, float)) else None
            if fmt:
                cel.number_format = fmt
            if isinstance(v, str) and v.startswith(("barrido: ", "fijo: ")):   # matriz G02
                cel.fill = PatternFill("solid", fgColor="90CAF9" if v.startswith("barrido")
                                       else "E0E0E0" if v == "fijo: profesor" else "FFB74D")
            clase = CLASES_TXT.get(str(v)) if isinstance(v, str) else None
            if clase:
                cel.fill = PatternFill("solid", fgColor=COLOR_CLASE[clase])
    fin = h + len(df)
    if len(df):
        ws.auto_filter.ref = f"A{h}:{get_column_letter(len(df.columns))}{fin}"
    cuadricula(ws, h, 1, fin, max(len(df.columns), 1))
    nota = "Notas: " + str(meta["notas_nuevas"])
    ws.cell(fin + 2, 1, nota).font = Font(italic=True)
    # sin paneles inmovilizados: el usuario prefiere el desplazamiento normal (2026-10-11)
    _anchos(ws, len(df.columns), h, fin)
    return ws.title


def hoja_figuras(wb, clave, titulo):
    """Hoja con cada PNG de Matplotlib (ancho ~16 cm) y su pie debajo. Desde el 2026-10-11 hay
    una sola, al final del libro (clave "ref"): los gráficos del libro son nativos de Excel."""
    ws = wb.create_sheet(nombre_hoja(f"{clave}_figuras") if clave != "ref"
                         else "Figuras_matplotlib_ref")
    ws["A1"] = titulo; ws["A1"].font = TITULO
    pngs = []
    for pref in FIGURAS_DE.get(clave, []):
        pngs += sorted(glob.glob(os.path.join(FIG, "*", pref + "*.png")))
    if not pngs:
        ws["A3"] = "Este caso no lleva figura (ver su ficha)."
        return ws.title
    fila = 3
    for p in dict.fromkeys(pngs):
        img = XLImage(p)
        esc = 605 / img.width
        img.width, img.height = int(img.width * esc), int(img.height * esc)
        ws.add_image(img, f"A{fila}")
        fila += int(img.height / 20) + 2
        pie = p[:-4] + ".pie.md"
        texto = open(pie, encoding="utf-8").read() if os.path.exists(pie) else ""
        for linea in texto.splitlines():
            ws.cell(fila, 1, linea).font = Font(size=9, italic=True)
            fila += 1
        fila += 2
    ws.column_dimensions["A"].width = 120
    return ws.title


def hoja_leeme(wb):
    ws = wb.active; ws.title = "LEEME"
    filas = ["Ciclo Kalina KCS-11 — resultados de la campaña final de barridos (M1–M10b)", "",
             "Para qué es este libro: presentar las tablas y gráficos del análisis, ordenados por pregunta.",
             "Cómo se numera: pregunta (1–6) → estudio (2.4) → parte (2.4a). Ver la hoja Nomenclatura. "
             "Solo se comparan puntos de la misma parte.",
             "Cómo leer cada hoja: las primeras filas dicen qué caso es, el punto base, los datos fijos y "
             "barridos (parte por parte si hay varias), el filtro y la fuente; la tabla empieza después "
             "de una fila en blanco (con filtros); las notas van al final.",
             "Hoja Tablero: elija un estudio o una parte en la celda B2 y se actualizan las cifras, las condiciones y "
             "los gráficos.",
             "Hojas 'N Gráficos': gráficos nativos de Excel de cada pregunta (pase el cursor para ver los "
             "valores); cada uno lleva arriba su ID y título y abajo su pie con los datos fijos.",
             "Hoja Figuras_matplotlib_ref (al final): las 30 figuras de Matplotlib, solo como referencia.",
             "", "Glosario:"]
    filas += [f"• {k}: {v}" for k, v in GLOSARIO.items()]
    filas += ["", "Colores de clase:"]
    for i, t in enumerate(filas, 1):
        ws.cell(i, 1, t)
    ws["A1"].font = TITULO
    r = len(filas) + 1
    for clase, txt in (("KALINA", "KALINA: cumple todos los criterios"),
                       ("CORREGIBLE", "CORREGIBLE: falla un criterio operativo (casi siempre O2)"),
                       ("INVIABLE", "INVIABLE: falla un criterio estructural (O3, O5…)"),
                       ("NO_CONVERGIO", "NO_CONVERGIO: el modelo no encontró un ciclo"),
                       ("N2", "Falso convergido: el solver dijo que convergió, pero no es un ciclo real (no cuenta)")):
        c = ws.cell(r, 1, txt); c.fill = PatternFill("solid", fgColor=COLOR_CLASE[clase]); r += 1
    ws.column_dimensions["A"].width = 130


def hoja_profesor(wb):
    ws = wb.create_sheet("Punto_profesor")
    datos = [("P_alta", "kPa", 3000, 3000), ("P_baja", "kPa", 400, 563.64), ("x_b", "–", 0.5, 0.5),
             ("T_fuente", "K", 623.15, 403), ("T_sumidero", "K", 300.032917, 300.032917),
             ("η_t", "–", 0.85, 0.85), ("η_p", "–", 0.75, 0.75), ("ε_HRVG", "–", 0.85, 0.85),
             ("ε_reg", "–", 0.75, 0.75), ("ε_cond", "–", 0.80, 0.80), ("m_b", "kg/s", 1, 1)]
    ws.append(["Dato", "Unidad", "Punto del profesor (PP)", "Punto ancla A*"])
    for d in datos:
        ws.append(list(d))
    for c in ws[1]:
        c.font = NEGRITA; c.fill = GRIS
    cuadricula(ws, 1, 1, 1 + len(datos), 4)
    ws.append([]); ws.append(["A* = PP con solo dos cambios (T_fuente y P_baja). Es KALINA, η = 8.48 %. "
                              "PP no converge (M1)."])
    _anchos(ws, 4, 1, 12)


def hoja_df(wb, nombre, df, oculta=False):
    """Hoja con un DataFrame plano como tabla con filtros (Datos, Condiciones, Bloques)."""
    ws = wb.create_sheet(nombre)
    ws.append(list(df.columns))
    for fila in df.itertuples(index=False):
        ws.append([None if (isinstance(v, float) and pd.isna(v)) else v for v in fila])
    for c in ws[1]:
        c.font = NEGRITA; c.fill = GRIS
    ws.auto_filter.ref = ws.dimensions
    cuadricula(ws, 1, 1, ws.max_row, ws.max_column)
    if oculta:
        ws.sheet_state = "hidden"
    return ws
