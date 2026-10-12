# -*- coding: utf-8 -*-
"""A2 · Hoja Tablero (§4B.3): se elige un bloque en B2 y cambian KPIs, condiciones y gráficos.

Escrito por el director (Claude) el 2026-10-10 (OpenCode sin modelo disponible).
Solo fórmulas clásicas (INDEX/MATCH/IF/NA), sin macros ni FILTRAR: funciona en cualquier Excel.
Los valores por bloque se precalculan en las hojas ocultas Aux_tablero y Aux_puntos.
"""
import pandas as pd
from openpyxl.chart import BarChart, Reference, ScatterChart, Series
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

import pies
from a2_excel_hojas import cuadricula
from nombres import NOMBRE

PRIORIDAD = ["P_alta", "P_baja", "T_fuente", "x_b", "eps_cond", "T_sumidero", "eps_hrvg",
             "eps_reg", "eta_t", "eta_p"]
SERIES = [("K", "KALINA", ["KALINA"], "2E7D32", "circle"),
          ("C", "Corregible", ["CORREGIBLE"], "EF8A17", "triangle"),
          ("I", "Inviable", ["INVIABLE"], "C9A227", "square"),
          ("N", "No convergió o N2", ["NO_CONVERGIO", "N2"], "C62828", "x")]
CLASES = [("KALINA", "KALINA"), ("Corregible", "CORREGIBLE"), ("Inviable", "INVIABLE"),
          ("No convergió", "NO_CONVERGIO"), ("Falso convergido (N2)", "N2")]
CRIT = [("O1 título", "viol_O1"), ("O2 cavitación", "viol_O2"), ("O3 trabajo neto", "viol_O3"),
        ("O5 separador", "viol_O5")]
MAXP, F_ini = 900, 5          # puntos por serie y fila inicial de los datos del gráfico
KPI = ["Corridas", "Genuinas", "KALINA", "KALINA confiables", "η máx. KALINA [%]",
       "Criterio que más bloquea"]


def _cond_texto(c, var):
    g = c[c["variable"] == var]
    if g.empty or (g["estado"] == "sin_dato").all():
        return "sin dato"
    u = "" if g["unidad"].iloc[0] in ("–", None) or pd.isna(g["unidad"].iloc[0]) else f" {g['unidad'].iloc[0]}"
    if (g["es_oat"] == "sí").all() and (g["estado"] == "barrido").all():
        return f"Una a la vez (punto base {g['valor_base'].iloc[0]}{u}; rango {g['min'].min():g}–{g['max'].max():g})"
    if (g["estado"] == "fijo").all() and g["valor_fijo"].nunique() == 1:
        return f"Fija: {g['valor_fijo'].iloc[0]}{u}"
    if (g["estado"] == "barrido").all():
        return f"Barrida: {g['min'].min():g} – {g['max'].max():g}{u}"
    return "Distinta según el subcaso (ver hoja de condiciones del caso)"


def _ejes(c):
    barr = set(c[c["estado"] == "barrido"]["variable"])
    v = [x for x in PRIORIDAD if x in barr]
    if v and (c["es_oat"] == "sí").all():      # una a la vez: η contra la 1.ª variable
        return v[0], "eta_pct"
    if len(v) >= 2:
        return v[0], v[1]
    return (v[0] if v else "P_baja"), "eta_pct"


def construir(wb, d, cond, bloques):
    aux, aup = wb.create_sheet("Aux_tablero"), wb.create_sheet("Aux_puntos")
    aux["A1"] = "clave"
    filas = ["corridas", "genuinas", "kalina", "confiables", "eta_max", "criterio", "eje_x", "eje_y",
             "pie", "descripcion"] + [f"clase_{k}" for _, k in CLASES] + [f"crit_{k}" for _, k in CRIT] \
        + [f"cond_{v}" for v in cond["variable"].unique()]
    for i, f in enumerate(filas, 2):
        aux.cell(i, 1, f)
    variables = list(cond["variable"].unique())
    for j, b in enumerate(bloques, 2):
        g = d[d["bloque_director"] == b]
        cb = cond[cond["bloque_director"] == b]
        subs = list(dict.fromkeys(cb["id_subcaso"]))
        gen = g[g["genuino"].astype(bool)]
        nok = gen[gen["clasificacion"] != "KALINA"]
        crit = {k: int((pd.to_numeric(nok[k], errors="coerce") > 0).sum()) for _, k in CRIT}
        kal = g[g["clasificacion"] == "KALINA"]
        x, y = _ejes(cb)
        val = [len(g), len(gen), len(kal), int(g["kalina_confiable"].astype(bool).sum()),
               round(float(kal["eta_pct"].max()), 2) if len(kal) else None,
               (max(crit, key=crit.get).replace("viol_", "") if any(crit.values()) else "—"),
               NOMBRE.get(x, x), NOMBRE.get(y, y), pies.condiciones_texto(subs).replace("\n", "  ·  "),
               ", ".join(subs)] + [int((g["clase_trazo"] == k).sum()) for _, k in CLASES] \
            + [crit[k] for _, k in CRIT] + [_cond_texto(cb, v) for v in variables]
        aux.cell(1, j, b)
        for i, v in enumerate(val, 2):
            aux.cell(i, j, v)
        for s, (cod, _, clases, _, _) in enumerate(SERIES):
            sg = g[g["clase_trazo"].isin(clases) & ~g["omitido"].astype(bool)].dropna(subset=[x, y])
            if (cb["es_oat"] == "sí").all():           # solo los puntos donde se mueve x
                sg = sg[sg["variable"] == x]
            sg = sg.head(MAXP)
            cx = 1 + (j - 2) * 8 + 2 * s
            aup.cell(1, cx, f"{b}|x{cod}"); aup.cell(1, cx + 1, f"{b}|y{cod}")
            for r, (vx, vy) in enumerate(zip(sg[x], sg[y]), 2):
                aup.cell(r, cx, float(vx)); aup.cell(r, cx + 1, float(vy))
    aux.sheet_state = aup.sheet_state = "hidden"
    _hoja(wb, bloques, filas, variables, len(bloques), aup.max_column)


def _idx(fila, ultima):
    return (f"=INDEX(Aux_tablero!$B${fila}:${L(ultima)}${fila},"
            f"MATCH($B$2,Aux_tablero!$B$1:${L(ultima)}$1,0))")


def _hoja(wb, bloques, filas, variables, nb, ncol_p):
    ws = wb.create_sheet("Tablero", 1)
    ult = nb + 1
    ws["A1"] = "Tablero de resultados — elija un bloque en B2"; ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "Bloque:"; ws["A2"].font = Font(bold=True)
    ws["B2"] = bloques[0]; ws["B2"].fill = PatternFill("solid", fgColor="FFF2CC")
    ws["B2"].font = Font(bold=True, size=12)
    dv = DataValidation(type="list", formula1=f"=Aux_tablero!$B$1:${L(ult)}$1", allow_blank=False)
    ws.add_data_validation(dv); dv.add("B2")
    ws["C2"] = "=" + _idx(filas.index("descripcion") + 2, ult)[1:]
    for j, k in enumerate(KPI, 4):            # cifras en D..I, columnas propias (no cortan)
        c = ws.cell(4, j, k); c.font = Font(bold=True, size=9); c.alignment = Alignment(wrap_text=True)
        c = ws.cell(5, j, _idx(j - 2, ult)); c.font = Font(bold=True, size=14, color="2E7D32")
        ws.column_dimensions[L(j)].width = 16
    ws["K4"], ws["K5"] = "Eje X:", "Eje Y:"
    ws["L4"], ws["L5"] = _idx(filas.index("eje_x") + 2, ult), _idx(filas.index("eje_y") + 2, ult)
    ws["A7"] = "Condiciones de este bloque"; ws["A7"].font = Font(bold=True)
    for i, v in enumerate(variables):
        ws.cell(8 + i, 1, NOMBRE.get(v, v))
        ws.cell(8 + i, 2, _idx(filas.index(f"cond_{v}") + 2, ult))
    fp = 9 + len(variables)
    cuadricula(ws, 4, 4, 5, 3 + len(KPI)); cuadricula(ws, 8, 1, 7 + len(variables), 2)
    ws.cell(fp, 1, "Pie del bloque:").font = Font(bold=True)
    c = ws.cell(fp + 1, 1, _idx(filas.index("pie") + 2, ult))
    c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=fp + 1, start_column=1, end_row=fp + 6, end_column=2)   # pie completo
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 44, 62
    ws.column_dimensions["C"].width = 2
    # datos auxiliares de los gráficos (columnas ocultas desde Z)
    for i, (etq, k) in enumerate(CLASES):
        ws.cell(F_ini + i, 26, etq); ws.cell(F_ini + i, 27, _idx(filas.index(f"clase_{k}") + 2, ult))
    for i, (etq, k) in enumerate(CRIT):
        ws.cell(F_ini + 7 + i, 26, etq); ws.cell(F_ini + 7 + i, 27, _idx(filas.index(f"crit_{k}") + 2, ult))
    rng = f"Aux_puntos!$A$2:${L(ncol_p)}${MAXP + 1}"
    for s, (cod, _, _, _, _) in enumerate(SERIES):
        for e, eje in enumerate("xy"):
            col = 29 + 2 * s + e
            ws.cell(F_ini - 1, col, f"{eje}{cod}")
            for r in range(MAXP):
                m = f'MATCH($B$2&"|{eje}{cod}",Aux_puntos!$1:$1,0)'
                ws.cell(F_ini + r, col, f'=IFERROR(IF(INDEX({rng},{r + 1},{m})="",NA(),'
                                        f'INDEX({rng},{r + 1},{m})),NA())')
    for col in range(26, 29 + 8):
        ws.column_dimensions[L(col)].hidden = True
    _graficos(ws)


def _graficos(ws):
    b1 = BarChart(); b1.title = "Puntos por clase"; b1.y_axis.title = "Número de puntos"
    b1.add_data(Reference(ws, min_col=27, min_row=F_ini, max_row=F_ini + 4), titles_from_data=False)
    b1.set_categories(Reference(ws, min_col=26, min_row=F_ini, max_row=F_ini + 4))
    b1.legend = None; b1.height, b1.width = 7, 13
    from openpyxl.chart.marker import DataPoint      # una barra por clase, con su color
    for i, col in enumerate(["2E7D32", "EF8A17", "C9A227", "C62828", "9E9E9E"]):
        pt = DataPoint(idx=i); pt.graphicalProperties.solidFill = col
        b1.series[0].dPt.append(pt)
    ws.add_chart(b1, "D8")
    b2 = BarChart(); b2.title = "Criterios que bloquean (genuinos no KALINA)"
    b2.add_data(Reference(ws, min_col=27, min_row=F_ini + 7, max_row=F_ini + 10), titles_from_data=False)
    b2.set_categories(Reference(ws, min_col=26, min_row=F_ini + 7, max_row=F_ini + 10))
    b2.legend = None; b2.height, b2.width = 7, 13
    b2.series[0].graphicalProperties.solidFill = "C62828"
    ws.add_chart(b2, "D24")
    sc = ScatterChart(); sc.title = "Dispersión del bloque (ejes en L4:L5)"
    sc.style = 13; sc.height, sc.width = 14, 18
    sc.x_axis.title, sc.y_axis.title = "Eje X (ver L4)", "Eje Y (ver L5)"
    for s, (cod, etq, _, color, simb) in enumerate(SERIES):
        xs = Reference(ws, min_col=29 + 2 * s, min_row=F_ini, max_row=F_ini + MAXP - 1)
        ys = Reference(ws, min_col=30 + 2 * s, min_row=F_ini, max_row=F_ini + MAXP - 1)
        ser = Series(ys, xs, title=etq)
        ser.marker.symbol, ser.marker.size = simb, 6
        ser.marker.graphicalProperties.solidFill = color
        ser.marker.graphicalProperties.line.solidFill = color
        ser.graphicalProperties.line.noFill = True
        sc.series.append(ser)
    for g in (b1, b2, sc):                    # columnas auxiliares ocultas: graficarlas igual
        g.visible_cells_only = False
        g.x_axis.delete = g.y_axis.delete = False
    ws.add_chart(sc, "K8")
