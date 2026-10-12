# -*- coding: utf-8 -*-
"""A2 · Arma datos/Excel_kalina_final_final_final.xlsx (§6.3). Se puede volver a correr.

Escrito por el director (Claude) el 2026-10-10 (OpenCode sin modelo disponible).
No abre para escribir datos/Excel_final_kalina.xlsx: comprueba su hash antes y después.
"""
import hashlib
import os

import pandas as pd
from openpyxl import Workbook

import pies
from a2_comun import EXCEL_VIEJO, INV, LIBRO, MNN, leer_canon, registro
from a2_excel_hojas import (cuadricula, hoja_df, hoja_figuras, hoja_leeme, hoja_profesor,
                            hoja_tabla)
from a2_graf_casos_b import construir_graficos
from a2_excel_nomenclatura import hoja_nomenclatura
import codigos
from nombres import ETIQUETA_CLASE
from a2_excel_tablero import construir

ORIGEN = {"estudio": "Estudio", "parte_estudio": "Parte", "parte_nombre": "Nombre de la parte",
          "parte": "Parte", "id_caso": "Caso de origen", "id_subcaso": "Subcaso de origen",
          "bloque_director": "Bloque de origen"}
COLS_DATOS = ["id_punto", "estudio", "parte_estudio", "parte_nombre", "id_caso", "id_subcaso", "bloque_director", "tipo_punto", "clase_trazo",
              "clasificacion", "genuino", "kalina_confiable", "fuera_rango", "omitido", "T_fuente",
              "T_sumidero", "P_alta", "P_baja", "x_b", "eta_t", "eta_p", "eps_hrvg", "eps_reg",
              "eps_cond", "T_amb_diseno", "eta_pct", "Wnet", "dP", "rP", "codigos_falla", "variable"]


def sha(p):
    return hashlib.sha256(open(p, "rb").read()).hexdigest()


def _clave(id_nuevo):
    """Orden de las tablas: S primero, luego 1.1-1, 1.1-2 … 6.2-5."""
    est, n = str(id_nuevo).split("-")
    return (0 if est == "S" else 1, [int(x) for x in est.split(".") if x.isdigit()], int(n))


def main():
    h0 = sha(EXCEL_VIEJO)
    reg = registro()
    cond = pd.read_csv(os.path.join(INV, "condiciones_por_subcaso.csv"))
    d = leer_canon()
    wb = Workbook()
    hoja_leeme(wb)
    indice = []                                      # (hoja, estudio, partes, qué muestra)
    hoja_profesor(wb)
    indice.append(("Punto_profesor", "—", "—", "Datos del profesor y punto ancla A*"))
    reg = reg.assign(_k=reg["id_nuevo"].map(_clave)).sort_values("_k")
    hojas_graf = construir_graficos(wb)                              # crea las 6 hojas de gráficos al final
    for _, m in reg.iterrows():
        subs = [s for s in str(m["subcasos"]).split("; ") if s]
        cond_txt = "" if m["carpeta_nueva"] == "S" else pies.condiciones_por_bloque(subs)
        h = hoja_tabla(wb, m, cond_txt)
        est = str(m["carpeta_nueva"])
        indice.append((h, est, ", ".join(dict.fromkeys(codigos.parte(s)[0] for s in subs)) or "—",
                       m["titulo_nuevo"]))
    for h, est, que in hojas_graf:
        indice.append((h, est, "—", que))
    _ordenar(wb)
    ws = wb.create_sheet("A.1 Anexo fuera de rango")
    ws["A1"] = "ANEXO A.1 — equipos fuera del rango realista (no entra en las conclusiones)"
    from openpyxl.styles import Font, PatternFill
    ws["A1"].fill = PatternFill("solid", fgColor="F8B4B4"); ws["A1"].font = Font(bold=True, size=14)
    m8 = d[d["id_caso"] == "KF-M08"]["clase_trazo"].value_counts()
    ws["A3"] = ("Estudio A.1 (código de origen M08): η y ε de 0.20 a 0.99, fuera de [0.40, 0.95]. "
                "No es un resultado de diseño.")
    for i, (k, v) in enumerate(m8.items(), 5):
        ws.cell(i, 1, ETIQUETA_CLASE.get(k, k)); ws.cell(i, 2, int(v))
    cuadricula(ws, 5, 1, 4 + len(m8), 2)
    indice.append((ws.title, "A.1", "—", "Curiosidad: equipos fuera del rango realista"))
    datos = d[d["id_caso"] != "KF-M08"][COLS_DATOS]
    hoja_df(wb, "Datos", datos.rename(columns=ORIGEN))
    cond_p = cond.assign(estudio=cond["id_caso"].map(codigos.estudio),
                         parte=cond["id_subcaso"].map(lambda s: codigos.parte(s)[0]))
    hoja_df(wb, "Condiciones", cond_p[["estudio", "parte"] + list(cond.columns)].rename(columns=ORIGEN))
    bl = pd.read_csv(os.path.join(INV, "bloques.csv"))
    hoja_df(wb, "Partes", bl.assign(parte=bl["bloque_director"].map(codigos.bloque))
            [["parte"] + list(bl.columns)].rename(columns=ORIGEN))
    bloques = [b for b in dict.fromkeys(datos["bloque_director"]) if b != "BX"]
    construir(wb, datos.assign(viol_O1=d["viol_O1"], viol_O2=d["viol_O2"], viol_O3=d["viol_O3"],
                               viol_O5=d["viol_O5"]), cond[cond["bloque_director"] != "BX"], bloques)
    hoja_figuras(wb, "ref", "Las 30 figuras de Matplotlib (solo referencia; los gráficos del libro "
                            "son nativos de Excel)")
    indice.append(("Figuras_matplotlib_ref", "todos", "—", "las 30 figuras de Matplotlib, de referencia"))
    wb.move_sheet("Aux_graficos", offset=len(wb.sheetnames))
    hoja_nomenclatura(wb, 1)
    wi = wb.create_sheet("Indice", 3)
    wi.append(["Hoja", "Estudio", "Parte(s)", "Qué muestra"])
    for h, c, b, q in indice:
        wi.append([h, c, b, q])
        wi.cell(wi.max_row, 1).hyperlink = f"#'{h}'!A1"
        wi.cell(wi.max_row, 1).style = "Hyperlink"
    cuadricula(wi, 1, 1, wi.max_row, 4)
    for col, w in zip("ABCD", (34, 10, 40, 70)):
        wi.column_dimensions[col].width = w
    os.makedirs(os.path.dirname(LIBRO), exist_ok=True)
    wb.save(LIBRO)
    h1 = sha(EXCEL_VIEJO)
    print(f"Libro: {LIBRO} · {len(wb.sheetnames)} hojas")
    print(f"Excel_final_kalina.xlsx sin cambios: {'OK' if h0 == h1 else 'FALLA'} ({h0[:12]})")


def _ordenar(wb):
    """Tablas de cada pregunta seguidas de su hoja de gráficos (S, 1, 2 … 6)."""
    fijas = [ws for ws in wb._sheets if ws.title in ("LEEME", "Punto_profesor")]
    resto = [ws for ws in wb._sheets if ws not in fijas]

    def clave(ws):
        t = ws.title
        p = t.split(" ")[0].split(".")[0].split("-")[0]
        orden = {"S": 0, "1": 1, "2": 2, "3": 3, "4": 4, "5": 5, "6": 6}.get(p, 9)
        return (orden, 1 if "Gráficos" in t else 0, t)
    wb._sheets = fijas + sorted(resto, key=clave)


if __name__ == "__main__":
    main()
