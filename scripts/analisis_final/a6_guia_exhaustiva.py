# -*- coding: utf-8 -*-
"""A6 · Genera la guía de estudio exhaustiva del análisis de sensibilidad (Claude, 2026-10-12).

Lee la plantilla `a6_guia_plantilla.md` (la narrativa, escrita a mano) y sustituye marcadores por
contenido calculado desde los CSV de `02_tablas`, los pies de `03_figuras` y el propio libro Excel,
para que ninguna cifra de tablas, figuras u hojas se transcriba a mano.

Marcadores:
  {{T id}}                          tabla completa (id = 3.2-3)
  {{T id | cols=a;b;c | n=12 | filtro=Columna=valor}}
  {{COLS id}}                       definición de cada columna de la tabla
  {{P id}}                          pie de la figura (título, datos fijos, filtro, n)
  {{HOJAS}} {{CAT_TABLAS}} {{CAT_FIGURAS}}   catálogos automáticos
Uso (desde la raíz del repo):  python scripts/analisis_final/a6_guia_exhaustiva.py
"""
import glob
import os
import re
import sys

import numpy as np
import pandas as pd

RAIZ = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
AF = os.path.join(RAIZ, "analisis_final")
TAB = os.path.join(AF, "02_tablas")
FIG = os.path.join(AF, "03_figuras")
PLANTILLA = os.path.join(os.path.dirname(os.path.abspath(__file__)), "a6_guia_plantilla.md")
SALIDA = os.path.join(AF, "GUIA_DE_ESTUDIO_EXHAUSTIVA_ANALISIS_SENSIBILIDAD.md")
EXCEL = os.path.join(RAIZ, "datos", "Excel_kalina_final_final_final.xlsx")

# ---------------------------------------------------------------- definiciones de columnas
D = {
    "Parte": "Subconjunto del estudio con sus propios datos fijos (p. ej. 2.4a). Solo se comparan puntos de la misma parte.",
    "Variable": "Dato de entrada (o del modelo) al que se refiere la fila.",
    "Estado": "Papel de la variable en la parte: «Una a la vez» (se mueve ella sola, con las demás fijas), «Barrida» (se mueve junto con otras), «Fija» (queda en un valor) o «Sin dato» (no aplica).",
    "Valor en el punto base": "Valor de la variable en el punto base de la parte (PP en 2.1; A* en 3.1).",
    "Valor o rango": "Valor fijo, o rango [mínimo – máximo] barrido.",
    "n valores": "Cantidad de valores distintos probados de esa variable en la parte.",
    "¿Igual al profesor?": "«sí» si el valor fijo coincide con el del enunciado.",
    "Unidad": "Unidad de la variable.",
    "Corridas": "Puntos calculados en la parte.",
    "Omitidos": "Puntos que el modelo no llegó a evaluar; no aparecen ni en las tablas de resultados ni en los gráficos.",
    "Genuinos": "Puntos que el modelo resolvió de verdad: KALINA + Corregibles + Inviables.",
    "KALINA": "Puntos que cumplen todos los criterios.",
    "KALINA confiables": "KALINA dentro del rango validado de propiedades (IAPWS G4-01) y con η y ε entre 0.40 y 0.95.",
    "Corregibles": "Puntos que fallan un criterio operativo (casi siempre O2, cavitación).",
    "Inviables": "Puntos que fallan un criterio estructural (O3, O5…).",
    "No convergidos": "Puntos en que el modelo no encontró un ciclo.",
    "Falsos convergidos": "El solver dijo que convergió pero el clasificador lo descartó (solución espuria); no cuentan como genuinos.",
    "Fuera de rango": "Puntos con estados fuera del rango validado de propiedades (p. ej. líquido por encima de 420 K); su η se toma con cautela.",
    # 1.1-3
    "Grupo": "Qué se movió en ese punto (base, P_alta, x_b, T_sumidero); en 3.2-5 y 4.1-4, el grupo comparado (KALINA o no KALINA).",
    "Valor": "Valor de la variable del grupo.",
    "Clase (antes)": "Clase que tuvo el punto en la Fase 1 de septiembre.",
    "Clase (ahora)": "Clase con el código actual.",
    "Misma clase": "«sí» si ambas clases coinciden.",
    "η (antes) [%]": "Eficiencia de la Fase 1 (vacía si no convergió).",
    "η (ahora) [%]": "Eficiencia con el código actual (vacía si no convergió).",
    "Δη [pp]": "Diferencia de eficiencia entre las dos comparaciones, en puntos porcentuales.",
    "Mensaje (antes)": "Texto con que el solver explicó el fallo en la Fase 1.",
    "Mensaje (ahora)": "Texto con que el solver explica el fallo ahora.",
    # 2.1-3
    "Tipo": "«diseño» (dato de diseño del ciclo) o «sitio» (condición del lugar: T_fuente, T_sumidero).",
    "Convergen": "Puntos de la variable en que el solver reportó convergencia (incluye los falsos convergidos).",
    "Valor convergido más cercano al PP": "Valor de la variable más cercano al del profesor entre los puntos que convergen.",
    "Valor KALINA más cercano al PP": "Valor de la variable más cercano al del profesor entre los KALINA (vacío si no hay).",
    "d normalizado": "Distancia al profesor del KALINA más cercano: Σ|v − v_prof| / (b − a). Vacío si no hay KALINA.",
    "D mínimo": "Menor índice de cercanía D de la fila (0 = cumple todos los criterios operativos).",
    "Valor de D mínimo": "Valor de la variable en el que D es mínimo.",
    "Criterios bloqueantes frecuentes": "Códigos de restricción que más fallan entre los puntos de la fila, con su conteo entre paréntesis (ver el glosario de códigos).",
    # 2.2-3
    "Ruta": "A = conserva la fuente del profesor (ε_HRVG × P_baja); B = baja la fuente (T_fuente × P_baja).",
    "Variable 1": "Primera variable movida en la fila.",
    "Valor 1": "Valor de la primera variable.",
    "Variable 2": "Segunda variable movida en la fila.",
    "P_baja mínima KALINA [kPa]": "Menor P_baja con la que el punto es KALINA.",
    "η mínima KALINA [%]": "η del punto KALINA de P_baja mínima (no es la menor η de la fila).",
    "Fuera de rango en la mínima KALINA": "Si el punto de P_baja mínima está fuera del rango validado.",
    "P_baja mínima con O2 [kPa]": "Menor P_baja con margen de cavitación O2 ≥ 0 (por bisección).",
    "P_baja del mejor D [kPa]": "P_baja en la que D es mínimo.",
    # 2.3-3
    "d menor": "Menor distancia d al profesor entre los KALINA de la fila.",
    "Valor 1 del d menor": "Valor de la variable 1 en el KALINA más cercano al profesor.",
    "Valor 2 del d menor": "Valor de la variable 2 en el KALINA más cercano al profesor.",
    "η en el d menor [%]": "η de ese punto.",
    "η mayor [%]": "Mayor η entre los KALINA de la fila.",
    "Valor 1 de la η mayor": "Valor de la variable 1 en el KALINA de mayor η.",
    "Valor 2 de la η mayor": "Valor de la variable 2 en el KALINA de mayor η.",
    "Bloqueantes": "Criterios que fallan entre los genuinos de la fila, con su conteo.",
    # 2.4-3
    "Base": "Base del barrido: B1 = 2.4a (fuente 403 K), B2 = 2.4b (420 K), B3 = 2.4c (ε_HRVG 0.40), B4 = 2.4d (profesor exacto).",
    "B verificado en la mínima KALINA": "Si el punto de P_baja mínima pasó la verificación B (el motor real recalcula h4s; marca «no verificado» si |Δh4s| > 2.0 kJ/kg).",
    "η máxima KALINA [%]": "Mayor η entre los KALINA de la fila.",
    "P_baja en la η máxima KALINA [kPa]": "P_baja en la que se da esa η máxima.",
    "Bloqueantes (genuinos)": "Clases de los genuinos que no son KALINA, con su conteo.",
    "Códigos de bloqueo": "Criterios que fallan en esos genuinos.",
    # 3.2-3
    "Separación": "Cómo se partió la muestra en dos grupos: «KALINA» (KALINA contra el resto), «converge», «cumple_O2» o «cumple_O1».",
    "n en B (KALINA)": "Puntos del grupo que cumple (B).",
    "n en B̄ (no KALINA)": "Puntos del grupo que no cumple (B̄), incluidos los que no convergieron.",
    "n excluidos": "Puntos que no entran en la comparación de esa variable.",
    "KS": "Estadístico de Kolmogorov–Smirnov: máxima separación entre las dos curvas acumuladas (0 = iguales, 1 = totalmente distintas).",
    "Valor p": "Probabilidad de ver una diferencia así por azar (prueba KS de dos muestras, o t del PRCC).",
    "IC inferior": "Límite inferior del intervalo de confianza al 95 % (bootstrap de 1000 remuestreos).",
    "IC superior": "Límite superior del intervalo de confianza al 95 %.",
    "Rango favorable, p05": "Percentil 5 de la variable dentro del grupo que cumple: extremo inferior de «donde viven» los KALINA.",
    "Rango favorable, p95": "Percentil 95: extremo superior del rango favorable.",
    # 3.2-4
    "Salida": "Magnitud explicada: η, o el margen de un criterio (m_O1, m_O2, m_O5).",
    "N": "Cantidad de puntos usados (n).",
    "PRCC": "Coeficiente de correlación parcial de rangos, de −1 a +1: signo y fuerza del efecto de la variable sobre la salida, sin el efecto de las demás.",
    "t de Student": "Estadístico con que se prueba que el PRCC no es cero.",
    # 3.2-5
    "Mínimo": "Menor valor del grupo.",
    "Percentil 5": "Valor bajo el cual queda el 5 % del grupo.",
    "Mediana": "Valor central del grupo.",
    "Percentil 95": "Valor bajo el cual queda el 95 % del grupo.",
    "Máximo": "Mayor valor del grupo.",
    # 3.2-6..3.3-4
    "Franja": "Intervalo de la variable en que se agrupan los puntos (número de franja o rango en kPa).",
    "Límite inferior": "Extremo inferior de la franja.",
    "Límite superior": "Extremo superior de la franja.",
    "KALINA [%]": "Porcentaje de los puntos de la franja que son KALINA.",
    "P_baja mínima [kPa]": "Menor P_baja de los KALINA de la franja.",
    "P_baja máxima [kPa]": "Mayor P_baja de los KALINA de la franja.",
    "ΔP mínima [kPa]": "Menor ΔP = P_alta − P_baja entre los KALINA de la franja.",
    "ΔP máxima [kPa]": "Mayor ΔP entre los KALINA de la franja.",
    "r_P mínima [–]": "Menor razón de presiones r_P = P_alta / P_baja entre los KALINA de la franja.",
    "r_P máxima [–]": "Mayor r_P entre los KALINA de la franja.",
    "Fuente": "Muestra de la que viene la franja (3.2 o 3.3).",
    "P_baja máxima KALINA [kPa]": "Mayor P_baja con la que hay KALINA en la franja.",
    # 4.1-3
    "Escenario": "Escenario de conservación (1 = 4.1a … 5 = 4.1e).",
    "Evaluados": "Puntos evaluados (corridas − omitidos).",
    "Factible": "True si el escenario tiene al menos un KALINA.",
    "Criterio bloqueante": "Criterio que más falla entre los genuinos no KALINA.",
    "Veces que bloquea": "Cuántas veces falla ese criterio.",
    "Identificador del D mínimo": "Punto con el menor D del escenario.",
    "Identificador del más cercano": "Punto KALINA más cercano al profesor.",
    "d del más cercano": "Distancia d de ese punto al profesor.",
    "η del más cercano [%]": "η de ese punto.",
    "Identificador de la η máxima": "Punto KALINA de mayor η.",
    "η máxima [%]": "Mayor η entre los KALINA del escenario (en 6.2-4, de los KALINA de la variable).",
    "Tiempo total [s]": "Tiempo de cómputo total del escenario.",
    # 4.2-3
    "T_fuente mínima [K]": "Menor T_fuente con KALINA en ese clima y ε_cond.",
    "T_fuente máxima [K]": "Mayor T_fuente con KALINA.",
    "η en T_fuente mínima [%]": "η en el extremo inferior de la franja.",
    "η en T_fuente máxima [%]": "η en el extremo superior de la franja.",
    # 6.1 / 4.3-4
    "Identificador": "Etiqueta del punto de control (P01…P12; A*, A4, HUS).",
    "Origen": "Estudio y punto del que viene.",
    "Clase con teqp": "Clase con el motor rápido (teqp).",
    "Clase con el motor real": "Clase con el motor de referencia (AmmoniaWaterAdapter, más lento y riguroso).",
    "Coincide la clase": "Si ambos motores dan la misma clase.",
    "η con teqp [%]": "Eficiencia con teqp.",
    "η con el motor real [%]": "Eficiencia con el motor de referencia (vacía si se cortó por tiempo).",
    "Dentro de la banda": "Si |Δη| queda dentro de la banda de incertidumbre (±0.7 pp).",
    "Margen O1 con teqp [–]": "q4 − 0.90 con teqp (positivo = cumple).",
    "Margen O1 con el motor real [–]": "Ídem con el motor de referencia.",
    "Margen O2 con teqp [K]": "T_burbuja − T de entrada a la bomba con teqp (positivo = cumple).",
    "Margen O2 con el motor real [K]": "Ídem con el motor de referencia.",
    "B verificado": "Si pasó la verificación B: el motor real recalcula la entalpía isentrópica h4s y el punto queda marcado como no verificado si |Δh4s| > 2.0 kJ/kg (se señala, no se reclasifica).",
    "Subenfriamiento, ΔT_sub [K]": "Margen de subenfriamiento exigido a la entrada de la bomba.",
    "Alcanzable": "Si existe una P_baja que lo cumple.",
    "P_baja original [kPa]": "P_baja del punto sin exigir margen.",
    "P_baja requerida [kPa]": "P_baja necesaria para tener ese margen.",
    "ΔP_baja [kPa]": "Aumento de P_baja que cuesta el margen.",
    "η original [%]": "η del punto original.",
    "η requerida [%]": "η con la P_baja requerida.",
    "Razón de presiones requerida, r_P [–]": "r_P con la P_baja requerida.",
    "Margen O1 requerido [–]": "Margen de O1 con la P_baja requerida.",
    "Iteraciones": "Pasos de bisección usados.",
    "Nota": "Advertencia sobre ese cálculo.",
    # 5.1-3
    "Arranque": "Punto de partida de la búsqueda (A1 = 5.1a … A5 = 5.1e).",
    "Clase del ciclo": "Clase del punto óptimo.",
    "Índice de cercanía D [–]": "D del punto (0 = cumple los criterios operativos).",
    "Eficiencia térmica del ciclo, η [%]": "η del ciclo.",
    "Trabajo neto, W_net [kW] (m_b = 1 kg/s)": "Trabajo neto por kg/s de mezcla base.",
    "Margen O1 [–]": "q4 − 0.90.",
    "Margen O2 [K]": "T_burbuja − T de entrada a la bomba.",
    "Margen O5 [–]": "Margen del separador (distancia de q2 al borde del intervalo bifásico).",
    "Evaluaciones": "Evaluaciones usadas por la búsqueda.",
    # 4.3-3
    "Ronda": "Ronda de la escalera (0 a 3).",
    "Datos conservados": "Grupo de datos del profesor que se conserva en ese conjunto.",
    "Datos del profesor conservados": "Cantidad de datos del profesor que se conservan exactamente en ese conjunto.",
    "Evaluadas": "Evaluaciones del conjunto.",
    "Criterio más roto": "Criterio que más falla entre los genuinos, con su conteo.",
    "Elegido": "True si fue el conjunto elegido de esa ronda.",
    "KALINA por búsqueda": "Marca interna de si los KALINA del conjunto salen de la búsqueda dirigida; en esta campaña vale «no» en las 18 filas.",
    "Piso de O2 (T de diseño de la bomba) [K]": "Temperatura con que se evalúa la cavitación.",
    "Tiempo [s]": "Tiempo del motor de referencia.",
    "Mensaje": "Mensaje del clasificador.",
    # 6.2
    "Familia": "Familia de puntos de Húsavík (F3-A…F3-E de septiembre; B1, B2 nuevas) o, en S-1, familia de método del estudio.",
    "Variables movidas": "Variables que se mueven en la familia.",
    "Piso de O2 [K]": "Temperatura con que se evalúa la cavitación.",
    "KALINA (septiembre)": "KALINA que tenía la familia en septiembre.",
    "KALINA (ahora)": "KALINA con el código actual.",
    "Confiables ahora": "Cuántos de ellos son confiables.",
    "Coincide": "Cuántos mantienen la clase de septiembre.",
    "|Δη| medio [pp]": "Diferencia media de η con septiembre.",
    "|Δη| máximo [pp]": "Diferencia máxima de η con septiembre.",
    "Mínimo KALINA": "Menor valor de la variable con KALINA.",
    "Máximo KALINA": "Mayor valor de la variable con KALINA.",
    "Criterio que rompe al bajar la variable": "Primer criterio que falla más allá del mínimo KALINA.",
    "Criterio que rompe al subir la variable": "Primer criterio que falla más allá del máximo KALINA.",
    "η mínima [%]": "Menor η entre los KALINA de la variable.",
    "Rango explorado": "Intervalo barrido de la variable.",
    "Valor en A*": "Valor de la variable en el ancla.",
    "Elasticidad E en A*": "E = (Δη/η) / (Δv/v) por diferencia central con ±1 % del rango.",
    "Δη (genuinos) [pp]": "η máxima − η mínima en la malla, solo con genuinos de W_net > 0.",
    "η máx. KALINA en la malla [%]": "Mayor η entre los KALINA de la malla de esa variable.",
    "Valor de esa η máx.": "Valor de la variable en ese punto.",
    "Zona KALINA": "Tramo contiguo a A* donde la clase sigue siendo KALINA.",
    "Criterios que bloquean": "Criterios que fallan fuera de la zona KALINA, con su conteo.",
    "Distancia a la frontera de O2 [–]": "|valor de cruce − valor en A*| / rango. Cuanto menor, más cerca de la frontera.",
    "x_b": "Fracción de amoníaco de la mezcla base.",
    "η en ese punto [%]": "η en la P_baja mínima KALINA.",
    "Criterio que más bloquea": "Criterio que más falla en esa fila.",
    # S
    "Estudio": "Estudio (N.M).",
    "Nombre": "Nombre del estudio.",
    "Partes": "Partes del estudio.",
    "Punto base": "Punto del que parte el barrido (PP, A*, caja…).",
    "Diseño": "Diseño del barrido (malla, OAT, LHS, búsqueda…).",
    "Genuinas": "Corridas genuinas.",
    "Flujo másico base, m_b [kg/s]": "Flujo másico de mezcla base (siempre 1 kg/s).",
    "motor": "Motor de propiedades (teqp).",
    "Datos fijos": "Datos que quedan fijos en la parte.",
    "n KALINA confiables": "Cantidad de KALINA confiables en la parte.",
    "p5": "Percentil 5 de la variable entre los KALINA confiables.",
    "p95": "Percentil 95 de la variable entre los KALINA confiables.",
    "Dato": "Dato del enunciado del profesor.",
    "Valor del profesor": "Valor que fija el enunciado.",
    "¿Se conserva con KALINA?": "Si el dato se conserva en los KALINA confiables.",
    "Respaldo": "Estudio que lo sostiene.",
    "Genuinos no KALINA": "Genuinos que no son KALINA (los que tienen algún criterio fallado).",
}
for _v in ("O1", "O2", "O3", "O5"):
    D[f"{_v} (n)"] = f"Genuinos no KALINA en que falla {_v}."
    D[f"{_v} [%]"] = f"Porcentaje de los genuinos no KALINA en que falla {_v}."
# Definiciones que dependen de la tabla (la misma columna significa cosas distintas)
DT = {("1.1-3", "Δη [pp]"): "η de ahora − η de la Fase 1, en puntos porcentuales (vacía si alguna no existe).",
      ("6.1-3", "Δη [pp]"): "η con el motor real − η con teqp, en puntos porcentuales.",
      ("4.3-4", "Δη [pp]"): "η con el motor real − η con teqp, en puntos porcentuales.",
      ("6.1-4", "Δη [pp]"): "η con la P_baja requerida − η original, en puntos porcentuales (negativo = cuesta eficiencia).",
      ("3.2-4", "Valor p"): "Probabilidad de ver un PRCC así si no hubiera relación (prueba t); «< 0.001» = prácticamente cero.",
      ("3.2-3", "Valor p"): "Probabilidad de ver esa separación por azar (prueba KS de dos muestras).",
      ("3.2-5", "Valor p"): "Probabilidad de ver esa separación por azar (prueba KS de dos muestras).",
      ("4.1-4", "Valor p"): "Probabilidad de ver esa separación por azar (prueba KS de dos muestras).",
      ("6.2-3", "Origen"): "«sept» = punto de la Fase 3 de septiembre; «nuevo» = barrido nuevo de esta campaña.",
      ("4.3-4", "Origen"): "Punto del que viene (ancla A*, óptimo 5.1d o base de Húsavík).",
      ("6.1-3", "Origen"): "Estudio y punto del que viene (P01…P12).",
      ("6.1-4", "Origen"): "Estudio y punto del que viene.",
      ("3.2-5", "Grupo"): "Grupo comparado: KALINA o no KALINA genuino.",
      ("4.1-4", "Grupo"): "Grupo comparado: KALINA o no KALINA.",
      ("6.1-4", "Identificador"): "Etiqueta del punto de control (P01…P12).",
      ("4.3-4", "Identificador"): "Etiqueta del punto (A* = ancla, A4 = óptimo 5.1d, HUS = Húsavík).",
      ("S-1", "Familia"): "Familia de método del estudio: R reproducibilidad, F factibilidad, L local, G global, C conservación, V verificación, O optimización, H caso real."}
VAR_COLS = ("Temperatura de la fuente", "Temperatura del sumidero", "Presión de alta", "Presión de baja",
            "Fracción de amoníaco", "Eficiencia isentrópica", "Efectividad del")
FALTAN = set()


def definicion(col, tid=None):
    if (tid, col) in DT:
        return DT[(tid, col)]
    if col in D:
        return D[col]
    if col.startswith(VAR_COLS):
        return "Valor de la variable en ese punto (la columna lleva su unidad)."
    FALTAN.add(col)
    return "(sin definición)"


# ---------------------------------------------------------------- lectura
def registro():
    return pd.read_csv(os.path.join(TAB, "registro_tablas.csv"))


def leer_tabla(tid):
    r = registro()
    m = r[r["id_nuevo"] == tid].iloc[0]
    return m, pd.read_csv(os.path.join(TAB, str(m["carpeta_nueva"]), m["archivo"] + ".csv"))


def fmt(v):
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return "—"
    if isinstance(v, (bool, np.bool_)):
        return "sí" if v else "no"
    if isinstance(v, float):
        if v.is_integer() and abs(v) < 1e7:
            return str(int(v))
        return f"{v:.3f}".rstrip("0").rstrip(".")
    return str(v).replace("|", "\\|").replace("\n", " ")


def md_tabla(df):
    cab = "| " + " | ".join(str(c).replace("|", "\\|") for c in df.columns) + " |"
    sep = "|" + "|".join(["---"] * len(df.columns)) + "|"
    def celda(c, v):
        if str(c) == "Valor p" and isinstance(v, (int, float)) and not pd.isna(v) and v < 0.0005:
            return "< 0.001"
        return fmt(v)

    filas = ["| " + " | ".join(celda(c, v) for c, v in zip(df.columns, fila)) + " |"
             for fila in df.itertuples(index=False)]
    return "\n".join([cab, sep] + filas)


def marcador_T(arg):
    partes = [p.strip() for p in arg.split("|")]
    tid, opts = partes[0], dict(p.split("=", 1) for p in partes[1:])
    m, df = leer_tabla(tid)
    if "filtro" in opts:
        col, val = opts["filtro"].split("=", 1) if "=" in opts["filtro"] else (None, None)
        if col:
            df = df[df[col].astype(str) == val]
    if "cols" in opts:
        df = df[[c for c in opts["cols"].split(";") if c in df.columns]]
    total = len(df)
    if "n" in opts:
        df = df.head(int(opts["n"]))
    out = md_tabla(df)
    nota = f"\n\n*Tabla {tid} · hoja «{m['hoja']}» · {total} filas" + (
        f" (se muestran {len(df)})" if len(df) < total else "") + "."
    return out + nota + "*"


def marcador_COLS(tid):
    m, df = leer_tabla(tid)
    return "\n".join(f"- **{c}.** {definicion(c, tid)}" for c in df.columns)


def pie_fig(fid):
    f = glob.glob(os.path.join(FIG, "*", f"Figura_{fid}_*.pie.md"))
    if not f:
        return None, []
    return f[0], open(f[0], encoding="utf-8").read().splitlines()


def marcador_P(fid):
    ruta, lineas = pie_fig(fid)
    if not ruta:
        return f"*(figura {fid}: sin pie)*"
    out = [f"> **{lineas[0]}**"]
    for l in lineas[1:]:
        out.append(f"> {l}")
    return "\n".join(out)


def hojas_excel():
    import openpyxl
    wb = openpyxl.load_workbook(EXCEL)
    filas = []
    for ws in wb.worksheets:
        filas.append((ws.title, ws.max_row, ws.max_column, len(ws._charts), ws.sheet_state))
    return filas


def describir_hoja(nombre, reg):
    t = nombre
    if t in ("LEEME", "Nomenclatura", "Indice", "Punto_profesor", "Tablero"):
        return {"LEEME": "Cómo leer el libro, glosario y colores de clase.",
                "Nomenclatura": "Equivalencias entre la numeración nueva (pregunta → estudio → parte) y los códigos antiguos (M05, KF-M05-3, B09-3).",
                "Indice": "Lista de todas las hojas con vínculo, estudio, partes y qué muestra.",
                "Punto_profesor": "Los 11 datos del enunciado y el ancla A* lado a lado.",
                "Tablero": "Selector (celda B2) que cambia cifras, condiciones y tres gráficos según el estudio o la parte elegidos."}[t]
    m = re.match(r"^(S-\d|\d\.\d-\d)\b", t)
    if m:
        r = reg[reg["id_nuevo"] == m.group(1)]
        if len(r):
            return f"Tabla {m.group(1)}: {r.iloc[0]['titulo_nuevo']}."
    if re.match(r"^\d\.\d-R\b", t):
        return "Todos los puntos del estudio, un bloque por barrido (y por variable); con sus gráficos a la derecha."
    if t.endswith(" Datos de gráficos"):
        return "Los números con los que Excel dibuja cada gráfico de la pregunta, un bloque por figura."
    if t.endswith(" Gráficos"):
        return "Gráficos nativos de Excel de la pregunta, con su título arriba y su pie abajo."
    return {"A.1 Anexo fuera de rango": "Conteo por clase del anexo (η y ε de 0.20 a 0.99); no entra en las conclusiones.",
            "Datos": "Tabla canónica: un punto por fila (4 560), con sus 10 variables, clase, η, W_net, ΔP, r_P y códigos de falla; con filtros.",
            "Condiciones": "Qué quedó fijo y qué se barrió, por parte y variable (insumo del Tablero).",
            "Partes": "Lista de partes con sus datos fijos y diferencias internas.",
            "Aux_tablero": "Hoja oculta: listas y fórmulas que alimentan el Tablero.",
            "Aux_puntos": "Hoja oculta: puntos ordenados por parte para los gráficos del Tablero."}.get(t, "")


def marcador_HOJAS():
    reg = registro()
    filas = hojas_excel()
    out = ["| # | Hoja | Filas × columnas | Gráficos | Qué contiene |", "|---|---|---|---|---|"]
    for i, (n, r, c, g, est) in enumerate(filas, 1):
        d = describir_hoja(n, reg)
        oculta = " *(oculta)*" if est != "visible" else ""
        out.append(f"| {i} | `{n}`{oculta} | {r} × {c} | {g or ''} | {d} |")
    return "\n".join(out)


def marcador_CAT_TABLAS():
    reg = registro()
    out = []
    for _, m in reg.iterrows():
        df = pd.read_csv(os.path.join(TAB, str(m["carpeta_nueva"]), m["archivo"] + ".csv"))
        out.append(f"#### Tabla {m['id_nuevo']} · {m['titulo_nuevo']}")
        out.append(f"- **Hoja:** `{m['hoja']}` · **Filas:** {len(df)} · **Columnas:** {len(df.columns)}")
        out.append(f"- **Fuente:** `{m['fuente']}` · **Filtro:** {m['filtro']}")
        out.append("- **Columnas:** " + "; ".join(f"*{c}*" for c in df.columns))
        out.append("")
    return "\n".join(out)


def marcador_CAT_FIGURAS():
    out = []
    def clave(p):
        fid = os.path.basename(p).split("_")[1]
        return (fid.startswith("S"), [int(x) for x in re.split(r"[.-]", fid) if x.isdigit()])

    for ruta in sorted(glob.glob(os.path.join(FIG, "*", "*.pie.md")), key=clave):
        lineas = open(ruta, encoding="utf-8").read().splitlines()
        fid = os.path.basename(ruta).split("_")[1]
        out.append(f"#### {lineas[0]}")
        for l in lineas[1:]:
            out.append(f"- {l}")
        out.append("")
    return "\n".join(out)


MARC = re.compile(r"\{\{\s*(T|COLS|P|HOJAS|CAT_TABLAS|CAT_FIGURAS)\s*([^}]*)\}\}")


def sustituir(m):
    tipo, arg = m.group(1), m.group(2).strip()
    return {"T": marcador_T, "COLS": marcador_COLS, "P": marcador_P,
            "HOJAS": lambda a: marcador_HOJAS(), "CAT_TABLAS": lambda a: marcador_CAT_TABLAS(),
            "CAT_FIGURAS": lambda a: marcador_CAT_FIGURAS()}[tipo](arg)


def main():
    texto = open(PLANTILLA, encoding="utf-8").read()
    texto = MARC.sub(sustituir, texto)
    open(SALIDA, "w", encoding="utf-8").write(texto)
    print(f"Guía: {SALIDA} · {len(texto.splitlines())} líneas")
    if FALTAN:
        print("Columnas sin definición:", sorted(FALTAN))
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
