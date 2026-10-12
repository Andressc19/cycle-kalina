# -*- coding: utf-8 -*-
"""A2 · Verificación del libro Excel_kalina_final_final_final.xlsx (TASK_CONTEXT A2)."""
import random

import pandas as pd
from openpyxl import load_workbook

from a2_comun import EXCEL_VIEJO, INV, LIBRO, registro
from a2_excel_libro import sha

H_REF = "70cd44edfa9b9d6946b7d74f2833700c60be953cc74c12c84f39234bfe25d2ea"


def main():
    fallas = 0
    wb = load_workbook(LIBRO)
    print(f"[1] {len(wb.sheetnames)} hojas: {', '.join(wb.sheetnames[:8])} …")
    reg = registro()
    conteo = pd.read_csv(f"{INV}/conteo_por_subcaso.csv").set_index("id_subcaso")
    hojas01 = [m for _, m in reg.iterrows() if m["tabla"].endswith("-01") and m["carpeta"] != "G"]
    random.seed(1)
    for m in random.sample(hojas01, 5):
        ws = wb[m["hoja"][:31]]
        # nomenclatura nueva: columna "Parte" con etiquetas; se compara la suma del estudio
        h = next(r for r in range(1, 80) if ws.cell(r, 1).value == "Parte")
        cab = [c.value for c in ws[h]]
        filas = [f for f in ws.iter_rows(min_row=h + 1, values_only=True)
                 if f[0] and not str(f[0]).startswith("Notas")]
        ref = conteo.loc[str(m["subcasos"]).split("; ")]
        ok = sum(int(f[cab.index("Corridas")]) for f in filas) == int(ref["n_puntos"].sum()) and \
            sum(int(f[cab.index("KALINA")]) for f in filas) == int(ref["kalina"].sum())
        print(f"[2] {m['hoja']}: conteos = conteo_por_subcaso → {'OK' if ok else 'FALLA'}")
        fallas += not ok
    for _, m in reg.iterrows():
        if m["carpeta"] == "G" or m["tabla"].endswith("-00") or m["tabla"].endswith("-01"):
            continue
        b = str(m["bloques"]).split(", ")
        if len(b) > 1 and not str(m["bloques"]).startswith("B13"):
            print(f"[3] INFO {m['hoja']}: tabla de resultado con varios bloques ({m['bloques']}); "
                  "cada fila lleva su subcaso")
    h = sha(EXCEL_VIEJO)
    print(f"[4] Excel_final_kalina.xlsx sin cambios: {'OK' if h == H_REF else 'FALLA'}")
    fallas += h != H_REF
    t = wb["Tablero"]
    dv = t.data_validations.dataValidation
    graf = len(t._charts)
    ok = bool(dv) and graf == 3
    print(f"[5] Tablero: validación en B2 = {bool(dv)} ({dv[0].formula1 if dv else '—'}); "
          f"gráficos = {graf} → {'OK' if ok else 'FALLA'}")
    fallas += not ok
    graf = {h: len(wb[h]._charts) for h in wb.sheetnames if h.endswith(" Gráficos")}
    viejas = [h for h in wb.sheetnames if h.endswith(("_figuras", "_graficos")) and h != "Aux_graficos"]
    ok = len(graf) == 6 and all(graf.values()) and not viejas and "Figuras_matplotlib_ref" in wb.sheetnames
    print(f"[6] Gráficos nativos: {sum(graf.values())} en {len(graf)} hojas; hojas de imágenes por caso: "
          f"{len(viejas)}; hoja de referencia: {'Figuras_matplotlib_ref' in wb.sheetnames} → "
          f"{'OK' if ok else 'FALLA'}")
    fallas += not ok
    print(f"\nResultado: {'TODO OK' if not fallas else str(fallas) + ' FALLA(S)'}")
    return 1 if fallas else 0


if __name__ == "__main__":
    raise SystemExit(main())
