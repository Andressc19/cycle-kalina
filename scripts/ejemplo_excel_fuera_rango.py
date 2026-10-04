"""Excel de ejemplo de la marca "fuera de rango confiable" (modo A).

Filas sintéticas (no corre el ciclo): mezcla de puntos dentro y fuera de rango,
con al menos un KALINA y un NO_CONVERGIO marcados. Usa el mismo formateo de
celdas que `formatear_excel_barridos.py`.
"""
from __future__ import annotations

import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill

REPO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO))
from src._rango_barrido import anotar_rango  # noqa: E402
from src.excel_rango import GRIS_RANGO, LEYENDA_RANGO, marcar_tabla  # noqa: E402
from src.state import EstadoTermo  # noqa: E402

SALIDA = REPO / "resultados" / "2026-10-03_ejemplo_fuera_rango" / "ejemplo_fuera_rango.xlsx"
COLOR_CLASE = {"KALINA": "C6EFCE", "CORREGIBLE": "FFD9A0",
               "NO_CONVERGIO": "F8B4B4", "INVIABLE": "FFF2A8"}


def _res(**cambios):
    est = {}
    for i in range(1, 11):
        d = dict(T=350.0, P=3000.0, h=0.0, s=0.0, x=0.5, fase="liquido")
        d.update(cambios.get(f"e{i}", {}))
        est[f"e{i}"] = EstadoTermo(etiqueta=str(i), **d)
    return {"estados": est}


CASOS = [  # (P_alta, P_baja, x_b, clase, eta, Wnet, resultado | None, nota)
    (3000.0, 400.0, 0.50, "KALINA", 0.112, 112.0, _res(), "dentro de rango"),
    (5000.0, 700.0, 0.65, "KALINA", 0.131, 140.5,
     _res(e2=dict(T=431.0, fase="bifasico")), "T2 líquido > 420 K"),
    (12000.0, 900.0, 0.70, "CORREGIBLE", 0.140, 150.2,
     _res(e2=dict(P=12000.0, T=470.0, fase="vapor"),
          e3=dict(P=12000.0, T=470.0, fase="vapor")), "vapor > 10 MPa"),
    (4000.0, 420.0, 0.60, "INVIABLE", 0.080, 70.1, _res(), "dentro de rango"),
    (45000.0, 400.0, 0.50, "NO_CONVERGIO", None, None, None, "P_alta > 40 MPa"),
    (3000.0, 400.0, 0.50, "NO_CONVERGIO", None, None, None, "dentro de rango"),
]


def main() -> None:
    cab = ["P_alta", "P_baja", "x_b", "T_sumidero", "clasificacion", "eta", "Wnet",
           "nota", "fuera_rango", "detalle_rango", "columnas_rango"]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Ejemplo_barrido"
    ws.append(cab)
    for P_a, P_b, x, clase, eta, W, res, nota in CASOS:
        combo = dict(P_alta=P_a, P_baja=P_b, x_b=x, T_sumidero=300.032917)
        fila = dict(combo, clasificacion=clase, eta=eta, Wnet=W, nota=nota)
        fila = anotar_rango(fila, combo, res, "teqp", None, "extrapolar_marcar")
        ws.append([fila.get(c) for c in cab])
    for r in range(2, ws.max_row + 1):
        c = ws.cell(r, cab.index("clasificacion") + 1)
        c.fill = PatternFill("solid", fgColor=COLOR_CLASE[c.value])
    n = marcar_tabla(ws, cab)
    for j in range(1, len(cab) + 1):
        ws.cell(1, j).font = Font(bold=True)
    ws.column_dimensions["J"].width = 90
    r = ws.max_row + 2
    ws.cell(r, 1, "FUERA_RANGO").fill = PatternFill("solid", fgColor=GRIS_RANGO)
    ws.cell(r, 2, LEYENDA_RANGO)
    SALIDA.parent.mkdir(parents=True, exist_ok=True)
    wb.save(SALIDA)
    print(f"{SALIDA.name}: {ws.max_row - 3} filas, {n} celdas grises")


if __name__ == "__main__":
    main()
