# -*- coding: utf-8 -*-
"""Sensibilidad 1 - x7: ensambla el libro en el orden indicado y lo guarda."""

import os

from openpyxl import Workbook

from x1_cargar import cargar_todo, ESTUDIOS, SALIDA
from x2_formato import FUENTE_SUB, tabla
from x3_hojas_datos import hoja_datos
from x4_resumenes import hoja_resumen
from x5_graficos import dibujar_todos
from x6_sintesis import hoja_sintesis

# Evitamos escribir hojas por defecto
wb = Workbook()
wb.remove(wb.active)

# Bandas de T_fuente (grados C) de la tabla de la seccion "Que es el rango experimental".
BANDAS = [(60, 120), (120, 130), (130, 140), (140, 150), (150, 160), (160, 170),
          (170, 180), (180, 200), (200, 250), (250, 300)]

TEXTO_RANGO = (
    "La ecuación del motor (IAPWS G4-01, §6) fue contrastada con datos experimentales de líquido "
    "hasta 420 K y de vapor hasta 10 MPa. Un punto queda «fuera del rango experimental» cuando "
    "alguno de sus estados de líquido o bifásico supera 420 K (en esta campaña, T2 o T5; ocurre "
    "desde unos 150 °C de fuente). El motor lo calcula y lo clasifica igual, y la guía declara la "
    "ecuación válida hasta el lugar crítico; la marca solo informa que en esa zona no hay datos "
    "experimentales con que contrastar el cálculo. Estos puntos forman parte de los resultados y "
    "de las conclusiones. El lugar crítico no se vigila en el motor (la guía advierte que su "
    "ubicación es incierta); con P_alta <= 5 MPa en esta campaña no se espera cruzarlo, pero no se "
    "comprobó punto por punto.")


def _tabla_bandas(df):
    """Filas [banda, KALINA, fuera del rango experimental, % fuera] y los totales globales."""
    k = df[(df["estudio"] >= 1) & df["kalina"]]
    tfc = k["T_fuente"] - 273.15
    filas, n_tot, f_tot = [], 0, 0
    for a, b in BANDAS:
        s = k[(tfc > a) & (tfc <= b)]
        n = int(s.shape[0])
        f = int(s["fuera_rango"].sum())
        n_tot += n
        f_tot += f
        filas.append([f"{a}-{b}", n, f, round(100.0 * f / n, 1) if n else 0.0])
    return filas, n_tot, f_tot


def _seccion_rango(ws, fila, df):
    """Seccion 'Que es el rango experimental' con la tabla por banda de T_fuente."""
    ws.cell(fila, 1, "¿Qué es el rango experimental?").font = FUENTE_SUB
    fila += 1
    ws.cell(fila, 1, TEXTO_RANGO)
    fila += 2
    filas, n_tot, f_tot = _tabla_bandas(df)
    _, ultima, _ = tabla(ws, fila,
                         ["Banda de T_fuente (°C)", "KALINA",
                          "fuera del rango experimental", "% fuera"],
                         filas, anchos=[22, 10, 30, 10])
    ultima += 1
    ws.cell(ultima, 1,
            f"Global: {f_tot} de {n_tot} KALINA (estudios 1 a 11, sin el estudio 0) "
            "fuera del rango experimental.")


def construir():
    df = cargar_todo()
    # 1. LEEME
    ws = wb.create_sheet("LEEME")
    ws.cell(1, 1, "Sensibilidad 1: Kalina_anal_sensible_1.xlsx")
    ws.cell(2, 1, "Un solo libro con los 12 estudios. Todos los datos salen de los CSV de resultados/.")
    # Guía de hojas
    ws.cell(4, 1, "Guia de hojas")
    fila = 5
    ws.cell(fila, 1, "Conclusion consolidada"); fila += 1
    for e in ESTUDIOS:
        ws.cell(fila, 1, e["hoja"]); fila += 1
        ws.cell(fila, 1, e["datos"]); fila += 1
    # Seccion de terminologia y tabla por banda de T_fuente
    _seccion_rango(ws, fila + 1, df)
    # 2. Conclusion consolidada (con sintesis)
    hoja_sintesis(wb, df)
    # 3. Por cada estudio: resumen y datos
    for e in ESTUDIOS:
        hoja_resumen(wb, e, df)
        hoja_datos(wb, e, df)
    # Dibujar graficos registrados
    dibujar_todos(wb)
    # Asegurar que LEEME sea la primera
    sheets = wb.sheetnames
    if sheets[0] != "LEEME":
        idx = sheets.index("LEEME")
        sheets.insert(0, sheets.pop(idx))
        wb._sheets = [wb[s] for s in sheets]
    # Guardar
    os.makedirs(os.path.dirname(SALIDA), exist_ok=True)
    wb.save(SALIDA)
    print("OK: guardado", SALIDA)


if __name__ == "__main__":
    construir()
