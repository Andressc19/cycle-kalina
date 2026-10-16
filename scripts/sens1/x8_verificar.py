# -*- coding: utf-8 -*-
"""Sensibilidad 1 - x8: verificaciones obligatorias. Imprime OK/FALLA por control."""

import os
import sys

from openpyxl import load_workbook

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from x1_cargar import cargar_todo, ESTUDIOS, SALIDA
from x8b_graficos import control4, control5, x8b_textos

ERRORES = []

df = cargar_todo()


def ok(msg):
    print("OK:", msg)


def fail(msg):
    print("FALLA:", msg)
    ERRORES.append(msg)


# 1
try:
    wb = load_workbook(SALIDA)
except Exception as e:
    fail(f"1.1 no abre: {e}")
    sys.exit(1)
for s in wb.sheetnames:
    if len(s) > 31:
        fail(f"1.2 hoja con >31 chars: {s}")
    ws = wb[s]
    if ws.sheet_state != "visible":
        fail(f"1.3 hoja no visible: {s}")
    if ws.freeze_panes:
        fail(f"1.4 freeze panes en {s}: {ws.freeze_panes}")
ok("1. libro abierto, sin hojas ocultas, sin freeze panes")


# 2
for e in ESTUDIOS:
    ws = wb[e["datos"]] if e["datos"] in wb else None
    if ws is None:
        fail(f"2.0 falta hoja de datos {e['datos']}")
        continue
    marker = None
    for i in range(4, ws.max_row + 1):
        v0 = ws[i][0].value if ws[i][0].value is not None else None
        if isinstance(v0, str) and "Descripci" in v0:
            marker = i
            break
    stop = marker if marker else ws.max_row + 1
    count = 0
    for i in range(4, stop):
        vals = [c.value for c in ws[i]]
        if all(v is None or v == "" for v in vals):
            continue
        count += 1
    real = len(df[df["estudio"] == e["n"]])
    if count != real:
        fail(f"2.{e['n']} {e['datos']}: filas {count} vs {real}")
    else:
        ok(f"2.{e['n']} {e['datos']}: {real} filas correctas")
ok("2. conteo de filas verificado")


# 3
esperados = {1:186,2:398,3:249,4:159,5:17,6:47,7:45,8:188,9:43,10:218,11:203}
for n, exp in esperados.items():
    sub = df[df["estudio"] == n]
    real_k = int(sub[sub["clasificacion"] == "KALINA"].shape[0])
    if exp != real_k:
        fail(f"3.{n} KALINA esperado {exp} vs real {real_k}")
ok("3. chequeo KALINA principales hecho")


# 4 y 5 - graficos nativos: series y valores visibles, y sin solapes con tablas ni otros graficos
print("Graficos del libro:")
control4(wb, ok, fail, [e["hoja"] for e in ESTUDIOS])
control5(wb, ok, fail)


# 6 - tabla por banda de T_fuente (estudios 1 a 11, sin el estudio 0)
BANDAS = [(60, 120), (120, 130), (130, 140), (140, 150), (150, 160), (160, 170),
          (170, 180), (180, 200), (200, 250), (250, 300)]
k = df[(df["estudio"] >= 1) & df["kalina"]]
tfc = k["T_fuente"] - 273.15
n_tot = f_tot = 0
print("Tabla por banda de T_fuente (estudios 1-11):")
for a, b in BANDAS:
    s = k[(tfc > a) & (tfc <= b)]
    n = int(s.shape[0])
    f = int(s["fuera_rango"].sum())
    n_tot += n
    f_tot += f
    pct = round(100.0 * f / n, 1) if n else 0.0
    print("   %3d-%3d  KALINA %4d  fuera %4d  %5.1f %%" % (a, b, n, f, pct))
if n_tot == 1753 and f_tot == 275:
    ok(f"6. bandas T_fuente: {f_tot} de {n_tot} KALINA fuera del rango experimental")
else:
    fail(f"6. bandas T_fuente: total {n_tot} (esperado 1753), fuera {f_tot} (esperado 275)")


# 7 - ninguna celda del libro (ni titulo/etiqueta de grafico) con terminos prohibidos
PROHIBIDOS = ["confiable", "fuera de rango", "sin respaldo", "no citable", "indicio", "extrapol"]


def _prohibido(texto):
    low = str(texto).lower()
    return any(p in low for p in PROHIBIDOS)


malas = []
for s in wb.sheetnames:
    for row in wb[s].iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.strip().lower() != "fuera_rango" and _prohibido(c.value):
                malas.append((s, c.coordinate, c.value[:40]))
    for ch in wb[s]._charts:
        for texto in x8b_textos(ch):
            if _prohibido(texto):
                malas.append((s, "titulo/etiqueta de grafico", texto[:40]))
if not malas:
    ok("7. ninguna celda ni titulo de grafico con terminos prohibidos (confiable/fuera de rango/sin respaldo/no citable/indicio/extrapol)")
else:
    fail(f"7. {len(malas)} celdas/titulos con terminos prohibidos: {malas[:10]}")


if ERRORES:
    print("== FIN FALLA ==")
    for e in ERRORES: print(e)
    sys.exit(1)
else:
    print("== FIN OK ==")
    sys.exit(0)
