# -*- coding: utf-8 -*-
"""Sensibilidad 1 - x3: hojas `E{n} Datos ...` con todos los puntos de cada estudio.

Una fila por punto y una columna por dato (con su unidad en el encabezado); filtro en la fila
de encabezados, cuadricula completa y, al final, dos filas de separacion y la descripcion corta
(que se itero, que quedo fijo y el conteo por clase).
"""
import math

from openpyxl.styles import Alignment

from x1_cargar import ESTADOS
from x2_formato import BORDE, FILL_ENC, FUENTE_ENC, anchos_col, descripcion


def _fmt(v, tipo):
    if v is None:
        return None
    if isinstance(v, float) and math.isnan(v):
        return None
    if tipo == "si":
        return "sí" if bool(v) else "no"
    if tipo == "txt":
        return str(v)
    if tipo == "n2":
        return round(float(v), 2)
    if tipo == "n3":
        return round(float(v), 3)
    if tipo == "n4":
        return round(float(v), 4)
    if tipo == "n5":
        return round(float(v), 5)
    return v


# (encabezado, clave del DataFrame, tipo de formato)
CAMPOS = [
    ("Punto", "id_punto", "txt"),
    ("Centro", "centro", "txt"),
    ("Variable o par barrido", "varpar", "txt"),
    ("T_fuente (K)", "T_fuente", "n2"),
    ("T_sumidero (K)", "T_sumidero", "n2"),
    ("P_alta (kPa)", "P_alta", "n2"),
    ("P_baja (kPa)", "P_baja", "n2"),
    ("x_b (-)", "x_b", "n4"),
    ("Eficiencia de turbina, eta_t (-)", "eta_t", "n3"),
    ("Eficiencia de bomba, eta_p (-)", "eta_p", "n3"),
    ("Efectividad del HRVG, eps_HRVG (-)", "eps_hrvg", "n3"),
    ("Efectividad del regenerador, eps_reg (-)", "eps_reg", "n3"),
    ("Efectividad del condensador, eps_cond (-)", "eps_cond", "n3"),
    ("m_b (kg/s)", "m_b", "n2"),
    ("T_amb_diseno (K)", "T_amb_diseno", "n2"),
    ("Clasificación", "clasificacion", "txt"),
    ("Convergió (sí/no)", "convergio", "si"),
    ("Fuera del rango experimental (sí/no)", "fuera_rango", "si"),
    ("KALINA dentro del rango experimental (sí/no)", "confiable", "si"),
    ("η (%)", "eta_pct", "n4"),
    ("Wnet (kW)", "Wnet", "n4"),
    ("Margen O1, m_O1 (-)", "m_O1", "n4"),
    ("Margen O2 (cavitación), m_O2 (K)", "m_O2", "n4"),
    ("Margen O3, m_O3 (kW)", "m_O3", "n4"),
    ("Margen O5, m_O5 (-)", "m_O5", "n4"),
    ("Calor de la fuente, Qi (kW)", "Qi", "n3"),
    ("Calor rechazado, Qout (kW)", "Qout", "n3"),
    ("Trabajo de turbina, Wt (kW)", "Wt", "n3"),
    ("Trabajo de bomba, Wp (kW)", "Wp", "n3"),
    ("Calor del regenerador, Qreg (kW)", "Qreg", "n3"),
    ("Salto entálpico de turbina, dh_turb (kJ/kg)", "dh_turb", "n3"),
    ("Qreg/Qi (-)", "qreg_qi", "n4"),
    ("P_alta/P_baja (-)", "rp", "n3"),
    ("Tiempo de cómputo, t_s (s)", "t_s", "n2"),
    ("Mensaje del motor", "mensaje", "txt"),
]


def _campos_estados():
    return ([(f"T_e{i} (K)", f"T_e{i}", "n3") for i in range(1, 11)] +
            [(f"h_e{i} (kJ/kg)", f"h_e{i}", "n3") for i in range(1, 11)])


ANCHOS = [16, 18, 20, 10, 11, 10, 10, 8, 12, 12, 13, 13, 13, 9, 12, 14, 9, 12, 12, 9, 10,
          10, 12, 10, 10, 11, 11, 11, 10, 11, 12, 10, 10, 11, 60] + [9, 9, 9, 9, 9, 9, 9, 9, 9, 9,
                                                                    10, 10, 10, 10, 10, 10, 10, 10, 10, 10]


def hoja_datos(wb, est, df):
    """Escribe la hoja de datos completa de un estudio y devuelve su nombre."""
    ws = wb.create_sheet(est["datos"])
    campos = CAMPOS + _campos_estados()
    ws.cell(1, 1, f"{est['hoja']} · datos de todos los puntos").font = FUENTE_ENC
    r = 3
    for j, (t, _, _) in enumerate(campos, 1):
        c = ws.cell(r, j, t)
        c.font = FUENTE_ENC
        c.fill = FILL_ENC
        c.border = BORDE
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    anchos_col(ws, ANCHOS[:len(campos)], 1)
    sub = df[df["estudio"] == est["n"]].reset_index(drop=True)
    enc = r
    for _, fila in sub.iterrows():
        r += 1
        for j, (_, clave, tipo) in enumerate(campos, 1):
            v = fila.get(clave)
            if clave == "varpar":
                v = fila.get("variable") or fila.get("par") or ""
            c = ws.cell(r, j, _fmt(v, tipo))
            c.border = BORDE
    ultima = r
    ws.auto_filter.ref = f"A{enc}:{_letra(len(campos))}{ultima}"
    clases = sub["clasificacion"].value_counts().to_dict()
    fuera = int(sub["fuera_motor"].sum())
    conf = int(sub["confiable"].sum())
    lineas = [
        f"Qué se iteró: {est['nombre']} ({est['grupo'] or 'sin subdivisiones'}).",
        "Qué quedó fijo: las variables de cada barrido están en la columna «Variable o par barrido»; "
        "m_b = 1.0 kg/s salvo en el estudio 0; el piso de O2 es el T_sumidero de cada punto.",
        "Conteo por clase: " + ", ".join(f"{k} {v}" for k, v in sorted(clases.items())) +
        f"; KALINA dentro del rango experimental {conf}; KALINA fuera del rango experimental "
        f"{fuera}; total {len(sub)} puntos.",
    ]
    descripcion(ws, ultima, lineas)
    return ws.title


def _letra(n):
    s = ""
    while n:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


__all__ = ["CAMPOS", "ESTADOS", "hoja_datos"]
