# -*- coding: utf-8 -*-
"""Sensibilidad 1 - x2: estilos y escritura de tablas (cuadricula, colores de clase, anchos,
encabezados con unidad). Sin paneles inmovilizados (freeze panes) en ninguna hoja.

El color de cada fila sale de su clase: KALINA verde, CORREGIBLE naranja, NO_CONVERGIO rojo e
INVIABLE amarillo, todos en tonos claros. La hoja de datos completa no se colorea por fila.
"""
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L

CM_FILA, CM_COL = 0.53, 1.7            # cm por fila / columna (promedio, igual que a2_graf_base)

# Colores de clase, en tono claro (regla 3 del usuario).
COLOR_CLASE = {"KALINA": "C6EFCE", "CORREGIBLE": "FCE4D6", "NO_CONVERGIO": "FFC7CE",
               "INVIABLE": "FFEB9C", "KALINA fuera del rango experimental": "E2EFDA"}
FILL_CLASE = {k: PatternFill("solid", fgColor=v) for k, v in COLOR_CLASE.items()}
FILL_ENC = PatternFill("solid", fgColor="DDEBF7")     # encabezado de tabla
FILL_TIT = PatternFill("solid", fgColor="1F4E79")     # banda de titulo de seccion
FILL_COND = PatternFill("solid", fgColor="F2F2F2")    # bloque de condiciones
FUENTE_TIT = Font(bold=True, size=14, color="FFFFFF")
FUENTE_SUB = Font(bold=True, size=12)
FUENTE_ENC = Font(bold=True, size=10)
FUENTE_NOTA = Font(italic=True, size=9)
_FINO = Side(style="thin", color="A6A6A6")
BORDE = Border(left=_FINO, right=_FINO, top=_FINO, bottom=_FINO)


def titulo(ws, fila, texto, ancho=10):
    """Banda de titulo de la hoja (fondo oscuro, texto blanco). Devuelve la fila siguiente."""
    c = ws.cell(fila, 1, texto)
    c.font = FUENTE_TIT
    for j in range(1, ancho + 1):
        ws.cell(fila, j).fill = FILL_TIT
    ws.row_dimensions[fila].height = 22
    return fila + 1


def subtitulo(ws, fila, texto):
    c = ws.cell(fila, 1, texto)
    c.font = FUENTE_SUB
    return fila + 1


def condiciones(ws, fila, lineas):
    """Bloque de condiciones: que quedo fijo (y en que valor) y que se barrio."""
    for linea in lineas:
        c = ws.cell(fila, 1, linea)
        c.alignment = Alignment(wrap_text=False, vertical="top")
        if linea and not linea.startswith("  "):
            c.font = Font(bold=True, size=10)
        else:
            c.font = Font(size=10)
        for j in range(1, 10):
            ws.cell(fila, j).fill = FILL_COND
        fila += 1
    return fila + 1


def encabezado(ws, fila, textos, anchos=None):
    for j, t in enumerate(textos, 1):
        c = ws.cell(fila, j, t)
        c.font = FUENTE_ENC
        c.fill = FILL_ENC
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        c.border = BORDE
    if anchos:
        anchos_col(ws, anchos, 1)
    ws.row_dimensions[fila].height = 30
    return fila + 1


def anchos_col(ws, anchos, desde=1):
    for j, a in enumerate(anchos, desde):
        ws.column_dimensions[L(j)].width = a


def fila_datos(ws, fila, valores, fila_clase=None, clases=None):
    """Escribe una fila con cuadricula; si se dan clases, pinta la fila segun su clase."""
    for j, v in enumerate(valores, 1):
        c = ws.cell(fila, j, v)
        c.border = BORDE
        c.alignment = Alignment(vertical="center", wrap_text=isinstance(v, str) and len(str(v)) > 40)
    if clases:
        relleno = FILL_CLASE.get(clases.get(fila_clase))
        if relleno:
            for j in range(1, len(valores) + 1):
                ws.cell(fila, j).fill = relleno
    return fila + 1


def nota(ws, fila, texto):
    c = ws.cell(fila, 1, texto)
    c.font = FUENTE_NOTA
    return fila + 1


def tabla(ws, fila0, encabezados, filas, anchos=None, clases=None, clave_fila=None):
    """Escribe encabezado + filas y devuelve (fila_encabezado, fila_ultima, n_cols).

    `clases` es un dict {clave_de_fila: clase_en_texto} para pintar las filas; `clave_fila` es
    la posicion (base 0) de la columna cuya celda identifica la fila.
    """
    r = encabezado(ws, fila0, encabezados, anchos)
    for fila in filas:
        clave = fila[clave_fila] if (clases and clave_fila is not None) else None
        r = fila_datos(ws, r, fila, fila_clase=clave, clases=clases)
    return fila0, r - 1, len(encabezados)


def descripcion(ws, fila, lineas):
    """Descripcion corta tras 2 filas de separacion al final de una hoja."""
    fila += 2
    c = ws.cell(fila, 1, "Descripción de la hoja")
    c.font = Font(bold=True, size=11)
    fila += 1
    for linea in lineas:
        ws.cell(fila, 1, linea).font = Font(size=10)
        fila += 1
    return fila
