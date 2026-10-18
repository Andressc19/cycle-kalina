# -*- coding: utf-8 -*-
"""Sensibilidad 1 - x5: motor de graficos nativos de Excel (barras y lineas).

Cada grafico toma sus datos de celdas visibles del propio libro: la tabla de datos que el
helper `bloque` escribe justo a su izquierda. El fallo de la primera version era que los
specs se registraban con `series` vacio, de modo que `_barras`/`_linea` no anadian ninguna
serie y el grafico salia con 0 series (vacio). Aqui `bloque` siempre construye las series a
partir de la tabla, y `dibujar_todos` las traza.

Se reutiliza el acabado de `scripts/analisis_final/a2_graf_estilo.py` (escalas redondas,
titulos sin superponer, un color por serie) para no repetir las trampas de openpyxl ya
resueltas alli (overlay=False, varyColors=False, invertIfNegative=False).
"""
import os
import sys

from openpyxl.chart import BarChart, Reference, ScatterChart, Series
from openpyxl.chart.series import SeriesLabel
from openpyxl.utils import get_column_letter as L

RAIZ = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, os.path.join(RAIZ, "scripts", "analisis_final"))
from a2_graf_estilo import acabado, escala  # noqa: E402

COLORES = ["2E7D32", "1F4E79", "C0392B", "B9770E", "6A1B9A", "00838F", "4E342E"]
GRAFICOS = []                      # specs registrados por x4/x6, dibujados por x7

ALTO_CM_FILA = 0.53                # cm por fila (igual que x2_formato)
FILAS_GRAF = 16                    # filas que ocupa un grafico de 8 cm (control 5 de x8)
COL_GRAF = 12                      # columna de anclaje (L), a la derecha de la tabla de datos


def registrar(spec):
    """Registra un grafico: se dibuja despues de escribir todas las tablas."""
    GRAFICOS.append(spec)
    return spec


def _ref(ws, c):
    return Reference(ws, min_col=c[0], min_row=c[1], max_col=c[2], max_row=c[3])


def _acabado(ch, spec):
    ch.width, ch.height = spec.get("ancho", 16.0), spec.get("alto", 8.0)
    ch.visible_cells_only = False
    ch.display_blanks = "gap"
    acabado(ch, chico=ch.width < 12)
    if isinstance(ch, BarChart):
        ch.x_axis.crosses = ch.y_axis.crosses = "autoZero"
        if spec.get("apiladas"):
            ch.grouping, ch.overlap = "stacked", 100
        elif len(spec.get("series", [])) > 1:
            ch.grouping, ch.overlap = "clustered", 0
    return ch


def _valores(spec):
    return [v for s in spec["series"] for v in s["valores"] if v is not None]


def _barras(ws, spec):
    b = BarChart()
    b.type = "bar" if spec.get("horizontal") else "col"
    b.title, b.style = spec["titulo"], 10
    b.x_axis.title, b.y_axis.title = spec["x"], spec["y"]
    for k, s in enumerate(spec["series"]):
        b.add_data(_ref(ws, s["ref"]), titles_from_data=False)
        ser = b.series[-1]
        col = s.get("color") or COLORES[k % len(COLORES)]
        ser.graphicalProperties.solidFill = col
        ser.graphicalProperties.line.solidFill = col
        ser.invertIfNegative = False
        ser.tx = SeriesLabel(v=s["nombre"])
    b.set_categories(_ref(ws, spec["cat"]))
    escala(b.y_axis, _valores(spec), desde_cero=True)
    if len(spec["series"]) > 1:
        b.legend.position = "b"
    else:
        b.legend = None
    return b


def _linea(ws, spec):
    sc = ScatterChart()
    sc.title, sc.style = spec["titulo"], 13
    sc.x_axis.title, sc.y_axis.title = spec["x"], spec["y"]
    for k, s in enumerate(spec["series"]):
        col = s.get("color") or COLORES[k % len(COLORES)]
        ser = Series(_ref(ws, s["ref"]), _ref(ws, spec["cat"]), title=s["nombre"])
        ser.graphicalProperties.line.solidFill = col
        ser.graphicalProperties.line.width = 22000
        ser.marker.symbol = "circle"
        ser.marker.size = 5
        ser.marker.graphicalProperties.solidFill = col
        ser.marker.graphicalProperties.line.solidFill = col
        ser.smooth = False
        sc.series.append(ser)
    escala(sc.x_axis, spec.get("x_vals", []))
    escala(sc.y_axis, _valores(spec))
    sc.legend.position = "b"
    return sc


def dibujar_todos(wb):
    """Crea todos los graficos registrados, anclados en la celda indicada de su hoja."""
    n = 0
    for spec in GRAFICOS:
        ws = wb[spec["hoja"]]
        ch = _linea(ws, spec) if spec.get("tipo", "barras") == "linea" else _barras(ws, spec)
        _acabado(ch, spec)
        ws.add_chart(ch, spec["celda"])
        n += 1
    return n


def celda_grafico(ws_nombre, fila, col_ref):
    return f"{L(col_ref)}{fila}"


def grafico_sobre_tabla(ws, fila_enc, n_filas, cols_series, *, titulo, x, y, cat_col=1,
                        tipo="barras", apiladas=False, colores=None, nota_txt=None,
                        ancho=16.0, alto=8.0):
    """Registra un grafico cuyos datos ya estan en una tabla escrita (columnas dispersas).

    `fila_enc` es la fila del encabezado y `cols_series` las columnas de las series (la
    columna de categoria es `cat_col`). Deja FILAS_GRAF filas de hueco y devuelve la fila
    siguiente libre.
    """
    ultima = fila_enc + n_filas
    series = []
    for k, c in enumerate(cols_series):
        s = {"nombre": ws.cell(fila_enc, c).value,
             "ref": (c, fila_enc + 1, c, ultima),
             "valores": [ws.cell(r, c).value for r in range(fila_enc + 1, ultima + 1)]}
        if colores and k < len(colores):
            s["color"] = colores[k]
        series.append(s)
    registrar({"tipo": tipo, "hoja": ws.title, "titulo": titulo, "x": x, "y": y,
               "celda": celda_grafico(ws.title, fila_enc, COL_GRAF), "ancho": ancho, "alto": alto,
               "cat": (cat_col, fila_enc + 1, cat_col, ultima), "series": series,
               "apiladas": apiladas, "horizontal": False,
               "x_vals": [ws.cell(r, cat_col).value for r in range(fila_enc + 1, ultima + 1)]})
    fila_nota = max(ultima + 2, fila_enc + FILAS_GRAF)
    if nota_txt:
        from x2_formato import nota
        nota(ws, fila_nota, nota_txt)
    return fila_nota + 2


def bloque(ws, fila, *, titulo, encabezados, filas, x, y, nota_txt,
           tipo="barras", apiladas=False, horizontal=False, x_vals=None,
           ancho=16.0, alto=8.0, colores=None):
    """Escribe la tabla de datos del grafico (con cuadricula) y registra el grafico a su derecha.

    `encabezados[0]` es la categoria y `encabezados[1:]` los nombres de serie. `filas` es una
    lista de filas [categoria, valor_serie1, ...]. El grafico se ancla en la columna L (12),
    junto a la tabla, y debajo queda una celda con el criterio (nota). Devuelve la fila
    siguiente libre (deja FILAS_GRAF filas de hueco para que no se solapen dos graficos).
    """
    from x2_formato import encabezado, fila_datos, nota
    encabezado(ws, fila, list(encabezados))
    r = fila + 1
    for f in filas:
        fila_datos(ws, r, list(f))
        r += 1
    ultima = fila + len(filas)
    series = []
    for k in range(len(encabezados) - 1):
        col = 2 + k
        serie = {"nombre": encabezados[k + 1],
                 "ref": (col, fila + 1, col, fila + len(filas)),
                 "valores": [f[k + 1] for f in filas]}
        if colores and k < len(colores):
            serie["color"] = colores[k]
        series.append(serie)
    registrar({"tipo": tipo, "hoja": ws.title, "titulo": titulo, "x": x, "y": y,
               "celda": celda_grafico(ws.title, fila, COL_GRAF), "ancho": ancho, "alto": alto,
               "cat": (1, fila + 1, 1, fila + len(filas)), "series": series,
               "apiladas": apiladas, "horizontal": horizontal,
               "x_vals": x_vals if x_vals is not None else [f[0] for f in filas]})
    fila_nota = max(ultima + 2, fila + FILAS_GRAF)
    nota(ws, fila_nota, nota_txt)
    return fila_nota + 2


__all__ = ["GRAFICOS", "registrar", "dibujar_todos", "bloque", "grafico_sobre_tabla",
           "celda_grafico", "COL_GRAF", "FILAS_GRAF"]
