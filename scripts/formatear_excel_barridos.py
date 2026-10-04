"""Formatea datos/barridos_e_iteraciones_kalina_consolidado.xlsx (salida de
consolidar_excel_barridos.py): hojas F1_/F2_/F3_, titulos con unidad y color,
anchos, filtro, codigo de colores de clasificacion y descripcion al final de
cada tabla. consolidar_excel_barridos.py lo invoca al final; si el archivo ya
esta formateado, aborta."""
from __future__ import annotations

import re
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

REPO = Path(__file__).resolve().parents[1]
import sys  # noqa: E402
sys.path.insert(0, str(REPO))
from src.excel_rango import GRIS_RANGO, LEYENDA_RANGO, marcar_tabla  # noqa: E402
XLSX = REPO / "datos" / "barridos_e_iteraciones_kalina_consolidado.xlsx"

COLOR_CLASE = {"KALINA": "C6EFCE", "CORREGIBLE": "FFD9A0",
               "NO_CONVERGIO": "F8B4B4", "INVIABLE": "FFF2A8"}
COLOR_TITULO = {"entrada": "BDD7EE", "resultado": "D9D2E9",
                "diagnostico": "B7E1DD", "clase": "F5E1C8", "texto": "D9D9D9"}
ENTRADAS = {"hora", "T_fuente", "T_sumidero", "P_alta", "P_baja", "x_b", "m_b",
            "eta_t", "eta_p", "eps_hrvg", "eps_reg", "eps_cond", "pinch",
            "eps_realista", "etiqueta_punto"}
ORDEN_ENTRADAS = ("T_fuente", "T_sumidero", "P_alta", "P_baja", "x_b", "m_b",
                  "eta_t", "eta_p", "eps_hrvg", "eps_reg", "eps_cond")
# Columnas de entrada ausentes en algunas hojas: constantes de su script.
CONSTANTES = {
    "Fase2_mapa_2d_pbaja_epscond": dict(
        P_alta=3000.0, T_fuente=470.0, T_sumidero=300.032917, m_b=1.0,
        eta_t=0.85, eta_p=0.75, eps_hrvg=0.85, eps_reg=0.75, x_b=0.40),
    "Fase2_ronda5_diagnostico": dict(T_fuente=470.0, T_sumidero=300.032917,
                                     m_b=1.0),
    "Fase2_ronda6_diagnostico": dict(T_sumidero=300.032917, m_b=1.0),
}
UNIDAD = {
    "hora": "h", "T_fuente": "K", "T_sumidero": "K", "P_alta": "kPa",
    "P_baja": "kPa", "m_b": "kg/s", "Wnet": "kW", "Qi": "kW", "Qout": "kW",
    "Wt": "kW", "Wp": "kW", "O2_margen": "K", "O2_margen_K": "K",
    "T_sat_L": "K", "T9_amb": "K", "pinch_K": "K", "pinch": "K", "T2": "K",
    "T_bub": "K", "T_dew": "K", "T9": "K", "T_burbuja": "K", "T_bur": "K",
    "margen": "K", "margen_O2": "K", "margen_S5": "K", "margen_O2_real": "K",
    "margen_ligante": "K o -", "margen_m": "K", "margen_0": "K",
    "margen_p": "K", "P_estrella": "kPa", "P_star": "kPa", "s3": "kJ/kg·K",
    "dEta": "-", "dH4": "kJ/kg", "dh4s": "kJ/kg", "dT4s": "K",
    "iteraciones_pbaja": "n", "n_evaluaciones": "n", "n_recurrencias": "n",
    "n_fallos": "n", "A_recurrencias": "n", "sobrecosto_pct": "%",
    "s_por_punto": "s", "s_30_puntos": "s", "m3": "kg/s", "h3": "kJ/kg",
    "h4": "kJ/kg", "T3": "K", "T4": "K", "motor_real_eta": "-",
    "B_dh4s": "kJ/kg", "Estado": "-", "x (NH3) [-]": "-",
}
RE_UNIDAD = [
    (r"^T\d+$", "K"), (r"^T4[s]?_(m|0|p|teqp|real)$", "K"),
    (r"^T4s_(teqp|real)$", "K"), (r"^h4s?_(m|0|p|teqp|real)$", "kJ/kg"),
    (r"^Wnet_(m|0|p|teqp|real)$", "kW"), (r"^margen_O2", "K"),
    (r"^margen_S5$", "K"), (r"^margen_O[15]$", "-"),
    (r"^(t_.*_s|tiempo_s|t_s|B_t_s|t_[A-Z]+_s)$", "s"),
    (r"^dh4s_", "kJ/kg"), (r"^(eta|q|x|eps)(_.*)?$", "-"), (r"^q\d.*$", "-"),
    (r"^x\d+$", "-"), (r"^eta_", "-"), (r"^d_eta$", "-"),
    (r"^Wnet_", "kW"),
]
TEXTO = re.compile(r"^(mensaje|fallas.*|detalle.*|error.*|err_.*|criterio_ligante"
                   r"|motor|caso|etiqueta.*|opcion|base|nota|fase4s|estado_filtro"
                   r"|B_error|rama_fisica|detalle_rango|columnas_rango)$")
BOOL = re.compile(r"^(convergio|pbaja_convergio|pbaja_encontrada|al_filo_O1"
                  r"|motor_confirmado|spot_check|estable|clases_distintas"
                  r"|cambia_clasificacion|verificacion_incompleta|verificado_.*"
                  r"|verif_incompleta_A|A_verificacion_incompleta|B_verificado"
                  r"|cumple_N2|fuera_rango)$")
RESULTADOS = {"eta", "Wnet", "O2_margen", "eta_real", "Wnet_real"}


def limpia(h: str) -> str:
    h = str(h).replace("�", "·").replace("[", "(").replace("]", ")")
    return re.sub(r"\s*\((?:-|K|kPa|kW|kg/s|s|h|n|%|kJ/kg|kJ/kg·K|K o -|texto"
                  r"|categoría|V/F)\)$", "", h)


def unidad(h: str) -> str | None:
    if "clasif" in h or h.startswith("clas_"):
        return "categoría"
    if TEXTO.match(h):
        return "texto"
    if BOOL.match(h):
        return "V/F"
    if h in UNIDAD:
        return UNIDAD[h]
    if h.startswith("T_") or h.startswith("T4"):
        return "K"
    for pat, u in RE_UNIDAD:
        if re.match(pat, h):
            return u
    return None


# En esta hoja eta_p/eta_m/eta_0 son resultados (P* +/- 0.5 kPa), no entradas.
HOJAS_ETA_RESULTADO = ("0924_estabilidad_22",)


def es_entrada(h: str, hoja: str) -> bool:
    return h in ENTRADAS and not (hoja in HOJAS_ETA_RESULTADO and h == "eta_p")


def tipo_col(h: str, hoja: str = "") -> str:
    if "clasif" in h or h.startswith("clas_"):
        return "clase"
    if TEXTO.match(h):
        return "texto"
    if es_entrada(h, hoja):
        return "entrada"
    if h in RESULTADOS or h.startswith(("eta", "Wnet")):
        return "resultado"
    return "diagnostico"


def fmt(v) -> str:
    if isinstance(v, float):
        return f"{v:.6g}"
    return str(v)


def rango(vals: list) -> str:
    u = sorted(set(vals))
    if len(u) <= 4:
        return ", ".join(fmt(x) for x in u)
    return f"{fmt(u[0])} a {fmt(u[-1])} ({len(u)} valores)"


def nuevo_nombre(n: str) -> str:
    m = re.match(r"^Fase(\d)_(.*)$", n)
    return f"F{m.group(1)}_{m.group(2)}" if m else n


def leer_tabla(ws):
    filas = []
    for r in ws.iter_rows(values_only=True):
        if all(v is None for v in r):
            break
        filas.append(list(r))
    cab = filas[0]
    while cab and cab[-1] is None:
        cab.pop()
    return cab, [f[:len(cab)] for f in filas[1:]]


def anade_constantes(nombre: str, cab: list, filas: list):
    cte = CONSTANTES.get(nombre)
    if not cte:
        return cab, filas, []
    faltan = [c for c in ORDEN_ENTRADAS if c not in cab and c in cte]
    if not faltan:
        return cab, filas, []
    nuevos = [c for c in ORDEN_ENTRADAS if c in cab or c in faltan]
    resto = [c for c in cab if c not in ORDEN_ENTRADAS]
    orden = nuevos + resto
    out = []
    for f in filas:
        d = dict(zip(cab, f))
        d.update({c: cte[c] for c in faltan})
        out.append([d[c] for c in orden])
    return orden, out, faltan


def descripcion(cab, filas, agregadas, extra, hoja) -> list[str]:
    col = {c: [f[i] for f in filas] for i, c in enumerate(cab)}
    itera, fijo = [], []
    for c in cab:
        if not es_entrada(c, hoja) or c == "etiqueta_punto":
            continue
        vals = [v for v in col[c] if v is not None]
        if not vals:
            continue
        txt = f"{c} ({unidad(c) or '-'}): "
        if len(set(vals)) > 1:
            itera.append(txt + rango(vals))
        else:
            fijo.append(txt + fmt(vals[0]))
    lin = ["Descripción"]
    lin.append("Se iteró: " + ("; ".join(itera) if itera else "no aplica"))
    for i in range(0, len(fijo), 4):
        lin.append(("Fijo: " if i == 0 else "      ") + "; ".join(fijo[i:i + 4]))
    if agregadas:
        lin.append("Columnas añadidas con su valor constante: "
                   + ", ".join(agregadas))
    cl = [c for c in cab if "clasif" in c or c.startswith("clas_")]
    if cl:
        for c in cl[:2]:
            n = {}
            for v in col[c]:
                if v is not None:
                    n[v] = n.get(v, 0) + 1
            lin.append(f"{c}: " + " · ".join(f"{k} {v}" for k, v in n.items()))
    lin.extend(extra)
    return lin


def estilo_titulo(cell, tipo):
    cell.fill = PatternFill("solid", fgColor=COLOR_TITULO[tipo])
    cell.font = Font(bold=True, color="1F1F1F")
    cell.alignment = Alignment(horizontal="center", vertical="center",
                               wrap_text=True)
    cell.border = Border(*(Side(style="thin", color="8C8C8C"),) * 4)


def escribe_hoja(wb, nombre, cab, filas, extra, bloques):
    ws = wb.create_sheet(nuevo_nombre(nombre))
    cab = [limpia(c).replace("O2_margen_K", "O2_margen") for c in cab]
    tipos = [tipo_col(c, nombre) for c in cab]
    ws.append([f"{c} ({unidad(c) or '-'})" for c in cab])
    for f in filas:
        ws.append(f)
    for j, t in enumerate(tipos, 1):
        estilo_titulo(ws.cell(1, j), t)
        w = max(len(str(ws.cell(1, j).value)) + 2,
                *(len(fmt(f[j - 1])) + 2 for f in filas if f[j - 1] is not None),
                10)
        if t == "texto" and cab[j - 1] in ("mensaje", "detalle_error", "error",
                                           "detalle", "fallas"):
            w = 80
            for r in range(2, ws.max_row + 1):
                ws.cell(r, j).alignment = Alignment(wrap_text=True,
                                                    vertical="top")
        ws.column_dimensions[get_column_letter(j)].width = min(w, 80)
        if t == "clase":
            for r in range(2, ws.max_row + 1):
                c = ws.cell(r, j)
                if c.value in COLOR_CLASE:
                    c.fill = PatternFill("solid", fgColor=COLOR_CLASE[c.value])
    if marcar_tabla(ws, cab):
        extra = [*extra, LEYENDA_RANGO]
    ws.row_dimensions[1].height = 34
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cab))}{ws.max_row}"
    fin = ws.max_row
    r = fin + 3
    for k, txt in enumerate(descripcion(cab, filas, [], extra, nombre)):
        c = ws.cell(r, 1, txt)
        if k == 0:
            c.font = Font(bold=True, size=12)
        r += 1
    for hcab, hfilas in bloques:
        r += 2
        for j, h in enumerate(hcab, 1):
            estilo_titulo(ws.cell(r, j, limpia(h)), "texto")
        r += 1
        for f in hfilas:
            for j, v in enumerate(f, 1):
                ws.cell(r, j, v)
            r += 1
    return ws


# Sensibilidades de Fase 3 con el T_amb_diseno antiguo (30.4 C): descartadas.
HOJAS_DESCARTADAS = ("Fase3_sensibilidad_palta", "Fase3_sensibilidad_pbaja",
                     "Fase3_sensibilidad_xb")
# Fase 3 del Resumen tras descartarlas: hojas, puntos, KALINA, CORREGIBLE.
RESUMEN_F3 = (8, 180, 91, 84)


def main() -> None:
    wb = openpyxl.load_workbook(XLSX)
    if any(n.startswith("F1_") for n in wb.sheetnames):
        raise SystemExit("Ya formateado: regenera con consolidar_excel_barridos.py")
    for n in HOJAS_DESCARTADAS:
        if n in wb.sheetnames:
            del wb[n]
    for fila in wb["Resumen"].iter_rows(min_row=5):
        if fila[0].value == "Fase 3":
            for c, v in zip(fila[1:5], RESUMEN_F3):
                c.value = v
    anclas = {}
    for n in ("Fase2_ancla_balance", "Fase2_ancla_estados",
              "Fase3_ancla_balance", "Fase3_ancla_estados"):
        anclas[n] = leer_tabla(wb[n])
    sin_unidad = set()
    nuevas = {}
    for ws in list(wb):
        if ws.title == "Resumen" or "_ancla_" in ws.title:
            continue
        cab, filas = leer_tabla(ws)
        cab, filas, agregadas = anade_constantes(ws.title, cab, filas)
        extra = []
        if agregadas:
            extra.append("Columnas añadidas con su valor constante: "
                         + ", ".join(agregadas))
        ph = re.match(r"^Fase(\d)_sensibilidad_P_baja$", ws.title)
        bloques = []
        if ph:
            f = ph.group(1)
            bloques = [anclas[f"Fase{f}_ancla_balance"],
                       anclas[f"Fase{f}_ancla_estados"]]
            extra.append("Punto ancla (base de las sensibilidades de la fase): "
                         "balance y estados abajo.")
        elif re.match(r"^Fase[23]_sensibilidad_", ws.title):
            extra.append(f"Punto ancla: ver F{ws.title[4]}_sensibilidad_P_baja.")
        for c in cab:
            if unidad(limpia(c)) is None:
                sin_unidad.add((ws.title, c))
        nuevas[ws.title] = (cab, filas, extra, bloques)
    # reconstruye en orden
    resumen = wb["Resumen"]
    for n in list(wb.sheetnames):
        if n != "Resumen":
            del wb[n]
    orig = openpyxl.load_workbook(XLSX)
    for n in orig.sheetnames:
        if n in nuevas:
            escribe_hoja(wb, n, *nuevas[n])
    formatea_resumen(resumen)
    wb.save(XLSX)
    print("Hojas:", wb.sheetnames)
    print("Sin unidad asignada ('-' por defecto):", sorted(sin_unidad))


def formatea_resumen(ws) -> None:
    for j in range(1, 9):
        estilo_titulo(ws.cell(4, j), "entrada")
    for fila in range(5, 5 + 18):
        ws.cell(fila, 8).alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 30
    for L in "BCDEFG":
        ws.column_dimensions[L].width = 14
    ws.column_dimensions["H"].width = 110
    for r in range(5, ws.max_row + 1):
        if ws.cell(r, 1).value and str(ws.cell(r, 1).value).startswith(
                "- Las hojas *_ancla_*"):
            ws.cell(r, 1).value = ("- Los puntos ancla de F2 y F3 estan al final "
                                   "de F2_/F3_sensibilidad_P_baja (no son "
                                   "iteraciones).")
    r = ws.max_row + 2
    ws.cell(r, 1, "Código de colores de clasificación:").font = Font(bold=True)
    for k, (n, c) in enumerate(COLOR_CLASE.items()):
        cell = ws.cell(r + 1 + k, 1, n)
        cell.fill = PatternFill("solid", fgColor=c)
    k = len(COLOR_CLASE)
    ws.cell(r + 1 + k, 1, "FUERA_RANGO").fill = PatternFill("solid", fgColor=GRIS_RANGO)
    ws.cell(r + 1 + k, 2, LEYENDA_RANGO)


if __name__ == "__main__":
    main()
