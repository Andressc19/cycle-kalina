# -*- coding: utf-8 -*-
"""Sens2 - s2_f4t: hoja `F4 Ciclo vs T_fuente` (eficiencia del ciclo frente a T_fuente por bandas).

Lee `resultados/2026-10-18_sens2_F1/p6_ciclo_T.csv` y arma la hoja: tabla completa por banda,
grafico de lineas de la mejor eta, P90 y mediana frente a T_fuente y grafico de barras de
P(KALINA) con y sin NO_CONVERGIO. Copia el patron de `s2_f6.py`.
"""
import os
import sys

import pandas as pd

AQUI = os.path.dirname(os.path.abspath(__file__))
RAIZ = os.path.dirname(os.path.dirname(AQUI))
sys.path.insert(0, os.path.join(RAIZ, "scripts", "sens1"))
sys.path.insert(0, AQUI)
from x2_formato import condiciones, descripcion, nota, subtitulo, tabla, titulo  # noqa: E402
from x5_graficos import bloque  # noqa: E402
from s2_comun import numero  # noqa: E402

F1 = os.path.join(RAIZ, "resultados", "2026-10-18_sens2_F1")
MIN_KALINA = 10
PROHIBIDAS = ["confiable", "no citable", "fuera de rango del motor", "indicio", "extrapol", "sin respaldo"]
COLS = ["banda", "T_centro_C", "n", "n_kalina", "n_corregible", "n_inviable", "n_no_convergio",
        "p_kalina", "p_kalina_si_convergio", "eta_max", "id_punto", "T_fuente", "P_alta", "P_baja",
        "x_b", "eta_t", "eps_reg", "eps_cond", "fuera_rango", "eta_max_dentro", "eta_p90",
        "eta_mediana", "wnet_mb_max", "wnet_mb_mediana", "eta_mediana_lhs", "n_kalina_lhs", "n_kalina_dentro", "eta_p90_dentro", "eta_mediana_dentro"]


def _d():
    return pd.read_csv(os.path.join(F1, "p6_ciclo_T.csv"))


def _c(x):
    v = numero(x, 3)
    return "" if v is None else v


def _fila(x):
    out = []
    for col in COLS:
        if col in ("banda", "id_punto"):
            out.append(str(x[col]) if pd.notna(x[col]) else "")
        elif col == "fuera_rango":
            out.append(str(x[col]))
        else:
            out.append(_c(x[col]))
    return out


def hoja_f4_ciclo_T(wb, datos):
    ws = wb.create_sheet("F4 Ciclo vs T_fuente")
    d = _d()
    f = titulo(ws, 1, "F4 Ciclo vs T_fuente: mejor eficiencia del ciclo por banda de T_fuente", ancho=10)
    f = condiciones(ws, f, [
        "Base: %d puntos de p2_dataset_dedup.csv (todos los estudios); sin corridas nuevas del motor." % int(d["n"].sum()),
        "Bandas de T_fuente de 20 K ([340, 360) a [560, 580)); KALINA (total) con eta valida.",
        "El maximo por banda es el mayor valor muestreado, no el optimo: no hay optimizador.",
        "Dentro del rango experimental = fuera_rango == False (liquido > 420 K).",
        "Fuente: resultados/2026-10-18_sens2_F1/p6_ciclo_T.csv."])
    f = subtitulo(ws, f, "1. Tabla completa por banda de T_fuente")
    _, f, _ = tabla(ws, f, COLS, [_fila(x) for _, x in d.iterrows()])
    f += 1
    f = subtitulo(ws, f, "2. Eficiencia frente a T_fuente (bandas con al menos %d KALINA)" % MIN_KALINA)
    sel = d[(d["n_kalina_dentro"] >= MIN_KALINA) & d["eta_max_dentro"].notna()]
    xv = [numero(x["T_centro_C"], 1) for _, x in sel.iterrows()]
    fl = [[xv[i], _c(x["eta_max_dentro"]), _c(x["eta_p90_dentro"]), _c(x["eta_mediana_dentro"])]
          for i, (_, x) in enumerate(sel.iterrows())]
    f = bloque(ws, f, titulo="eta maxima, P90 y mediana de los KALINA dentro del rango experimental frente a T_fuente",
               encabezados=["T_centro_C (C)", "eta_max_dentro (%)", "eta_p90_dentro (%)", "eta_mediana_dentro (%)"],
               filas=fl, x="T_centro_C (C)", y="eta (%)", tipo="linea", x_vals=xv,
               nota_txt="Fuente: p6_ciclo_T.csv. Solo bandas con al menos %d KALINA dentro del rango experimental (columna n_kalina_dentro); "
                        "las bandas calientes tienen casi todos sus KALINA fuera del rango (liquido > 420 K) y no aparecen. El maximo es el mayor "
                        "valor muestreado y crece con el tamano del conjunto." % MIN_KALINA)
    f = subtitulo(ws, f, "3. P(KALINA) por banda de T_fuente (efecto de NO_CONVERGIO)")
    fl2 = [["%.0f C" % x["T_centro_C"], _c(100 * x["p_kalina"]), _c(100 * x["p_kalina_si_convergio"])]
           for _, x in d.iterrows()]
    f = bloque(ws, f, titulo="P(KALINA) por banda, con y sin NO_CONVERGIO en el denominador",
               encabezados=["Centro de la banda de T_fuente (C)", "p_kalina (%)", "p_kalina_si_convergio (%)"],
               filas=fl2, x="Centro de la banda de T_fuente (C)", y="P(KALINA) (%)",
               nota_txt="Fuente: p6_ciclo_T.csv. NO_CONVERGIO sesga hacia abajo la P(KALINA); "
                        "p_kalina_si_convergio usa (n - NO_CONVERGIO) como denominador y se separa mas "
                        "en las bandas frias.")
    f = nota(ws, f, "El maximo por banda es el mayor valor muestreado y no el optimo del ciclo (no hay "
                    "optimizador). El panorama de eta frente a T_fuente se aplana desde aproximadamente "
                    "150 C segun los barridos de una variable (conclusion E8).")
    descripcion(ws, f, ["Eficiencia del ciclo frente a la temperatura de la fuente, por bandas de 20 K,",
                        "sobre los puntos ya calculados. El maximo es el mayor valor muestreado."])
    return ws


HOJAS_F4T = [hoja_f4_ciclo_T]


# --- Prueba en proceso nuevo: crea la hoja, dibuja los graficos, guarda y verifica ---
def _refs(ch):
    out = []
    for s in ch.series:
        for attr in ("val", "cat", "xVal", "yVal"):
            o = getattr(s, attr, None)
            if o is None:
                continue
            for sub in ("numRef", "strRef"):
                r = getattr(o, sub, None)
                if r is not None and getattr(r, "f", None):
                    out.append(r.f)
    return out


def _vacios(ws):
    from openpyxl.utils import range_boundaries
    mal = []
    for ch in ws._charts:
        for f in _refs(ch):
            c1, r1, c2, r2 = range_boundaries(f.split("!")[-1].replace("$", ""))
            for rr in range(r1, r2 + 1):
                for cc in range(c1, c2 + 1):
                    v = ws.cell(rr, cc).value
                    if v is None or (isinstance(v, str) and v.strip() == ""):
                        mal.append((rr, cc))
    return mal


def _prohibidas(ws):
    mal = []
    for fila in ws.iter_rows():
        for c in fila:
            if isinstance(c.value, str):
                t = c.value.lower()
                for p in PROHIBIDAS:
                    if p in t:
                        mal.append((c.coordinate, p))
    return mal


def _probar():
    from openpyxl import Workbook, load_workbook

    from x5_graficos import GRAFICOS, dibujar_todos
    GRAFICOS.clear()
    wb = Workbook()
    wb.remove(wb.active)
    hoja_f4_ciclo_T(wb, {})
    n = dibujar_todos(wb)
    os.makedirs(F1, exist_ok=True)
    ruta = os.path.join(F1, "_prueba_f4t.xlsx")
    wb.save(ruta)
    wb2 = load_workbook(ruta)
    ws = wb2["F4 Ciclo vs T_fuente"]
    vacias, prohib = _vacios(ws), _prohibidas(ws)
    print("hoja F4 Ciclo vs T_fuente presente:", "F4 Ciclo vs T_fuente" in wb2.sheetnames)
    print("graficos dibujados:", n, "| graficos en la hoja:", len(ws._charts))
    print("celdas vacias en tablas de graficos:", len(vacias))
    print("palabras prohibidas:", len(prohib))
    os.remove(ruta)
    assert len(ws._charts) == 2, "se esperaban 2 graficos"
    assert not vacias, "hay celdas vacias en las tablas de los graficos"
    assert not prohib, "hay palabras prohibidas"
    print("PRUEBA OK")


if __name__ == "__main__":
    _probar()
