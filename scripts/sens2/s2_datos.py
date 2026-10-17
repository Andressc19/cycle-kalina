# -*- coding: utf-8 -*-
"""Sens2 - s2_datos: hojas de datos `D1` a `D4` (valores numericos como numeros, con filtro)."""
import math
import os
import sys

from openpyxl.styles import Alignment

RAIZ = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, os.path.join(RAIZ, "scripts", "sens1"))
from x2_formato import BORDE, FILL_ENC, FUENTE_ENC, anchos_col, descripcion  # noqa: E402

CAMPOS_D1 = [
    ("Punto", "id_punto", "txt"), ("Estudio", "estudio", "int"), ("Centro", "centro", "txt"),
    ("Clasificacion", "clasificacion", "txt"), ("KALINA (total) (si/no)", "kalina", "si"),
    ("Fuera del rango experimental (si/no)", "fuera_rango", "si"),
    ("T_fuente (K)", "T_fuente", "num"), ("T_sumidero (K)", "T_sumidero", "num"),
    ("P_alta (kPa)", "P_alta", "num"), ("P_baja (kPa)", "P_baja", "num"), ("x_b (-)", "x_b", "num"),
    ("eta_t (-)", "eta_t", "num"), ("eta_p (-)", "eta_p", "num"),
    ("eps_HRVG (-)", "eps_hrvg", "num"), ("eps_reg (-)", "eps_reg", "num"),
    ("eps_cond (-)", "eps_cond", "num"), ("eta (%)", "eta_pct", "num"), ("Wnet (kW)", "Wnet", "num"),
    ("D (-)", "D", "num"), ("viol_O1 (-)", "viol_O1", "num"), ("viol_O2 (-)", "viol_O2", "num"),
    ("viol_O3 (-)", "viol_O3", "num"), ("viol_O5 (-)", "viol_O5", "num"),
    ("r_P (-)", "rp", "num"), ("Q_reg/Q_i (-)", "qreg_qi", "num"),
    ("dh_turb (kJ/kg)", "dh_turb", "num"), ("Qi (kW)", "Qi", "num"), ("Qreg (kW)", "Qreg", "num"),
    ("q4, titulo de turbina (-)", "q4", "num"), ("t_s (s)", "t_s", "num"),
    ("m_b (kg/s)", "m_b", "num"), ("Wnet/m_b (kW/kg/s)", "wnet_mb", "num"),
]
CAMPOS_D2 = [
    ("Bloque", "centro", "txt"), ("Variable", "variable", "txt"), ("Forma", "forma", "txt"),
    ("Direccion", "direccion", "txt"), ("Valor de la frontera", "valor_frontera", "num"),
    ("Valor anterior", "valor_anterior", "num"), ("Clase anterior", "clase_anterior", "txt"),
    ("Fraccion anterior", "frac_anterior", "num"), ("n anterior", "n_anterior", "int"),
    ("Valor siguiente", "valor_siguiente", "num"), ("Clase siguiente", "clase_siguiente", "txt"),
    ("Fraccion siguiente", "frac_siguiente", "num"), ("n siguiente", "n_siguiente", "int"),
]
CAMPOS_D3 = [
    ("Punto", "id_punto", "txt"), ("Estudio", "estudio", "int"), ("Bloque", "centro", "txt"),
    ("D (-)", "D", "num"), ("Violacion dominante", "dominante", "txt"),
    ("Valor de la dominante", "valor_dominante", "num"), ("viol_O1 (-)", "viol_O1", "num"),
    ("viol_O2 (-)", "viol_O2", "num"), ("viol_O3 (-)", "viol_O3", "num"),
    ("viol_O5 (-)", "viol_O5", "num"),
]
CAMPOS_D4 = [
    ("Punto (CORREGIBLE)", "id_punto", "txt"), ("Bloque", "centro", "txt"), ("D (-)", "D", "num"),
    ("Violacion dominante", "viol_dominante", "txt"), ("Punto vecino (KALINA)", "vecino_id", "txt"),
    ("L1 al vecino", "l1", "num"), ("n variables distintas", "nvars_diff", "int"),
    ("Variable que cambia", "var_cambia", "txt"), ("Valor en el CORREGIBLE", "val_origen", "num"),
    ("Valor en el vecino", "val_destino", "num"), ("Direccion", "direccion", "txt"),
    ("eta CORREGIBLE (%)", "eta_cor", "num"), ("eta vecino (%)", "eta_vecino", "num"),
]


def _fmt(v, tipo):
    if v is None or (isinstance(v, float) and math.isnan(v)):
        return None
    if tipo == "si":
        return "si" if bool(v) else "no"
    if tipo == "txt":
        return str(v)
    if tipo == "int":
        return int(v)
    return float(v)


def _hoja(wb, nombre, encabezado_txt, campos, df, desc):
    ws = wb.create_sheet(nombre)
    ws.cell(1, 1, encabezado_txt).font = FUENTE_ENC
    r = 3
    for j, (t, _, _) in enumerate(campos, 1):
        c = ws.cell(r, j, t)
        c.font = FUENTE_ENC
        c.fill = FILL_ENC
        c.border = BORDE
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    anchos_col(ws, [max(10, min(len(t) + 2, 22)) for t, _, _ in campos], 1)
    enc = r
    for _, fila in df.iterrows():
        r += 1
        for j, (_, clave, tipo) in enumerate(campos, 1):
            c = ws.cell(r, j, _fmt(fila.get(clave), tipo))
            c.border = BORDE
    ultima = r
    ws.auto_filter.ref = "A%d:%s%d" % (enc, _letra(len(campos)), ultima)
    descripcion(ws, ultima, desc)
    return ws


def _letra(n):
    s = ""
    while n:
        n, m = divmod(n - 1, 26)
        s = chr(65 + m) + s
    return s


def hoja_d1(wb, datos):
    df = datos["dataset_dedup"]
    return _hoja(wb, "D1 Datos puntos", "D1. Todos los puntos analizados (sin duplicados)",
                 CAMPOS_D1, df, ["Fuente: p2_dataset_dedup.csv.",
                                 "Una fila por punto; %d puntos." % len(df),
                                 "KALINA (total), CORREGIBLE, NO_CONVERGIO e INVIABLE; D y las "
                                 "violaciones O1/O2/O3/O5 salen del motor."])


def hoja_d2(wb, datos):
    df = datos["fronteras"]
    return _hoja(wb, "D2 Datos fronteras", "D2. Fronteras de los barridos de una variable",
                 CAMPOS_D2, df, ["Fuente: p1_fronteras.csv.",
                                 "Una fila por frontera; %d fronteras." % len(df),
                                 "Forma = escalon, rampa o ventana; es el nivel donde cambia la "
                                 "clase, con el nivel anterior y el siguiente."])


def hoja_d3(wb, datos):
    df = datos["zonas_grises"]
    return _hoja(wb, "D3 Datos zonas grises", "D3. Zonas grises (CORREGIBLE con D bajo)",
                 CAMPOS_D3, df, ["Fuente: p1_zonas_grises.csv.",
                                 "Una fila por zona gris; %d puntos." % len(df),
                                 "Zona gris = CORREGIBLE con D <= cuantil 25 % de D entre los "
                                 "CORREGIBLE; dominante es la violacion de mayor magnitud."])


def hoja_d4(wb, datos):
    df = datos["distancia"]
    return _hoja(wb, "D4 Datos distancia", "D4. Distancia a la frontera de cada CORREGIBLE",
                 CAMPOS_D4, df, ["Fuente: p2_distancia.csv.",
                                 "Una fila por CORREGIBLE; %d puntos." % len(df),
                                 "vecino_id es el KALINA mas cercano en variables normalizadas; "
                                 "var_cambia y la direccion indican el cambio minimo que lo rescata."])


HOJAS_DATOS = [hoja_d1, hoja_d2, hoja_d3, hoja_d4]
