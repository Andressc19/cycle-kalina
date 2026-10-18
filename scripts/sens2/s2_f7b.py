# -*- coding: utf-8 -*-
"""Sens2 - s2_f7b: prueba de `s2_f7.hoja_f7_rigor` (libro de prueba, no toca el principal).

Crea un libro vacio, escribe la hoja F7, dibuja los graficos registrados y comprueba sobre el
archivo guardado: que la hoja existe, cuantos graficos tiene, que ninguna celda de las tablas
de esos graficos esta vacia y que ninguna cadena usa palabras prohibidas. Se ejecuta en un
proceso nuevo porque `x5_graficos.GRAFICOS` es una lista global que acumula specs.
"""
import os
import sys

from openpyxl import Workbook, load_workbook
from openpyxl.utils import range_boundaries

AQUI = os.path.dirname(os.path.abspath(__file__))
RAIZ = os.path.dirname(os.path.dirname(AQUI))
sys.path.insert(0, os.path.join(RAIZ, "scripts", "sens1"))
sys.path.insert(0, AQUI)
from s2_f7 import R, hoja_f7_rigor  # noqa: E402
from x5_graficos import dibujar_todos  # noqa: E402

PROHIBIDAS = ("confiable", "no citable", "fuera de rango del motor", "indicio", "extrapol", "sin respaldo")


def _refs(ch):
    out = []
    for s in ch.series:
        for src in (s.val, s.cat):
            nr = getattr(src, "numRef", None) or getattr(src, "strRef", None)
            if nr is not None and nr.f:
                out.append(nr.f)
    return out


def main():
    ruta = os.path.join(R, "_prueba_f7.xlsx")
    wb = Workbook()
    wb.remove(wb.active)
    hoja_f7_rigor(wb, None)
    n = dibujar_todos(wb)
    wb.save(ruta)
    wsl = load_workbook(ruta)
    ws = wsl["F7 Rigor estadistico"]
    vacias = []
    for ch in ws._charts:
        for ref in _refs(ch):
            c1, r1, c2, r2 = range_boundaries(ref.split("!")[-1].replace("$", ""))
            for rr in range(r1, r2 + 1):
                for cc in range(c1, c2 + 1):
                    if ws.cell(rr, cc).value in (None, ""):
                        vacias.append("%s%d" % (chr(64 + cc), rr))
    textos = [str(c.value).lower() for row in ws.iter_rows() for c in row if isinstance(c.value, str)]
    malas = [p for p in PROHIBIDAS if any(p in t for t in textos)]
    print("hoja existe:", "F7 Rigor estadistico" in wsl.sheetnames)
    print("graficos:", len(ws._charts), "| dibujar_todos:", n)
    print("filas de la hoja:", ws.max_row)
    print("celdas vacias en tablas de graficos:", len(vacias), vacias[:5])
    print("frases prohibidas:", malas)
    assert "F7 Rigor estadistico" in wsl.sheetnames
    assert len(ws._charts) >= 5
    assert not vacias
    assert not malas
    print("PRUEBA OK")


if __name__ == "__main__":
    main()
