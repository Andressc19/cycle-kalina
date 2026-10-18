# -*- coding: utf-8 -*-
"""Sens2 - s2_f5: hoja `F5 Superficies de respuesta` (mapas de calor de eta y ajuste cuadratico)."""
import os
import sys

import numpy as np
from openpyxl.chart import Reference, ScatterChart, Series
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment

AQUI = os.path.dirname(os.path.abspath(__file__))
RAIZ = os.path.dirname(os.path.dirname(AQUI))
sys.path.insert(0, os.path.join(RAIZ, "scripts", "sens1"))
sys.path.insert(0, os.path.join(RAIZ, "scripts", "analisis_final"))
sys.path.insert(0, AQUI)
from x2_formato import (BORDE, COLOR_CLASE, FILL_CLASE, condiciones, descripcion, encabezado, fila_datos,  # noqa: E402
                        nota, subtitulo, tabla, titulo)
from x5_graficos import COL_GRAF, FILAS_GRAF, celda_grafico  # noqa: E402
from s2_calc import BLOQUES, ajuste_cuadratico, malla_clases, num  # noqa: E402
from s2_comun import COLORES  # noqa: E402
from a2_graf_estilo import acabado, escala  # noqa: E402

PARES = {"B2": ("P_baja", "kPa", "x_b", "-"), "B3": ("P_alta", "kPa", "P_baja", "kPa")}
CORTO = {"C1 Husavik": "C1", "C2 ancla A*": "C2", "C3 caliente 150 C": "C3", "C4 frio 110 C": "C4"}
CLASE = {"C": "CORREGIBLE", "NC": "NO_CONVERGIO", "I": "INVIABLE"}


def _matriz(ws, f, eta, cls, av, bv, a, ua, c, uc):
    """Mapa de calor: eta (%) en las celdas KALINA y C / NC / I en las demas, con el color de su clase."""
    encabezado(ws, f, ["filas: %s (%s) - columnas: %s (%s)" % (a, ua, c, uc)] + [str(x) for x in bv])
    f0 = f + 1
    for i, x in enumerate(av):
        ws.cell(f0 + i, 1, float(x)).border = BORDE
        for j in range(len(bv)):
            cel = ws.cell(f0 + i, 2 + j)
            cel.border = BORDE
            cel.alignment = Alignment(horizontal="center")
            if cls[i, j] == "K":
                cel.value = num(eta[i, j], 2)
            else:
                cel.value = cls[i, j]
                if cls[i, j] in CLASE:
                    cel.fill = FILL_CLASE[CLASE[cls[i, j]]]
    rango = "B%d:%s%d" % (f0, chr(65 + len(bv)), f0 + len(av) - 1)
    ws.conditional_formatting.add(rango, ColorScaleRule(start_type="min", start_color="F8E5A8", mid_type="percentile",
                                                        mid_value=50, mid_color="A9D18E", end_type="max",
                                                        end_color="1E7B3C"))
    return f0 + len(av)


def _largo(eta, av, bv):
    """Tabla larga (nivel de la 2.a variable, valor de la 1.a, eta) solo de celdas KALINA, ordenada por nivel."""
    filas = []
    for j, y in enumerate(bv):
        for i, x in enumerate(av):
            if np.isfinite(eta[i, j]):
                filas.append([float(y), float(x), num(eta[i, j], 3)])
    return filas


def _grafico(ws, enc, filas, a, c, titulo_g):
    ch = ScatterChart()
    ch.title, ch.style = titulo_g, 13
    ch.x_axis.title, ch.y_axis.title = a, "eta (%)"
    niveles = list(dict.fromkeys(x[0] for x in filas))
    r = enc + 1
    for k, n in enumerate(niveles):
        m = sum(1 for x in filas if x[0] == n)
        if m >= 2:
            s = Series(Reference(ws, min_col=3, min_row=r, max_row=r + m - 1),
                       Reference(ws, min_col=2, min_row=r, max_row=r + m - 1), title="%s = %s" % (c, n))
            col = COLORES[k % len(COLORES)]
            s.graphicalProperties.line.solidFill = col
            s.graphicalProperties.line.width = 20000
            s.marker.symbol = "circle"
            s.marker.size = 4
            s.marker.graphicalProperties.solidFill = col
            s.marker.graphicalProperties.line.solidFill = col
            s.smooth = False
            ch.series.append(s)
        r += m
    escala(ch.x_axis, [x[1] for x in filas])
    escala(ch.y_axis, [x[2] for x in filas])
    ch.legend.position = "b"
    ch.width, ch.height = 18.0, 8.0
    acabado(ch)
    ws.add_chart(ch, celda_grafico(ws.title, enc, COL_GRAF))


def hoja_f5_superficies(wb, datos):
    ws = wb.create_sheet("F5 Superficies de respuesta")
    f = titulo(ws, 1, "F5 Superficies de respuesta de eta (mallas 7 x 7 de pares)", ancho=10)
    f = condiciones(ws, f, [
        "Base: mallas de pares de los estudios 2 (C1 y C2) y 11 (C3 y C4), contadas por pertenencia a la malla.",
        "Celdas: eta (%) donde la clase es KALINA (escala de color del minimo al maximo de la malla); C = CORREGIBLE, "
        "NC = NO_CONVERGIO, I = INVIABLE (sin eta).",
        "Ajuste cuadratico eta ~ a + b + a^2 + b^2 + ab sobre las celdas KALINA, con variables estandarizadas; R^2 y RMSE en esas celdas.",
        "El maximo del ajuste se evalua en toda la malla: si cae fuera de la region KALINA, el optimo con restricciones esta en la frontera.",
        "Fuente: CSV de Sensibilidad 1 (x1_cargar)."])
    ajustes = []
    for par, (a, ua, c, uc) in PARES.items():
        f = subtitulo(ws, f, "%s: %s x %s" % (par, a, c))
        for b in BLOQUES:
            eta, cls, av, bv = malla_clases(b, par)
            f = subtitulo(ws, f, "%s - %s x %s" % (CORTO[b], a, c))
            f = _matriz(ws, f, eta, cls, av, bv, a, ua, c, uc) + 1
            r = ajuste_cuadratico(eta, av, bv)
            if r:
                ajustes.append([par, CORTO[b], r["n"], num(r["r2"], 3), num(r["r2_aj"], 3), num(r["rmse"], 2),
                                "%s / %s / %s" % (r["max_malla"][0], r["max_malla"][1], num(r["max_malla"][2], 2)),
                                "%s / %s / %s" % (r["max_ajuste"][0], r["max_ajuste"][1], num(r["max_ajuste"][2], 2))])
            larga = _largo(eta, av, bv)
            enc, _, _ = tabla(ws, f, ["Nivel de %s" % c, "%s (%s)" % (a, ua), "eta (%)"], larga,
                              anchos=[28, 12, 12, 12, 12, 12, 12, 12, 12, 12])
            _grafico(ws, enc, larga, "%s (%s)" % (a, ua), "%s" % c, "eta frente a %s por nivel de %s (%s, %s)"
                     % (a, c, CORTO[b], par))
            f = max(enc + len(larga) + 2, enc + FILAS_GRAF) + 1
    f = subtitulo(ws, f, "Ajuste cuadratico por malla")
    _, f, _ = tabla(ws, f, ["Par", "Bloque", "Celdas KALINA", "R^2", "R^2 ajustado", "RMSE (pp de eta)",
                            "Maximo de la malla KALINA (a / b / eta %)", "Maximo del ajuste en toda la malla (a / b / eta %)"],
                    ajustes, anchos=[28, 12, 12, 12, 12, 14, 34, 38])
    f = nota(ws, f + 0, "B2: a = P_baja (kPa), b = x_b. B3: a = P_alta (kPa), b = P_baja (kPa). El maximo del ajuste que cae "
                        "fuera de la region KALINA es lo que predice el polinomio fuera de las celdas ajustadas: indica la direccion, no un diseno factible.")
    r3 = [x[3] for x in ajustes if x[0] == "B3"]
    descripcion(ws, f + 1, ["Mapas de calor de eta y ajuste cuadratico sobre las mallas P_baja x x_b y P_alta x P_baja.",
                            "En las mallas B3 el ajuste explica casi toda la variacion de eta (R^2 de %.2f a %.2f)." % (min(r3), max(r3))])
    return ws


HOJAS_F5 = [hoja_f5_superficies]
