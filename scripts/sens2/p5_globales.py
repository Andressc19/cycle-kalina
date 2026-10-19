# -*- coding: utf-8 -*-
"""Sens2 - p5_globales: cifras globales del trabajo calculadas de los archivos (una sola fuente para los documentos).

Calcula: hojas y graficos del libro, conclusiones, puntos del motor por etapa, controles independientes de cifras y poder del trazado.
Escribe `resultados/2026-10-17_sens2_P5/p5_globales.csv` (clave, valor, fuente). Los documentos escritos a mano marcan esas cifras como
`<!--g:clave-->valor<!--/g-->` y `docs_sync.py` las mantiene al dia; `tests/test_sens2_docs.py` falla si alguna no coincide.
"""
import glob
import os
import re

import pandas as pd

AQUI = os.path.dirname(os.path.abspath(__file__))
RAIZ = os.path.dirname(os.path.dirname(AQUI))
R = os.path.join(RAIZ, "resultados")
P5 = os.path.join(R, "2026-10-17_sens2_P5")
B = os.path.join(R, "2026-10-18_sens2_B")
LIBRO = os.path.join(RAIZ, "datos", "Kalina_anal_sensible_2.xlsx")


def _filas(ruta):
    return len(pd.read_csv(ruta))


def puntos_motor():
    """Puntos nuevos del motor: U1 (campana de incertidumbre) y etapa B (pisos, reintento, densidades, simetrico y optimo por T_fuente)."""
    u1 = _filas(os.path.join(R, "2026-10-17_sens2_U1", "u1_filas.csv"))
    b = {"pisos_reintento": _filas(os.path.join(B, "s2b_filas.csv")), "densidades": _filas(os.path.join(B, "s2b_dens_filas.csv")),
         "simetrico": _filas(os.path.join(B, "s2b_sim_filas.csv")), "optimo_t": sum(_filas(f) for f in glob.glob(os.path.join(B, "s2b_opt_*.csv")))}
    return u1, b


def globales(con_libro=True):
    g = {}
    u1, b = puntos_motor()
    g["puntos_u1"] = (u1, "u1_filas.csv")
    for k, v in b.items():
        g["puntos_b_" + k] = (v, "s2b_*.csv")
    g["puntos_b_total"] = (sum(b.values()), "s2b_*.csv")
    g["n_conclusiones"] = (_filas(os.path.join(P5, "p5_conclusiones.csv")), "p5_conclusiones.csv")
    cc = _filas(os.path.join(P5, "p5_cifras_check.csv"))
    cb = _filas(os.path.join(P5, "p5_cifras_check_c.csv"))
    g["controles_cifras"], g["controles_bh"], g["controles_total"] = (cc, "p5_cifras_check.csv"), (cb, "p5_cifras_check_c.csv"), (cc + cb, "p5_cifras_check*.csv")
    with open(os.path.join(P5, "p5_trazado_fuente.md"), encoding="utf-8") as fh:
        m = re.search(r"se habria trazado en el ([0-9.]+) %", fh.read())
    g["poder_trazado_pct"] = (float(m.group(1)) if m else float("nan"), "p5_trazado_fuente.md")
    if con_libro and os.path.exists(LIBRO):
        from openpyxl import load_workbook
        wb = load_workbook(LIBRO)
        g["n_hojas"] = (len(wb.sheetnames), "Kalina_anal_sensible_2.xlsx")
        g["n_graficos"] = (sum(len(ws._charts) for ws in wb), "Kalina_anal_sensible_2.xlsx")
    return g


def valor_texto(clave, v):
    """Texto con que aparece la cifra en un documento."""
    return ("%.1f" % v) if clave == "poder_trazado_pct" else "%d" % v


def main():
    g = globales()
    os.makedirs(P5, exist_ok=True)
    pd.DataFrame([dict(clave=k, valor=v[0], fuente=v[1]) for k, v in g.items()]).to_csv(os.path.join(P5, "p5_globales.csv"), index=False, encoding="utf-8", lineterminator="\r\n")
    print("globales: " + "; ".join("%s=%s" % (k, valor_texto(k, v[0])) for k, v in g.items()))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
