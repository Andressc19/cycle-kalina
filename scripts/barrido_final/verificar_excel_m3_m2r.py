"""Verifica las hojas M3_/M2R_ de `datos/Excel_final_kalina.xlsx` (TASK_CONTEXT
de excel_m3_m2r): nombres de hoja, filas de datos frente a los CSV de origen,
la columna `fase` de las hojas de nivel, las filas de `Resumen` y la huella de
todo el libro (para comparar entre dos ejecuciones y comprobar idempotencia).
Imprime cada comprobación y termina con OK si todas pasan.
"""
from __future__ import annotations

import hashlib
import sys
from pathlib import Path

_DIR = Path(__file__).resolve().parent
_RAIZ = _DIR.parents[1]
for _p in (str(_RAIZ), str(_DIR), str(_DIR.parent)):   # parent = scripts/
    if _p not in sys.path:
        sys.path.insert(0, _p)

import pandas as pd  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

from formatear_excel_barridos import COLOR_CLASE  # noqa: E402

XLSX = _RAIZ / "datos" / "Excel_final_kalina.xlsx"
M3 = _RAIZ / "resultados" / "2026-10-06_final_M3"
M2R = _RAIZ / "resultados" / "2026-10-06_final_M2R"
# hoja -> [csv de origen] (mismo mapa que excel_m3_m2r.py)
HOJAS = {
    "M3_ofat": [M3 / "m3_ofat.csv"],
    "M3_resumen": [M3 / "m3_resumen.csv"],
    "M3_cruces": [M3 / "m3_cruces.csv"],
    "M2R_n1": [M2R / "m2r_n1_malla.csv", M2R / "m2r_n1_biseccion.csv"],
    "M2R_n2a": [M2R / "m2r_n2a_malla.csv", M2R / "m2r_n2a_biseccion.csv"],
    "M2R_n2b": [M2R / "m2r_n2b_malla.csv", M2R / "m2r_n2b_biseccion.csv"],
    "M2R_n3": [M2R / "m2r_n3_malla.csv"],
    "M2R_fronteras": [M2R / "m2r_fronteras.csv"],
    "M2R_kalina": [M2R / "m2r_kalina.csv"],
    "M2R_resumen": [M2R / "m2r_resumen.csv"],
}
FIGURA = "M3_figuras"
RUTA_TORNADO = (M3 / "m3_tornado.png").relative_to(_RAIZ).as_posix()
HOJAS_ESPERADAS = ["M0_protocolo", "M1_profesor_P_alta", "M1_profesor_x_b",
                   "M1_profesor_T_amb", "M1_profesor_base", "M1_comparacion",
                   "Resumen", "M2_N1_malla", "M2_N1_biseccion",
                   "M2_N1_fronteras", "M2_N1_resumen", "M2_N2_malla",
                   "M2_N2_biseccion", "M2_N2_fronteras", "M2_N2_kalina",
                   "M2_N2_resumen"] + list(HOJAS) + [FIGURA]
NIVELES = ("M2R_n1", "M2R_n2a", "M2R_n2b", "M2R_n3")


def filas_datos(ws) -> int:
    """Filas de la tabla: se corta en la primera vacía o en la descripción."""
    n = 0
    for r in range(2, ws.max_row + 1):
        v = ws.cell(r, 1).value
        if v in (None, "") or str(v).startswith("Descripción"):
            break
        n += 1
    return n


def resumen(ws, hoja: str):
    """[puntos, KALINA, CORREGIBLE, INVIABLE, NO_CONVERGIO, errores] de `hoja`."""
    for r in range(2, ws.max_row + 1):
        v = ws.cell(r, 1).value
        if v in (None, "") or str(v).startswith("Descripción"):
            return None
        if v == hoja:
            return [ws.cell(r, c).value for c in range(2, 8)]
    return None


def huella(wb) -> str:
    """sha256 de títulos y valores de celda de todo el libro (sin estilos)."""
    h = hashlib.sha256()
    for ws in wb.worksheets:
        h.update(ws.title.encode("utf-8"))
        for fila in ws.iter_rows(values_only=True):
            h.update(repr(fila).encode("utf-8"))
    return h.hexdigest()[:16]


def main() -> int:
    if not XLSX.exists():
        print(f"FALLO: no existe {XLSX}")
        return 1
    wb = load_workbook(XLSX)
    print("hojas:", wb.sheetnames)
    fallos = 0

    def check(ok: bool, txt: str) -> None:
        nonlocal fallos
        print(("  OK   " if ok else "  FALLO ") + txt)
        fallos += 0 if ok else 1

    faltan = [h for h in HOJAS_ESPERADAS if h not in wb.sheetnames]
    extra = [h for h in wb.sheetnames if h not in HOJAS_ESPERADAS]
    check(not faltan, f"nombres de hoja (faltan: {faltan or 'ninguno'})")
    check(not extra, f"sin hojas inesperadas (sobran: {extra or 'ninguna'})")
    filas_resumen = filas_datos(wb["Resumen"]) if "Resumen" in wb.sheetnames else -1

    for hoja, csvs in HOJAS.items():
        if hoja not in wb.sheetnames:
            check(False, f"{hoja}: no existe")
            continue
        ws = wb[hoja]
        esp = sum(len(pd.read_csv(c)) for c in csvs)
        n = filas_datos(ws)
        check(n == esp, f"{hoja}: {n} filas de datos = {esp} de los CSV "
                        f"({' + '.join(c.name for c in csvs)})")
        if hoja in NIVELES:
            check(str(ws.cell(1, 1).value).startswith("fase"),
                  f"{hoja}: la columna A de la cabecera es «{ws.cell(1, 1).value}»")
        fila = resumen(wb["Resumen"], hoja) if "Resumen" in wb.sheetnames else None
        check(fila is not None and fila[0] == n,
              f"{hoja}: fila en Resumen con puntos = {n} (fila: {fila})")
        cab = [str(ws.cell(1, c).value) for c in range(1, ws.max_column + 1)
               if ws.cell(1, c).value is not None]
        check(all(h.endswith(")") for h in cab),
              f"{hoja}: {len(cab)} cabeceras con unidad")
        check(bool(ws.auto_filter.ref), f"{hoja}: auto_filtro en {ws.auto_filter.ref}")
        desc = any(str(ws.cell(r, 1).value or "").startswith("Descripción")
                   for r in range(n + 2, ws.max_row + 1))
        check(desc, f"{hoja}: descripción al pie (2 filas vacías + texto)")
        jc = next((c for c in range(1, len(cab) + 1)
                   if cab[c - 1].startswith("clasificacion")), 0)
        if jc:   # color de clase del formateador en la primera fila de datos
            cel = ws.cell(2, jc)
            ok = (cel.value in COLOR_CLASE and cel.fill.patternType == "solid"
                  and str(cel.fill.fgColor.rgb or "").endswith(COLOR_CLASE[cel.value]))
            check(ok, f"{hoja}: color de clase en la fila 2 ({cel.value})")

    if FIGURA in wb.sheetnames:
        wsf = wb[FIGURA]
        check(wsf.cell(2, 2).value == RUTA_TORNADO,
              f"{FIGURA}: ruta = {wsf.cell(2, 2).value}")
        print(f"  info  {FIGURA}: imagen(es) insertada(s) = {len(wsf._images)} "
              f"(0 = solo la ruta, por falta de Pillow)")
    else:
        check(False, f"{FIGURA}: no existe")

    print(f"  info  filas de datos de Resumen = {filas_resumen}")
    print(f"huella del libro: {huella(wb)}")
    print("OK" if not fallos else f"FALLO ({fallos} comprobaciones)")
    return 0 if not fallos else 1


if __name__ == "__main__":
    sys.exit(main())
