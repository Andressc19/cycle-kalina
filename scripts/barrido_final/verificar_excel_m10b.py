"""M10b — verifica el Excel final tras `excel_m10b.py`.

Comprueba que las hojas M10b existen con el número de filas de datos de su CSV,
que las 78 hojas anteriores (M0...M10 incluidas) siguen presentes y que
`Resumen` tiene las filas M10b. Imprime OK/FALLA y sale con código 1 si falla.
"""
from __future__ import annotations

import csv  # noqa: E402
import sys  # noqa: E402
from pathlib import Path

_DIR = Path(__file__).resolve().parent
_RAIZ = _DIR.parents[1]
sys.path[:0] = [str(_RAIZ), str(_DIR)]

from openpyxl import load_workbook  # noqa: E402

import excel_final as ef  # noqa: E402
import excel_m10 as e10  # noqa: E402
import excel_m10b as e10b  # noqa: E402
import verificar_excel_m10 as v10  # noqa: E402

# 78 hojas presentes después de M10: M0...M9 + las 6 de M10.
PREVIAS = (set(v10.HOJAS_ANTES) | {h for h, _c, _d in e10.HOJAS}
           | {e10.FIGURA})
N_PREVIAS = 78
_falla = 0


def chk(nombre, ok, detalle="") -> None:
    global _falla
    if not ok:
        _falla += 1
    print(f'{"OK  " if ok else "FALLA"} {nombre}{" — " + detalle if detalle else ""}',
          flush=True)


def _n_csv(p: Path) -> int:
    with open(p, encoding="utf-8", newline="") as fh:
        return sum(1 for _ in csv.reader(fh)) - 1


def _n_datos(ws) -> int:
    """Filas de datos: desde la 2 hasta la primera celda A vacía o la descripción."""
    n = 0
    for r in range(2, ws.max_row + 1):
        v = ws.cell(r, 1).value
        if v in (None, "") or str(v).startswith("Descripción"):
            break
        n += 1
    return n


def main() -> None:
    if not ef.XLSX.exists():
        raise SystemExit(f"No existe {ef.XLSX}")
    wb = load_workbook(ef.XLSX, read_only=False)
    chk(f"conjunto de hojas previas = {N_PREVIAS}", len(PREVIAS) == N_PREVIAS,
        str(len(PREVIAS)))
    for h in sorted(PREVIAS):
        chk(f"hoja previa {h}", h in wb.sheetnames)
    for hoja, csvf, _d in e10b.HOJAS:
        if hoja not in wb.sheetnames:
            chk(f"hoja {hoja} existe", False)
            continue
        n = _n_datos(wb[hoja])
        ncsv = _n_csv(e10b.OUT / csvf)
        chk(f"{hoja}: filas", n == ncsv, f"{n} vs {ncsv}")
    chk(f"hoja {e10b.FIGURA} existe", e10b.FIGURA in wb.sheetnames)
    chk("imágenes M10b insertadas",
        e10b.FIGURA in wb.sheetnames
        and len(getattr(wb[e10b.FIGURA], "_images", [])) > 0)
    if "Resumen" in wb.sheetnames:
        ws = wb["Resumen"]
        m10b = [r for r in range(2, ws.max_row + 1)
                if str(ws.cell(r, 1).value or "").startswith("M10b_")]
        chk("Resumen tiene filas M10b", len(m10b) > 0, f"{len(m10b)} filas")
    print(f"\n{'TODO OK' if _falla == 0 else str(_falla) + ' FALLAS'}", flush=True)
    sys.exit(1 if _falla else 0)


if __name__ == "__main__":
    main()
