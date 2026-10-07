"""M2S/M4 — hojas de las campañas M2S (sumidero frío con P_baja libre) y M4/M4b
(LHS de 500 puntos + análisis de sensibilidad) en `datos/Excel_final_kalina.xlsx`.

Abre el libro que ya escribieron `excel_final.py`, `excel_m2.py` y
`excel_m3_m2r.py`, crea o reemplaza las hojas M2S_* y M4_* con los CSV de
`resultados/2026-10-06_final_M2S/`, `2026-10-06_final_M4/` y
`2026-10-07_final_M4b/` vía `excel_final.escribir_tabla` (mismo formato,
unidades, colores de clase, filtro y descripción al pie), añade o actualiza sus
filas en `Resumen` y monta `M4_figuras` con las tres PNG de M4b. `M4_lhs500`
junta los 300 de M4 con los 200 de M4b y la columna `lote` dice de cuál viene
cada fila. Es idempotente y se corre SIEMPRE el último de los cuatro.
"""
from __future__ import annotations

import sys
from pathlib import Path

_DIR = Path(__file__).resolve().parent
_RAIZ = _DIR.parents[1]
for _p in (str(_RAIZ), str(_DIR), str(_DIR.parent)):   # parent = scripts/
    if _p not in sys.path:
        sys.path.insert(0, _p)

import pandas as pd  # noqa: E402
from openpyxl import load_workbook  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

import excel_final as ef  # noqa: E402
from formatear_excel_barridos import estilo_titulo  # noqa: E402

M2S = _RAIZ / "resultados" / "2026-10-06_final_M2S"
M4 = _RAIZ / "resultados" / "2026-10-06_final_M4"
M4B = _RAIZ / "resultados" / "2026-10-07_final_M4b"
XLSX = ef.XLSX
FIGURA = "M4_figuras"
FIGS = ("m4b_plano_presiones.png", "m4b_hist_dP_rP.png", "m4b_rsa_ks.png")
KALINA_TODA = ("M2S_kalina", "M4_top_kalina")   # todas sus filas son KALINA
# hoja -> (rutas de sus CSV, campaña de la que salen sus filas, descripción al pie)
HOJAS = (
    ("M2S_malla", [M2S / "m2s_malla.csv"], "M2S", "M2S · malla: P_baja de 150 a 1500 kPa en 10 valores, para 4 bases × 5 climas de T_sumidero (t_amb_diseno = T_sumidero); 200 corridas."),
    ("M2S_biseccion", [M2S / "m2s_biseccion.csv"], "M2S", "M2S · bisección sobre P_baja (tolerancia 10 kPa, máximo 7 iteraciones) donde cambia «es KALINA» o el signo de m_O2 con ambos vecinos genuinos; `frontera` dice el criterio acotado y `iteracion` la ronda."),
    ("M2S_resumen", [M2S / "m2s_resumen.csv"], "M2S", "M2S: una fila por (base, T_sumidero); `n_genuinos`, `n_kalina` y `n_N2` suman malla + bisección, y `P_baja_min_kalina` con `eta_min_kalina` son el punto KALINA de menor P_baja de la fila."),
    ("M2S_kalina", [M2S / "m2s_kalina.csv"], "M2S", "M2S: los puntos KALINA sin duplicados de malla y bisección (123 filas); `B_verificado` y `B_dh4s` son la verificación B del motor real."),
    ("M4_lhs500", [M4 / "m4_lhs.csv", M4B / "m4b_lhs.csv"], "M4", "M4+M4b: los 500 puntos LHS unidos — 300 de M4 (semilla 2026-10-06) y 200 de M4b (semilla 2026-10-07); `lote` dice de cuál viene cada fila. La unión de dos LHS no es un LHS estricto de 500 (REPORTE_M4b §2)."),
    ("M4_rsa", [M4B / "m4b_rsa.csv"], "M4", "M4b · RSA sobre los 500 puntos: una fila por (separación, variable); `KS` mide cuánto separa KALINA del resto, `ic_lo`/`ic_hi` es el IC95 % por bootstrap y `fav_p05`/`fav_p95` el rango favorable (percentiles 5–95 dentro de KALINA)."),
    ("M4_prcc", [M4B / "m4b_prcc.csv"], "M4", "M4b · PRCC por (salida, variable) con la prueba t, su valor p y el IC95 %; salidas: eta, m_O1, m_O2 y m_O5."),
    ("M4_pawn", [M4B / "m4b_pawn.csv"], "M4", "M4b · PAWN por (salida, variable): KS de cada una de las 10 franjas de T, `franjas_vacias`, `T_mediana` y `T_max`; salidas: eta y D."),
    ("M4_presiones_rangos", [M4B / "m4b_rangos_presiones.csv"], "M4", "M4b: min, p05, mediana, p95 y max de P_alta, P_baja, dP y rP dentro de KALINA y en los genuinos que no lo son, con el KS y su valor p que separa los dos grupos."),
    ("M4_ventana_Palta", [M4B / "m4b_ventana_por_palta.csv"], "M4", "M4b: 6 franjas de P_alta de 2000 a 5000 kPa con `n`, `n_kalina` y `pct_kalina`, y los rangos (min–max) de P_baja, dP y rP solo de los KALINA de la franja."),
    ("M4_ventana_Tfuente", [M4B / "m4b_ventana_por_tfuente.csv"], "M4", "M4b: 6 franjas de T_fuente de 380 a 500 K; mismo contenido que M4_ventana_Palta."),
    ("M4_ventana_Tsumidero", [M4B / "m4b_ventana_por_tsumidero.csv"], "M4", "M4b: 4 franjas de T_sumidero de 278.15 a 303.55 K; mismo contenido que M4_ventana_Palta."),
    ("M4_bloqueo_rP", [M4B / "m4b_bloqueo_por_rP.csv"], "M4", "M4b: puntos genuinos por zona de rP — por debajo del p5 de los KALINA, dentro de [p5, p95] y por encima del p95 — y por código de falla."),
    ("M4_top_kalina", [M4B / "m4b_top_kalina.csv"], "M4", "M4b: los 10 puntos KALINA de mayor η de los 500, con sus 10 variables de la caja."),
    ("M4_convergencia", [M4B / "m4b_convergencia.csv"], "M4", "M4b: 300 vs 500 puntos, métrica por métrica y variable por variable; `rho_spearman` compara los dos rankings y `cumple_09` marca ρ ≥ 0.9."),
)
# Unidades nuevas del TASK. `lo`/`hi` las pone `_lo_hi` según la variable de la hoja.
UNIDADES = {"dP": "kPa", "dP_min": "kPa", "dP_max": "kPa", "rP": "\u2212",
            "rP_min": "\u2212", "rP_max": "\u2212", "dT_fs": "K", "KS": "\u2212",
            "p_valor": "\u2212", "prcc": "\u2212", "n_kalina": "\u2212",
            "pct_kalina": "%"}


def _leer(rutas: list) -> pd.DataFrame:
    """CSV de una hoja; con más de uno los junta con la columna `lote` (M4/M4b)."""
    dfs = [pd.read_csv(p) for p in rutas]
    if len(dfs) == 1:
        return dfs[0]
    if list(dfs[0].columns) != list(dfs[1].columns):
        raise SystemExit(f"Columnas distintas entre {[p.name for p in rutas]}")
    juntos = [d.assign(lote="M4b" if "M4b" in p.parent.name else "M4")
              for p, d in zip(rutas, dfs)]
    return pd.concat(juntos, ignore_index=True)[["lote"] + list(dfs[0].columns)]


def _campanas() -> dict:
    """Puntos de cada campaña: malla + bisección en M2S, los 500 LHS en M4."""
    return {"M2S": pd.concat([pd.read_csv(M2S / "m2s_malla.csv"),
                              pd.read_csv(M2S / "m2s_biseccion.csv")],
                             ignore_index=True),
            "M4": pd.concat([pd.read_csv(M4 / "m4_lhs.csv"),
                             pd.read_csv(M4B / "m4b_lhs.csv")],
                            ignore_index=True)}


def _lo_hi(df) -> None:
    """`lo`/`hi` toman la unidad de la `variable` de la hoja; sin una sola, «−»."""
    u = "\u2212"
    vars_ = set(df["variable"].dropna().unique())
    if len(vars_) == 1:
        u = ef.unidad(str(vars_.pop())) or "\u2212"
    ef.UNIDAD["lo"] = ef.UNIDAD["hi"] = u


def _fila_resumen(hoja: str, df, camp) -> list:
    """[hoja, filas, KALINA, CORREGIBLE, INVIABLE, NO_CONVERGIO, errores]."""
    if "clasificacion" in df:
        n = df["clasificacion"].value_counts()
        err = int(df["error_inesperado"].notna().sum())
    elif hoja in KALINA_TODA:      # como M2R_kalina: todas sus filas son KALINA
        n, err = {"KALINA": len(df)}, 0
    else:                          # tabla derivada: las clases son las de su campaña
        n = camp["clasificacion"].value_counts()
        err = int(camp["error_inesperado"].notna().sum())
    return [hoja, len(df)] + [int(n.get(c, 0)) for c in ef.CLASES] + [err]


def _ultima_fila(ws) -> int:
    """Última fila de datos de `Resumen` (deja fuera las 2 vacías y la descripción)."""
    ult = 1
    for r in range(2, ws.max_row + 1):
        v = ws.cell(r, 1).value
        if v in (None, "") or str(v).startswith("Descripción"):
            break
        ult = r
    return ult


def _reescribir_resumen(ws, filas: list) -> None:
    """Quita las filas M2S_/M4_ viejas de `Resumen` y reinserta las nuevas."""
    for r in range(ws.max_row, 1, -1):
        if str(ws.cell(r, 1).value or "").startswith(("M2S_", "M4_")):
            ws.delete_rows(r)
    pos = _ultima_fila(ws) + 1
    ws.insert_rows(pos, len(filas))
    for k, fila in enumerate(filas):
        for j, v in enumerate(fila, 1):
            ws.cell(pos + k, j, v)
    j = 1
    while ws.cell(1, j + 1).value is not None:
        j += 1
    ws.auto_filter.ref = f"A1:{get_column_letter(j)}{pos + len(filas) - 1}"


def _figuras(wb) -> str:
    """Hoja `M4_figuras` con las tres PNG de M4b; sin Pillow o sin PNG, solo rutas."""
    ws = wb.create_sheet(FIGURA)
    for j, t in enumerate(("figura", "ruta"), 1):
        estilo_titulo(ws.cell(1, j, t), "texto")
    rutas = [M4B / n for n in FIGS]
    for k, p in enumerate(rutas, 2):
        ws.cell(k, 1, p.name)
        ws.cell(k, 2, p.relative_to(_RAIZ).as_posix())
    c = len(rutas) + 3
    ws.cell(c, 1, "Figuras de M4b: (a) plano P_baja contra P_alta con la clase de cada punto y las líneas de rP constante; (b) histograma de dP y rP dentro y fuera de KALINA; (c) KS del RSA por variable.")
    ws.freeze_panes = "A2"
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 64
    faltan = [p.name for p in rutas if not p.exists()]
    try:
        from openpyxl.drawing.image import Image as XLImage
    except ImportError:
        ws.cell(c + 2, 1, "Pillow no disponible: no se insertó ninguna imagen; las rutas están arriba.")
        return f"{FIGURA}: sin Pillow, solo las rutas"
    if faltan:
        ws.cell(c + 2, 1, "Faltan: " + ", ".join(faltan))
        return f"{FIGURA}: faltan {', '.join(faltan)}"
    fila = c + 2
    for p in rutas:
        img = XLImage(str(p))
        ws.add_image(img, f"A{fila}")
        fila += -(-img.height // 20) + 1   # filas que ocupa la imagen (20 px) + 1 de aire
    return f"{len(rutas)} imágenes insertadas en {FIGURA}"


def main() -> None:
    faltan = [p for _h, rutas, _c, _d in HOJAS for p in rutas if not p.exists()]
    if not XLSX.exists():
        raise SystemExit(f"No existe {XLSX}: correr antes scripts/barrido_final/excel_final.py")
    if faltan:
        raise SystemExit("Faltan CSV: " + ", ".join(str(p) for p in faltan))
    ef.UNIDAD.update(UNIDADES)
    wb = load_workbook(XLSX)
    if "Resumen" not in wb.sheetnames:
        raise SystemExit("El libro no tiene la hoja Resumen: correr antes excel_final.py")
    viejas = [h for h, _r, _c, _d in HOJAS if h in wb.sheetnames]
    for h in viejas + ([FIGURA] if FIGURA in wb.sheetnames else []):
        del wb[h]
    camp = _campanas()
    filas = []
    for hoja, rutas, clave, desc in HOJAS:
        df = _leer(rutas)
        if {"lo", "hi"} <= set(df.columns):
            _lo_hi(df)
        ef.escribir_tabla(wb, hoja, list(df.columns), df.values.tolist(), extra=[desc])
        filas.append(_fila_resumen(hoja, df, camp[clave]))
    msg = _figuras(wb)
    _reescribir_resumen(wb["Resumen"], filas)
    wb.save(XLSX)
    print(f"escrito {XLSX}")
    print("hojas:", wb.sheetnames)
    print(msg)


if __name__ == "__main__":
    main()
