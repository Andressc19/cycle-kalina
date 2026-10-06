"""M2 — hojas de nivel 1 y nivel 2 en `datos/Excel_final_kalina.xlsx` (TASK_CONTEXT).

Abre el libro que ya escribió `excel_final.py`, crea o reemplaza las hojas
`M2_N1_malla`, `M2_N1_biseccion`, `M2_N1_fronteras`, `M2_N1_resumen`,
`M2_N2_malla`, `M2_N2_biseccion`, `M2_N2_fronteras`, `M2_N2_kalina` y
`M2_N2_resumen` a partir de los CSV de `resultados/2026-10-06_final_M2/`,
reutilizando `excel_final.escribir_tabla` (mismo formato y descripción), y añade
o actualiza las filas de esas hojas en `Resumen`. Es idempotente: al correrlo dos
veces el libro queda igual. Se corre SIEMPRE después de `excel_final.py`, que
reconstruye el libro desde cero y por eso borra estas hojas.
"""
from __future__ import annotations

import sys
from pathlib import Path

_DIR = Path(__file__).resolve().parent
_RAIZ = _DIR.parents[1]
for _p in (str(_RAIZ), str(_DIR), str(_DIR.parent)):   # parent = scripts/
    if _p not in sys.path:
        sys.path.insert(0, _p)

import pandas as pd  # noqa: E402
from openpyxl import load_workbook  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

import excel_final as ef  # noqa: E402

CAMPA = _RAIZ / "resultados" / "2026-10-06_final_M2"
XLSX = ef.XLSX
# hoja -> (csv de origen, descripción que va al pie de la tabla)
HOJAS = (
    ("M2_N1_malla", "m2n1_malla.csv",
     "M2 nivel 1 · Fase A: una sola variable cambia por fila sobre el punto del "
     "profesor (10 valores equiespaciados por variable; las demás, fijas)."),
    ("M2_N1_biseccion", "m2n1_biseccion.csv",
     "M2 nivel 1 · Fase B: puntos medios de la bisección; `iteracion` es la ronda "
     "y `frontera` el criterio que se acota (convergencia o KALINA)."),
    ("M2_N1_fronteras", "m2n1_fronteras.csv",
     "M2 nivel 1: una fila por frontera; `lado_si` es el extremo que cumple el "
     "criterio, `lado_no` el que no, y `ancho_final` = hi − lo tras la bisección."),
    ("M2_N1_resumen", "m2n1_resumen.csv",
     "M2 nivel 1: una fila por variable; los conteos suman malla + bisección y "
     "`d_normalizada` = |v − v_prof|/(b − a) del valor KALINA más cercano."),
    ("M2_N2_malla", "m2n2_malla.csv",
     "M2 nivel 2 · Fase A: dos variables cambian por fila sobre el punto del "
     "profesor —`P_baja` y una segunda (eps_hrvg en la ruta A, T_fuente en la "
     "B)— con las otras nueve fijas; 12 valores de P_baja por valor de la variable 1."),
    ("M2_N2_biseccion", "m2n2_biseccion.csv",
     "M2 nivel 2 · Fase B: puntos medios de la bisección sobre P_baja; "
     "`frontera` es el criterio que se acota (KALINA u O2) y `nota` anota los "
     "puntos medios que no son genuinos en una frontera O2."),
    ("M2_N2_fronteras", "m2n2_fronteras.csv",
     "M2 nivel 2: una fila por frontera de P_baja; `P_baja_si` es el extremo que "
     "cumple el criterio, `P_baja_no` el que no y `ancho_final` = hi − lo tras la bisección."),
    ("M2_N2_kalina", "m2n2_kalina.csv",
     "M2 nivel 2: todas las filas KALINA de malla y bisección; `d` = "
     "|val1 − val1_prof|/(b1 − a1) + |P_baja − 400|/(1500 − 150) con la escala del nivel 1."),
    ("M2_N2_resumen", "m2n2_resumen.csv",
     "M2 nivel 2: una fila por (ruta, val1); los conteos suman malla + "
     "bisección y solo cuentan como genuinos KALINA, CORREGIBLE e INVIABLE."),
)
# Unidades nuevas (TASK §4). `valor`, `ancho_final`, `lado_si` y `lado_no`
# mezclan kPa/K/(−) según la variable y en una sola columna no pueden variar,
# por eso van sin unidad: la columna `variable` de la misma fila la dice.
UNIDADES = {"valor": "\u2212", "ancho_final": "\u2212", "lado_si": "\u2212",
            "lado_no": "\u2212", "d_normalizada": "\u2212", "val1": "\u2212",
            "d": "\u2212", "mejor_D": "\u2212", "iteracion": "n",
            "iteraciones": "n", "n_genuinos": "n", "n_kalina": "n",
            "n_espurios_N2": "n", "P_baja_val": "kPa", "P_baja_si": "kPa",
            "P_baja_no": "kPa", "P_baja_min_kalina": "kPa",
            "P_baja_min_O2": "kPa", "P_baja_mejor_D": "kPa"}


def _fila_resumen(hoja: str, df, todo=None) -> list:
    """Fila de `Resumen`: [hoja, filas, KALINA, CORREGIBLE, INVIABLE,
    NO_CONVERGIO, errores_inesperados], con el conteo por clase de esa hoja."""
    if "clasificacion" in df:
        n = df["clasificacion"].value_counts()
        err = int(df["error_inesperado"].notna().sum())
    elif "clase_lado_si" in df:
        n, err = df["clase_lado_si"].value_counts(), 0
    elif "clase_si" in df:      # M2_N2_fronteras
        n, err = df["clase_si"].value_counts(), 0
    elif "n_espurios_N2" in df:  # M2_N2_resumen: los conteos son de malla + bisección
        if todo is None:
            raise SystemExit("M2_N2_resumen necesita m2n2_malla.csv y m2n2_biseccion.csv")
        n = todo["clasificacion"].value_counts()
        err = int(todo["error_inesperado"].notna().sum())
    elif "d" in df:             # M2_N2_kalina: todas sus filas son KALINA
        n, err = {"KALINA": len(df)}, 0
    else:   # M2_N1_resumen: sin columna de clase propia; suma sus n_* (malla+bisección)
        n = {c: int(df["n_" + c.lower()].fillna(0).sum()) for c in ef.CLASES}
        err = 0
    return [hoja, len(df)] + [int(n.get(c, 0)) for c in ef.CLASES] + [err]


def _ultima_fila(ws) -> int:
    """Última fila de datos de `Resumen` (deja fuera las 2 vacías y la descripción)."""
    ult = 1
    for r in range(2, ws.max_row + 1):
        v = ws.cell(r, 1).value
        if v in (None, "") or str(v).startswith("Descripción"):
            break
        ult = r
    return ult


def _reescribir_resumen(ws, filas: list) -> None:
    """Actualiza las filas M2_N1_/M2_N2_ de `Resumen`: quita las viejas y reinserta."""
    for r in range(ws.max_row, 1, -1):
        if str(ws.cell(r, 1).value or "").startswith(("M2_N1_", "M2_N2_")):
            ws.delete_rows(r)
    pos = _ultima_fila(ws) + 1
    ws.insert_rows(pos, len(filas))
    for k, fila in enumerate(filas):
        for j, v in enumerate(fila, 1):
            ws.cell(pos + k, j, v)
    j = 1
    while ws.cell(1, j + 1).value is not None:
        j += 1
    ws.auto_filter.ref = f"A1:{get_column_letter(j)}{pos + len(filas) - 1}"


def main() -> None:
    faltan = [c for _h, c, _e in HOJAS if not (CAMPA / c).exists()]
    if not XLSX.exists():
        raise SystemExit(f"No existe {XLSX}: correr antes scripts/barrido_final/excel_final.py")
    if faltan:
        raise SystemExit(f"Faltan en {CAMPA}: {', '.join(faltan)} (correr m2_nivel1.py)")
    ef.UNIDAD.update(UNIDADES)
    wb = load_workbook(XLSX)
    if "Resumen" not in wb.sheetnames:
        raise SystemExit("El libro no tiene la hoja Resumen: correr antes excel_final.py")
    for hoja in [h for h, _c, _e in HOJAS if h in wb.sheetnames]:
        del wb[hoja]
    # filas de malla + bisección del nivel 2, para los conteos de M2_N2_resumen
    todo = pd.concat([pd.read_csv(CAMPA / "m2n2_malla.csv"),
                      pd.read_csv(CAMPA / "m2n2_biseccion.csv")], ignore_index=True)
    filas_resumen = []
    for hoja, csv, extra in HOJAS:
        df = pd.read_csv(CAMPA / csv)
        ef.escribir_tabla(wb, hoja, list(df.columns), df.values.tolist(), extra=[extra])
        filas_resumen.append(_fila_resumen(hoja, df, todo))
    _reescribir_resumen(wb["Resumen"], filas_resumen)
    wb.save(XLSX)
    print(f"escrito {XLSX}")
    print("hojas:", wb.sheetnames)


if __name__ == "__main__":
    main()
