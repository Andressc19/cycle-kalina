"""M6 — verificación de las hojas M6 en `datos/Excel_final_kalina.xlsx` (§C5).

Comprueba que las 5 hojas M6 existen con tantas filas de datos como su CSV,
que las 59 hojas anteriores a M6 siguen presentes, y que `Resumen` tiene una
fila por hoja M6 con el conteo de filas correcto. Imprime OK/FALLA por cada
comprobación y sale con 1 si algo falla.
"""
from __future__ import annotations

import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

_DIR = Path(__file__).resolve().parent
_RAIZ = _DIR.parents[1]
sys.path[:0] = [str(_RAIZ), str(_DIR.parent), str(_DIR)]

import excel_final as ef  # noqa: E402

M6 = _RAIZ / "resultados" / "2026-10-08_final_M6"
HOJAS = (("M6_verificacion", "m6_verificacion.csv"),
         ("M6_comparacion", "m6_comparacion.csv"),
         ("M6_piso_30_4C", "m6_piso_30_4C.csv"),
         ("M6_subenfriamiento", "m6_subenfriamiento.csv"),
         ("M6_subenf_evals", "m6_subenfriamiento_evals.csv"))
# Las 59 hojas que el libro tenía antes de M6 (incluida Resumen).
ANTERIORES = (
    "M0_protocolo", "M1_profesor_P_alta", "M1_profesor_x_b", "M1_profesor_T_amb",
    "M1_profesor_base", "M1_comparacion", "Resumen", "M2_N1_malla",
    "M2_N1_biseccion", "M2_N1_fronteras", "M2_N1_resumen", "M2_N2_malla",
    "M2_N2_biseccion", "M2_N2_fronteras", "M2_N2_kalina", "M2_N2_resumen",
    "M3_ofat", "M3_resumen", "M3_cruces", "M2R_n1", "M2R_n2a", "M2R_n2b",
    "M2R_n3", "M2R_fronteras", "M2R_kalina", "M2R_resumen", "M3_figuras",
    "M2S_malla", "M2S_biseccion", "M2S_resumen", "M2S_kalina", "M4_lhs500",
    "M4_rsa", "M4_prcc", "M4_pawn", "M4_presiones_rangos", "M4_ventana_Palta",
    "M4_ventana_Tfuente", "M4_ventana_Tsumidero", "M4_bloqueo_rP",
    "M4_top_kalina", "M4_convergencia", "M4_figuras", "M5_lhs", "M5_resumen",
    "M5_rangos", "M5_cercano_profesor", "M5_escenarios", "M5_figuras",
    "M5c_malla", "M5c_frontera", "M4c_lhs", "M4c_techo", "M4c_techo_Palta",
    "M7_optimos", "M7_trayectorias", "M7_evaluaciones", "M8_ANECDOTICO_oat",
    "M5c_M4c_M7_M8_figuras")
FAL = 0


def chk(ok: bool, msg: str) -> None:
    global FAL
    FAL += 0 if ok else 1
    print(("OK    " if ok else "FALLO ") + msg, flush=True)


def _tabla(ws, n: int, cab: list) -> list:
    """Problemas de estructura de una hoja de `n` filas de datos."""
    mal = []
    for j, c in enumerate(cab, 1):
        h = str(ws.cell(1, j).value or "")
        if not h.startswith(str(c)):
            mal.append(f"cabecera {j} = {h!r} no empieza por {c!r}")
    if ws.cell(n + 1, 1).value is None:
        mal.append(f"la fila {n + 1} (último dato) está vacía")
    try:
        d = next(r for r in range(n + 2, n + 9)
                 if str(ws.cell(r, 1).value or "") == "Descripción")
    except StopIteration:
        return mal + ["no aparece la fila 'Descripción' tras los datos"]
    mal += [f"la fila {r} (entre datos y descripción) no está vacía"
            for r in range(n + 2, d) if ws.cell(r, 1).value is not None]
    return mal


def main() -> None:
    if not ef.XLSX.exists():
        chk(False, f"existe {ef.XLSX}")
        print(f"FIN: {FAL} fallo(s)", flush=True)
        sys.exit(1)
    wb = load_workbook(ef.XLSX)
    nombres = wb.sheetnames
    print(f"verificar_excel_m6 | {ef.XLSX} | {len(nombres)} hojas", flush=True)
    faltan = [h for h, _c in HOJAS if h not in nombres]
    chk(not faltan, f"están las 5 hojas M6 (faltan: {faltan})")
    previas = [h for h in ANTERIORES if h not in nombres]
    chk(len(ANTERIORES) == 59 and not previas,
        f"las 59 hojas anteriores siguen: {59 - len(previas)}/59 (faltan: {previas})")
    chk(len(nombres) == 59 + len(HOJAS),
        f"el libro tiene {59 + len(HOJAS)} hojas (59 + {len(HOJAS)}): {len(nombres)}")
    res = wb["Resumen"] if "Resumen" in nombres else None
    filas_res = {}
    if res is not None:
        for r in range(2, res.max_row + 1):
            v = res.cell(r, 1).value
            if v in (None, "") or str(v).startswith("Descripción"):
                break
            filas_res[str(v)] = (r, res.cell(r, 2).value)
    for hoja, csv in HOJAS:
        if hoja not in nombres:
            continue
        df = pd.read_csv(M6 / csv)
        ws = wb[hoja]
        mal = _tabla(ws, len(df), list(df.columns))
        chk(not mal, f"{hoja}: {len(df)} filas de datos como su CSV → {mal[:3]}")
        if res is not None:
            v = filas_res.get(hoja)
            chk(v is not None and v[1] == len(df),
                f"Resumen: {hoja} con puntos = {len(df)} → {v}")
    m6 = [f for f in filas_res if str(f).startswith("M6_")]
    chk(len(m6) == len(HOJAS), f"Resumen tiene las {len(HOJAS)} filas M6 → {sorted(m6)}")
    print(f"FIN: {FAL} fallo(s)", flush=True)
    sys.exit(1 if FAL else 0)


if __name__ == "__main__":
    main()
