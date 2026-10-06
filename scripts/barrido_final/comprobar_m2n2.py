"""Comprobación final de M2 nivel 2 (TASK_CONTEXT final_M2_N2, Paso 5).
Verifica (a) filas/columnas de cada CSV, (b) que resumen y kalina se recalculan
desde malla + bisección, (c) las fronteras contra los puntos evaluados, (d) las
hojas del xlsx, y (e) que cada número del REPORT sale de los CSV. rc=1 si falla."""
from __future__ import annotations

import re
import sys
from pathlib import Path

_DIR = Path(__file__).resolve().parent
_RAIZ = _DIR.parents[1]
sys.path[:0] = [str(_RAIZ), str(_DIR), str(_DIR.parent)]

import numpy as np  # noqa: E402
import pandas as pd  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

import excel_m2 as em  # noqa: E402
import m2_nivel2 as n2  # noqa: E402
import m2_nivel2_reporte as rep  # noqa: E402

OUT = _RAIZ / "resultados" / "2026-10-06_final_M2"
XLSX = _RAIZ / "datos" / "Excel_final_kalina.xlsx"
HOJAS = ("M2_N2_malla", "M2_N2_biseccion", "M2_N2_fronteras", "M2_N2_kalina", "M2_N2_resumen")
CLAVES = ("malla", "biseccion", "fronteras", "kalina", "resumen")
CLASES = ("KALINA", "CORREGIBLE", "INVIABLE", "NO_CONVERGIO")
FALLOS: list = []


def ck(ok, msg, det="") -> bool:
    print(f"OK   {msg}" if ok else f"FALLO {msg} :: {det}")
    if not ok:
        FALLOS.append(f"{msg} :: {det}")
    return bool(ok)


def s(x) -> str:
    return "" if x is None or (isinstance(x, float) and pd.isna(x)) else str(x)


def ig(a, b, tol=1e-9) -> bool:
    """Compara dos números tolerando None/NaN (de la columna vacía del CSV)."""
    a = None if a is None or (isinstance(a, float) and pd.isna(a)) else float(a)
    b = None if b is None or (isinstance(b, float) and pd.isna(b)) else float(b)
    return (a is None and b is None) or (a is not None and b is not None and abs(a - b) <= tol)


def multiset(cad) -> frozenset:
    """'C3(10) · N2(13)' -> frozenset de pares (conteo, código)."""
    return frozenset((int(n), k) for c in s(cad).split(" · ") if c for k, n in [c[:-1].rsplit("(", 1)])


def csvs() -> dict:
    return {k: pd.read_csv(OUT / f"m2n2_{k}.csv", float_precision="round_trip") for k in CLAVES}


def records(d: dict) -> list:
    """Malla + bisección como las vería el reporte: los NaN del CSV vuelven a None."""
    df = pd.concat([d["malla"], d["biseccion"]], ignore_index=True).astype(object)
    for c in df.columns:
        df[c] = df[c].where(df[c].notna(), None)
    return df.to_dict("records")


def parte_a(d: dict) -> None:
    """(a) columnas y número de filas de cada CSV, más la estructura de la malla."""
    m, b, f, k, r = (d[x] for x in CLAVES)
    n_grupos = sum(len(v[1]) for v in n2.RUTAS.values())
    n_espera = n_grupos * len(n2.P_MALLA)
    ck(list(m.columns) == n2.CMALLA and len(m) == n_espera, f"malla: {n_espera} filas x CMALLA", f"filas={len(m)}")
    ck(list(b.columns) == n2.CBISE and list(f.columns) == rep.CFRON, "columnas de biseccion y fronteras", str(list(b.columns)[:6]))
    ck(list(k.columns) == rep.CKAL and list(r.columns) == rep.CREST, "columnas de kalina y resumen", str(list(r.columns)[:6]))
    ck(len(r) == n_grupos, f"resumen: {n_grupos} filas (ruta x val1)", f"filas={len(r)}")
    ck(set(m.fase) == {"malla"} and set(b.fase) <= {"biseccion"}, "columna fase", f"{set(m.fase)} {set(b.fase)}")
    ck(len(k) == int((m.clasificacion == "KALINA").sum()) + int((b.clasificacion == "KALINA").sum()),
       "kalina: filas == KALINA de malla + biseccion", f"kalina={len(k)}")
    ck(int(f.iteraciones.sum()) == len(b), "fronteras: Σ iteraciones == filas de biseccion", f"Σ={int(f.iteraciones.sum())} bise={len(b)}")
    for ruta, (var1, vals) in n2.RUTAS.items():
        sub = m[m.ruta == ruta]
        ck(set(round(v, 9) for v in sub.val1) == set(vals), f"malla {ruta}: val1 de {var1}", str(sorted(set(sub.val1))))
        for v in vals:
            g = sorted(sub[np.isclose(sub.val1, v)].P_baja_val)
            ck(len(g) == len(n2.P_MALLA) and bool(np.isclose(g, n2.P_MALLA).all()),
               f"malla {ruta} {var1}={v:g}: los 12 P_baja de linspace(400, 2000, 12)", str(g))


def parte_b(d: dict) -> None:
    """(b) `m2n2_resumen.csv` y `m2n2_kalina.csv` recalculados desde las fases."""
    t = records(d)
    for a, (_, b) in zip(rep.filas_resumen(t), d["resumen"].iterrows()):
        fr = None if pd.isna(b.fuera_rango_min_kalina) else bool(b.fuera_rango_min_kalina)
        ok = (a["ruta"] == b.ruta and round(a["val1"], 9) == round(b.val1, 9)
              and a["n_genuinos"] == int(b.n_genuinos) and a["n_kalina"] == int(b.n_kalina)
              and a["n_espurios_N2"] == int(b.n_espurios_N2)
              and ig(a["P_baja_min_kalina"], b.P_baja_min_kalina) and ig(a["eta_min_kalina"], b.eta_min_kalina)
              and a["fuera_rango_min_kalina"] == fr and ig(a["P_baja_min_O2"], b.P_baja_min_O2)
              and ig(a["mejor_D"], b.mejor_D) and ig(a["P_baja_mejor_D"], b.P_baja_mejor_D)
              and multiset(a["criterios_bloqueantes_frecuentes"]) == multiset(b.criterios_bloqueantes_frecuentes))
        ck(ok, f"resumen recalculado: {a['ruta']} {a['var1']}={a['val1']:g}", str(a))
    kal = rep.filas_kalina(t)
    ck(len(kal) == len(d["kalina"]), "kalina: filas recalculadas == CSV", f"{len(kal)} vs {len(d['kalina'])}")
    for a, (_, b) in zip(kal, d["kalina"].iterrows()):
        fr = None if pd.isna(b.fuera_rango) else bool(b.fuera_rango)
        ok = (a["ruta"] == b.ruta and round(a["val1"], 9) == round(b.val1, 9)
              and ig(a["P_baja"], b.P_baja) and ig(a["eta"], b.eta) and ig(a["Wnet"], b.Wnet)
              and ig(a["m_O1"], b.m_O1) and ig(a["m_O2"], b.m_O2) and ig(a["d"], b.d)
              and a["fuera_rango"] == fr and s(a["detalle_rango"]) == s(b.detalle_rango))
        ck(ok, f"kalina recalculado: {a['ruta']} {a['var1']}={a['val1']:g} P={a['P_baja']:g}", f"d={a['d']} vs {b.d}")


def parte_c(d: dict) -> None:
    """(c) cada frontera, contra los puntos que la malla y la bisección evaluaron."""
    t = records(d)
    for _, f in d["fronteras"].iterrows():
        g = [r for r in t if r["ruta"] == f.ruta and round(r["val1"], 9) == round(f.val1, 9)]
        si = next((r for r in g if ig(r["P_baja_val"], f.P_baja_si)), None)
        no = next((r for r in g if ig(r["P_baja_val"], f.P_baja_no)), None)
        et = f'{f.ruta} {f.val1:g} {f.frontera} en {f.P_baja_si:g}..{f.P_baja_no:g}'
        if not ck(si is not None and no is not None, f"frontera {et}: extremos evaluados", "faltan puntos"):
            continue
        ck(ig(abs(f.P_baja_si - f.P_baja_no), f.ancho_final), f"frontera {et}: ancho == |si − no|", f"{f.ancho_final} vs {abs(f.P_baja_si - f.P_baja_no)}")
        ck(s(si["clasificacion"]) == s(f.clase_si) and s(si["codigos_falla"]) == s(f.codigos_falla_si), f"frontera {et}: clase/códigos del lado SÍ", f"{si['clasificacion']} vs {f.clase_si}")
        ck(n2.cumple(si, f.frontera) and not n2.cumple(no, f.frontera),
           f"frontera {et}: el lado SÍ cumple y el NO no", f"{si['clasificacion']} / {no['clasificacion']}")


def parte_d(d: dict) -> None:
    """(d) hojas del xlsx, filas de datos y las 5 filas de `Resumen`."""
    wb = load_workbook(XLSX)
    ck(wb.sheetnames[-5:] == list(HOJAS), "xlsx: 5 hojas M2_N2 al final", str(wb.sheetnames[-5:]))
    todo = pd.concat([d["malla"], d["biseccion"]], ignore_index=True)
    resumen_xl, ws_r = {}, wb["Resumen"]
    for r in range(2, ws_r.max_row + 1):
        v = ws_r.cell(r, 1).value
        if v in (None, "") or str(v).startswith("Descripción"):
            break
        resumen_xl[str(v)] = [ws_r.cell(r, j).value for j in range(1, 8)]
    for hoja, clave in zip(HOJAS, CLAVES):
        df = d[clave]
        ws = wb[hoja]
        cab = [re.sub(r"\s*\(.+\)$", "", str(ws.cell(1, j).value)) for j in range(1, ws.max_column + 1)]
        filas = 0
        for r in range(2, ws.max_row + 1):
            if ws.cell(r, 1).value in (None, ""):
                break
            filas += 1
        ck(cab == list(df.columns), f"{hoja}: cabecera == CSV", str(cab[:4]))
        ck(filas == len(df), f"{hoja}: {filas} filas de datos == {len(df)} del CSV", "")
        esp = em._fila_resumen(hoja, df, todo)
        fila = resumen_xl.get(hoja)
        got = None if fila is None else [fila[0], int(fila[1])] + [int(x) for x in fila[2:7]]
        ck(esp == got, f"Resumen {hoja}: puntos y conteos == CSV", f"esperado={esp} obtenido={got}")


def parte_e(d: dict) -> None:
    """(e) cada número del `REPORTE_M2_N2.md`, recalculado desde los CSV."""
    t = records(d)
    txt = (OUT / "REPORTE_M2_N2.md").read_text(encoding="utf-8")
    res, kal = rep.filas_resumen(t), rep.filas_kalina(t)
    for h in [f"## {i}." for i in range(1, 9)]:
        ck(h in txt, f"reporte: sección {h}", "")
    ck("(3 + 9) × 12 = 144 corridas" in txt, "reporte §2: 144 corridas", "")
    for ruta in n2.RUTAS:
        ck(rep._tabla_ruta(ruta, res) in txt, f"§3: tabla de la ruta {ruta}", "")
    ck(rep._tabla_kalina(kal) in txt, "§4: KALINA de menor d y de mayor η", "")
    nk = int(d["resumen"].n_kalina.sum())
    s6 = (f"Con dos cambios {'**sí**' if nk else '**no**'} se alcanza KALINA. "
          f"Ruta A (conservando la fuente del profesor): {rep._ruta_txt('A', res, t)}. "
          f"Ruta B: {rep._ruta_txt('B', res, t)}.")
    ck(s6 in txt, "§6: conclusión con la forma pedida", s6)
    sec5 = txt.split("## 5.")[1].split("## 6.")[0]
    m = re.search(r"Malla: (\d+) s · Bisección: (\d+) s · Total: (\d+) s", sec5)
    ck(m is not None and abs(int(m.group(3)) - int(m.group(1)) - int(m.group(2))) <= 1,
       "§5: total == malla + bisección", m.group(0) if m else "sin línea")
    ts = sorted(x["t_s"] for x in t if x.get("t_s") and not pd.isna(x["t_s"]))
    esp = (ts[0], ts[len(ts) // 2], ts[-1], sum(ts) / len(ts), float(len(ts)))
    m = re.search(r"mínimo ([\d.]+) s, mediana ([\d.]+) s, máximo ([\d.]+) s, media ([\d.]+) s \((\d+) puntos\)", sec5)
    got = tuple(float(x) for x in m.groups()) if m else ()
    ck(m is not None and all(abs(x - y) <= 0.05 for x, y in zip(esp, got)),
       "§5: t_s mínimo/mediana/máx/media/n", f"esperado={esp} obtenido={got}")
    fk = [k for k in kal if k["fuera_rango"]]
    sec7 = txt.split("## 7.")[1]
    ok = (f"**{len(fk)} de {len(kal)} puntos KALINA quedan fuera del rango confiable**" in sec7) if fk else (
        "- No se encontró ningún punto KALINA" in sec7 or f"Ninguno de los {len(kal)} puntos KALINA encontrados" in sec7)
    ck(ok, "§7: recuento de KALINA fuera del rango confiable", f"fk={len(fk)} de {len(kal)}")


def main() -> int:
    d = csvs()
    for parte in (parte_a, parte_b, parte_c, parte_d, parte_e):
        parte(d)
    print(f"\n{len(FALLOS)} fallos" if FALLOS else "\nTODO OK")
    for f in FALLOS:
        print("  -", f)
    return 1 if FALLOS else 0


if __name__ == "__main__":
    sys.exit(main())
