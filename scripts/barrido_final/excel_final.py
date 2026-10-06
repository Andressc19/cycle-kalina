"""M0/M1 — genera `datos/Excel_final_kalina.xlsx` con el formato del proyecto.

Hojas: `M0_protocolo` (la configuración congelada), `M1_profesor_P_alta`,
`M1_profesor_x_b`, `M1_profesor_T_amb`, `M1_profesor_base` (las filas de
`m1_filas.csv` de cada grupo), `M1_comparacion` y `Resumen`.

Formato (docs/METODOLOGIA_BARRIDOS_FINAL.md §3 y
`scripts/formatear_excel_barridos.py`, que se importa SIN modificar): unidad en
el título y color por tipo de columna, filtro en cada tabla, clase coloreada,
gris de fuera de rango, columnas de mensaje/error anchas con ajuste de texto y,
al final de cada tabla, dos filas vacías y una descripción (qué se iteró, qué
quedó fijo, conteo por clase). Es idempotente: lo vuelve a escribir COMPLETO a
partir de los CSV de `resultados/2026-10-06_final_M*` (hoy solo M1). El libro no
se edita a mano.
"""

from __future__ import annotations

import subprocess
import sys
from pathlib import Path

_DIR = Path(__file__).resolve().parent
_RAIZ = _DIR.parents[1]
for _p in (str(_RAIZ), str(_DIR), str(_DIR.parent)):   # parent = scripts/
    _p not in sys.path and sys.path.insert(0, _p)

import pandas as pd  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

from comun import PRESUPUESTO_S, PROFESOR, T_AMB_DISENO, TAU  # noqa: E402
from formatear_excel_barridos import (COLOR_CLASE, estilo_titulo,  # noqa: E402
                                      es_entrada, fmt, rango, tipo_col,
                                      unidad)
from src.excel_rango import marcar_tabla  # noqa: E402

XLSX = _RAIZ / "datos" / "Excel_final_kalina.xlsx"
M1 = _RAIZ / "resultados" / "2026-10-06_final_M1"
GRUPOS = (("P_alta", "M1_profesor_P_alta"), ("x_b", "M1_profesor_x_b"),
          ("T_amb", "M1_profesor_T_amb"), ("base", "M1_profesor_base"))
CLASES = ("KALINA", "CORREGIBLE", "INVIABLE", "NO_CONVERGIO")
# Unidades que las reglas del formateador no cubren (las pedidas en el
# TASK_CONTEXT van las primeras).
UNIDAD = {"m_O1": "-", "m_O2": "K", "m_O3": "kW", "m_O5": "-", "D": "-",
          "T9_amb": "K", "T_bub": "K", "t_s": "s", "valor": "-",
          "d_eta_pp": "pp", "nF": "n", "pid": "n", "bracket_lo": "K",
          "bracket_hi": "K", "A_recurrencias": "n"}
ANCHAS = {"mensaje", "err_margenes", "error_inesperado", "detalle_rango",
          "mensaje_antes", "mensaje_ahora", "B_error"}
DOCS = ("docs/METODOLOGIA_BARRIDOS_FINAL.md · "
        "docs/MATEMATICA_ITERACIONES_SENSIBILIDAD.md · "
        "docs/ESTUDIO_METODOS_SENSIBILIDAD.md · CONTEXT.md")


def _v(x):
    """Celda de openpyxl: `None` en vez de NaN, y tipos nativos de pandas."""
    if x is None or (isinstance(x, float) and pd.isna(x)):
        return None
    return x.item() if hasattr(x, "item") else x


def _conteo(vals) -> dict:
    c: dict = {}
    for v in vals:
        c[v] = c.get(v, 0) + 1
    return c


def _descripcion(ws, cab, filas, extra) -> None:
    """2 filas vacías y la descripción: qué se iteró, qué quedó fijo, clases."""
    col = {c: [f[i] for f in filas] for i, c in enumerate(cab)}
    itera, fijo = [], []
    for c in cab:
        if not es_entrada(c, ws.title):
            continue
        vals = [v for v in col[c] if v is not None]
        if not vals:
            continue
        u = UNIDAD.get(c) or unidad(c) or "-"
        txt = f"{c} ({u}): "
        if len(set(vals)) > 1:
            itera.append(txt + rango(vals))
        else:
            fijo.append(txt + fmt(vals[0]))
    lin = ["Descripción", "Se iteró: " + ("; ".join(itera) if itera else "no aplica")]
    for i in range(0, len(fijo), 4):
        lin.append(("Fijo: " if i == 0 else "      ") + "; ".join(fijo[i:i + 4]))
    for c in [c for c in cab if "clasif" in c or c.startswith("clas_")][:2]:
        n = _conteo(v for v in col[c] if v is not None)
        lin.append(f"{c}: " + " · ".join(f"{k} {v}" for k, v in n.items()))
    lin.extend(extra)
    r = ws.max_row + 3
    for k, t in enumerate(lin):
        ws.cell(r + k, 1, t)
    ws.cell(r, 1).font = Font(bold=True, size=12)


def escribir_tabla(wb, nombre, cab, filas, extra=()) -> None:
    """Una tabla con el formato del proyecto (ver el docstring del módulo)."""
    ws = wb.create_sheet(nombre)
    tipos = [tipo_col(c, nombre) for c in cab]
    ws.append([f"{c} ({UNIDAD.get(c) or unidad(c) or '-'})" for c in cab])
    for f in filas:
        ws.append([_v(v) for v in f])
    for j, (c, t) in enumerate(zip(cab, tipos), 1):
        estilo_titulo(ws.cell(1, j), t)
        w = max([len(str(ws.cell(1, j).value)) + 2, 10]
                + [len(fmt(_v(f[j - 1]))) + 2 for f in filas if _v(f[j - 1]) is not None])
        letra = get_column_letter(j)
        if t == "texto" or c in ANCHAS:
            w = 80
            for r in range(2, ws.max_row + 1):
                ws.cell(r, j).alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions[letra].width = min(w, 80)
        if t == "clase":
            for r in range(2, ws.max_row + 1):
                v = ws.cell(r, j).value
                if v in COLOR_CLASE:
                    ws.cell(r, j).fill = PatternFill("solid", fgColor=COLOR_CLASE[v])
    ws.row_dimensions[1].height = 34
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{letra}{ws.max_row}"
    if marcar_tabla(ws, cab):
        extra = [*extra, "Gris: fuera del rango confiable del motor (IAPWS G4-01 §6); el comentario de la celda lo dice."]
    _descripcion(ws, cab, filas, extra)


def commit() -> str:
    """`git rev-parse HEAD` del worktree, o un aviso si git no responde."""
    try:
        r = subprocess.run(["git", "rev-parse", "HEAD"], cwd=_RAIZ, capture_output=True, text=True, timeout=60)
        return r.stdout.strip() or f"sin salida de git (rc={r.returncode})"
    except (OSError, subprocess.SubprocessError) as exc:
        return f"no disponible ({type(exc).__name__})"


def protocolo() -> list:
    """Pares (clave, valor) de la configuración congelada de M0."""
    filas = [
        ("fecha", "2026-10-06"),
        ("commit", commit()),
        ("motor", "TeqpVerificado (protección A) + verificar_turbina (B) solo "
                  "en KALINA + bracket_tolerante=True (MA)"),
        ("motor_rango", "teqp, modo extrapolar_marcar (marca fuera_rango como "
                        "eje aparte; no reclasifica)"),
        ("presupuesto_s", PRESUPUESTO_S),
        ("T_AMB_DISENO (K, piso de diseño de O2)", T_AMB_DISENO),
        ("TAU", " · ".join(f"{k}={TAU[k]}" for k in ("O1", "O2", "O3", "O5"))),
        ("eps permitidas", "[0.75, 0.85]"),
        ("márgenes", "m_O1 = q4 − 0.90 · m_O2 = T_bub − T9 · m_O3 = Wnet · "
                     "m_O5 = min(q2 − 0.001, 0.999 − q2)"),
        ("D", "Σ viol_k / TAU_k, viol_k = max(0, −m_k)"),
        ("metodología", DOCS),
    ]
    filas += [(f"PUNTO PROFESOR · {k}", v) for k, v in PROFESOR.items()]
    return filas


def main() -> None:
    filas = pd.read_csv(M1 / "m1_filas.csv")
    cab = list(filas.columns)
    datos = filas.values.tolist()
    comparacion = pd.read_csv(M1 / "m1_comparacion.csv")
    wb = Workbook()
    wb.remove(wb.active)
    escribir_tabla(wb, "M0_protocolo", ["clave", "valor"], protocolo(),
                   extra=["La configuración congelada de todas las fases "
                          "(METODOLOGIA_BARRIDOS_FINAL.md §1)."])
    resumen = [["hoja", "puntos"] + list(CLASES) + ["errores_inesperados"]]
    for grupo, hoja in GRUPOS:
        sub = [f for f in datos if f[0] == grupo]
        if not sub:
            continue
        escribir_tabla(wb, hoja, cab, sub,
                       extra=[f"Grupo {grupo} de M1 (reproduce la Fase 1 del "
                              f"profesor del 2026-09-19)."])
        n = _conteo(f[cab.index("clasificacion")] for f in sub)
        i_err = cab.index("error_inesperado")
        resumen.append([hoja, len(sub)] + [n.get(c, 0) for c in CLASES]
                       + [sum(1 for f in sub if not pd.isna(f[i_err]))])
    escribir_tabla(wb, "M1_comparacion", list(comparacion.columns),
                   comparacion.values.tolist(),
                   extra=["Cada punto de M1 frente a su contraparte del "
                          "2026-09-19 (emparejamiento por el valor de la "
                          "variable barrida; por `hora` para T_amb)."])
    n = _conteo(comparacion["clase_ahora"])
    resumen.append(["M1_comparacion", len(comparacion)]
                   + [n.get(c, 0) for c in CLASES] + [0])
    escribir_tabla(wb, "Resumen", resumen[0], resumen[1:],
                   extra=["Conteo por clase de cada hoja del libro."])
    XLSX.parent.mkdir(parents=True, exist_ok=True)
    wb.save(XLSX)
    print(f"escrito {XLSX}")
    print("hojas:", wb.sheetnames)


if __name__ == "__main__":
    main()