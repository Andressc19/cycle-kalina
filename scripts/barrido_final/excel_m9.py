"""M9 — hojas finales en `datos/Excel_final_kalina.xlsx`.

Se corre SIEMPRE el último (después de `excel_final.py` … `excel_m6.py`): abre el
libro existente, crea o reemplaza las hojas `M9_husavik`, `M9_paso1`, `M9_paso2`,
`M9_busqueda`, `M9_escalera`, `M9_rangos`, `M9_motor_real` y `M9_figuras` con los
CSV de `resultados/2026-10-08_final_M9/` vía `excel_final.escribir_tabla` (mismo
formato, unidades, colores de clase, filtro y descripción al pie), y añade o
actualiza sus filas en `Resumen`. Es idempotente.
"""
from __future__ import annotations

import sys  # noqa: E402
from pathlib import Path

_DIR = Path(__file__).resolve().parent
_RAIZ = _DIR.parents[1]
for _p in (str(_RAIZ), str(_DIR), str(_DIR.parent)):
    if _p not in sys.path:
        sys.path.insert(0, _p)

import pandas as pd  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

import excel_final as ef  # noqa: E402
from m9_datos import OUT  # noqa: E402

XLSX = ef.XLSX
FIGURA = "M9_figuras"
FIGS = (OUT / "m9_escalera.png", OUT / "m9_plano_presiones.png")
PREFIJOS = ("M9_",)
# hoja -> (csv, campaña de sus clases (None = la propia), descripción)
HOJAS = (
    ("M9_husavik", "m9_husavik.csv", None,
     "M9 paso 0: el punto base de Húsavík con la malla de P_baja 500–1000 kPa "
     "y los dos pisos de O2 (283.15 y 278.15 K)."),
    ("M9_paso1", "m9_paso1.csv", None,
     "M9 paso 1: conservar UN grupo del profesor (6 escenarios × 100 LHS)."),
    ("M9_paso2", "m9_paso2.csv", None,
     "M9 paso 2: selección hacia adelante; `ronda` y `conservados` identifican "
     "cada conjunto."),
    ("M9_busqueda", "m9_busqueda.csv", None,
     "M9 paso 2.4: búsqueda dirigida por patrones sobre los conjuntos con 0 "
     "KALINA (aceptado = vecino que baja D)."),
    ("M9_escalera", "m9_escalera.csv", "camp",
     "M9: una fila por conjunto evaluado; `elegido` marca el camino."),
    ("M9_rangos", "m9_rangos.csv", "camp",
     "M9: rangos (mín/p5/mediana/p95/máx) de lo libre en los KALINA y en los "
     "KALINA confiables, con el KS contra los no KALINA genuinos."),
    ("M9_motor_real", "m9_motor_real.csv", None,
     "M9 paso 3: los 3 puntos con `motor=\"real\"` frente a su fila teqp; "
     "`d_eta_pp` = 100·(η_real − η_teqp)."),
)
UNIDADES = {"escenario": "\u2212", "conservados": "\u2212", "conjunto": "\u2212",
            "ambito": "\u2212", "variable": "\u2212", "id": "\u2212",
            "omitido": "V/F", "piso_tamb": "K", "P_baja_test": "kPa",
            "n": "n", "p05": "\u2212", "mediana": "\u2212", "p95": "\u2212",
            "min": "\u2212", "max": "\u2212", "ks_estadistico": "\u2212",
            "ks_p": "\u2212", "n_conservadas": "n", "n_eval": "n",
            "n_genuinos": "n", "n_kalina": "n", "n_kalina_confiable": "n",
            "D_min": "\u2212", "criterio_mas_roto": "texto",
            "elegido": "V/F", "kalina_por_busqueda": "V/F", "aceptado": "V/F",
            "origen": "\u2212", "clase_teqp": "\u2212", "clase_real": "\u2212",
            "eta_teqp": "\u2212", "eta_real": "\u2212", "coincide_clase": "V/F",
            "ronda": "n"}


def _campanas() -> dict:
    """Campaña combinada de paso 1 + paso 2 para las hojas sin clase propia."""
    p1 = pd.read_csv(OUT / "m9_paso1.csv")
    p2 = pd.read_csv(OUT / "m9_paso2.csv")
    return {"camp": pd.concat([p1, p2], ignore_index=True)}


def _fila_resumen(hoja, df, camp) -> list:
    if "clasificacion" in df.columns:
        n = df["clasificacion"].value_counts()
        err = int(df["error_inesperado"].notna().sum()) if "error_inesperado" in df.columns else 0
    elif "clase_real" in df.columns:
        n = df["clase_real"].value_counts()
        err = 0
    else:
        n = camp["clasificacion"].value_counts()
        err = int(camp["error_inesperado"].notna().sum())
    return [hoja, len(df)] + [int(n.get(c, 0)) for c in ef.CLASES] + [err]


def _reescribir_resumen(ws, filas) -> None:
    for r in range(ws.max_row, 1, -1):
        if str(ws.cell(r, 1).value or "").startswith(PREFIJOS):
            ws.delete_rows(r)
    pos = 2
    while ws.cell(pos, 1).value not in (None, "") and not str(
            ws.cell(pos, 1).value).startswith("Descripción"):
        pos += 1
    ws.insert_rows(pos, len(filas))
    for k, fila in enumerate(filas):
        for j, v in enumerate(fila, 1):
            ws.cell(pos + k, j, v)


def _figuras(wb) -> str:
    """Hoja con las 2 PNG de M9; sin Pillow o sin PNG, solo las rutas."""
    ws = wb.create_sheet(FIGURA)
    for k, p in enumerate(FIGS, 2):
        ws.cell(k, 1, p.name)
        ws.cell(k, 2, p.relative_to(_RAIZ).as_posix())
    c = len(FIGS) + 3
    ws.cell(c, 1, "Figuras: (a) escalera de conservación; (b) plano P_alta–P_baja.")
    ws.freeze_panes = "A2"
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 64
    faltan = [p.name for p in FIGS if not p.exists()]
    try:
        from openpyxl.drawing.image import Image as XLImage
    except ImportError:
        ws.cell(c + 2, 1, "Pillow no disponible: no se insertaron las imágenes.")
        return f"{FIGURA}: sin Pillow, solo las rutas"
    if faltan:
        ws.cell(c + 2, 1, "Faltan: " + ", ".join(faltan))
        return f"{FIGURA}: faltan {', '.join(faltan)}"
    fila = c + 2
    for p in FIGS:
        img = XLImage(str(p))
        ws.add_image(img, f"A{fila}")
        fila += -(-img.height // 20) + 1
    return f"{len(FIGS)} imágenes insertadas en {FIGURA}"


def main() -> None:
    faltan = [str(OUT / csv) for _h, csv, _k, _d in HOJAS
              if not (OUT / csv).exists()]
    if not XLSX.exists():
        raise SystemExit(f"No existe {XLSX}: correr antes excel_final.py")
    if faltan:
        raise SystemExit("Faltan: " + ", ".join(faltan))
    ef.UNIDAD.update(UNIDADES)
    wb = load_workbook(XLSX)
    if "Resumen" not in wb.sheetnames:
        raise SystemExit("El libro no tiene la hoja Resumen: correr antes excel_final.py")
    viejas = [h for h, _c, _k, _d in HOJAS if h in wb.sheetnames]
    for h in viejas + ([FIGURA] if FIGURA in wb.sheetnames else []):
        del wb[h]
    camp = _campanas()
    filas = []
    for hoja, csv, clave, desc in HOJAS:
        df = pd.read_csv(OUT / csv)
        ef.escribir_tabla(wb, hoja, list(df.columns), df.values.tolist(), extra=[desc])
        filas.append(_fila_resumen(hoja, df, camp.get(clave)))
    msg = _figuras(wb)
    _reescribir_resumen(wb["Resumen"], filas)
    wb.save(XLSX)
    print(f"escrito {XLSX}")
    print("hojas:", wb.sheetnames)
    print(msg)


if __name__ == "__main__":
    main()
