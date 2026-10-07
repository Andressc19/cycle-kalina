"""Verifica las hojas M2S_/M4_ de `datos/Excel_final_kalina.xlsx` (TASK_CONTEXT
de excel_m2s_m4): nombres de hoja, filas de datos frente a los CSV de origen
(500 en `M4_lhs500`), la columna `lote`, las filas de `Resumen`, las unidades
nuevas de las cabeceras, las tres figuras de `M4_figuras` y la huella de todo
el libro (para comparar entre dos ejecuciones y comprobar idempotencia).
Imprime cada comprobación y termina con OK si todas pasan.
"""
from __future__ import annotations

import hashlib
import sys
from pathlib import Path

_DIR = Path(__file__).resolve().parent
_RAIZ = _DIR.parents[1]
for _p in (str(_RAIZ), str(_DIR), str(_DIR.parent)):   # parent = scripts/
    if _p not in sys.path:
        sys.path.insert(0, _p)

import pandas as pd  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

XLSX = _RAIZ / "datos" / "Excel_final_kalina.xlsx"
M2S = _RAIZ / "resultados" / "2026-10-06_final_M2S"
M4 = _RAIZ / "resultados" / "2026-10-06_final_M4"
M4B = _RAIZ / "resultados" / "2026-10-07_final_M4b"
FIGURA = "M4_figuras"
FIGS = ("m4b_plano_presiones.png", "m4b_hist_dP_rP.png", "m4b_rsa_ks.png")
# hoja -> [csv de origen] (mismo mapa que excel_m2s_m4.py)
HOJAS = {
    "M2S_malla": [M2S / "m2s_malla.csv"],
    "M2S_biseccion": [M2S / "m2s_biseccion.csv"],
    "M2S_resumen": [M2S / "m2s_resumen.csv"],
    "M2S_kalina": [M2S / "m2s_kalina.csv"],
    "M4_lhs500": [M4 / "m4_lhs.csv", M4B / "m4b_lhs.csv"],
    "M4_rsa": [M4B / "m4b_rsa.csv"],
    "M4_prcc": [M4B / "m4b_prcc.csv"],
    "M4_pawn": [M4B / "m4b_pawn.csv"],
    "M4_presiones_rangos": [M4B / "m4b_rangos_presiones.csv"],
    "M4_ventana_Palta": [M4B / "m4b_ventana_por_palta.csv"],
    "M4_ventana_Tfuente": [M4B / "m4b_ventana_por_tfuente.csv"],
    "M4_ventana_Tsumidero": [M4B / "m4b_ventana_por_tsumidero.csv"],
    "M4_bloqueo_rP": [M4B / "m4b_bloqueo_por_rP.csv"],
    "M4_top_kalina": [M4B / "m4b_top_kalina.csv"],
    "M4_convergencia": [M4B / "m4b_convergencia.csv"],
}
# cabeceras exactas que el TASK pide ver con su unidad nueva
UNIDADES = {
    "M4_ventana_Palta": ["lo (kPa)", "hi (kPa)", "dP_min (kPa)", "rP_min (\u2212)",
                         "n_kalina (\u2212)", "pct_kalina (%)"],
    "M4_ventana_Tfuente": ["lo (K)", "hi (K)"],
    "M4_ventana_Tsumidero": ["lo (K)", "hi (K)"],
    "M4_rsa": ["KS (\u2212)", "p_valor (\u2212)"],
    "M4_prcc": ["prcc (\u2212)", "p_valor (\u2212)"],
    "M2S_resumen": ["n_kalina (\u2212)"],
}


def filas_datos(ws) -> int:
    """Filas de la tabla: se corta en la primera vacía o en la descripción."""
    n = 0
    for r in range(2, ws.max_row + 1):
        v = ws.cell(r, 1).value
        if v in (None, "") or str(v).startswith("Descripción"):
            break
        n += 1
    return n


def resumen(ws, hoja: str):
    """[puntos, KALINA, CORREGIBLE, INVIABLE, NO_CONVERGIO, errores] de `hoja`."""
    for r in range(2, ws.max_row + 1):
        v = ws.cell(r, 1).value
        if v in (None, "") or str(v).startswith("Descripción"):
            return None
        if v == hoja:
            return [ws.cell(r, c).value for c in range(2, 8)]
    return None


def huella(wb) -> str:
    """sha256 de títulos y valores de celda de todo el libro (sin estilos)."""
    h = hashlib.sha256()
    for ws in wb.worksheets:
        h.update(ws.title.encode("utf-8"))
        for fila in ws.iter_rows(values_only=True):
            h.update(repr(fila).encode("utf-8"))
    return h.hexdigest()[:16]


def main() -> int:
    if not XLSX.exists():
        print(f"FALLO: no existe {XLSX}")
        return 1
    wb = load_workbook(XLSX)
    print("hojas:", wb.sheetnames)
    fallos = 0

    def check(ok: bool, txt: str) -> None:
        nonlocal fallos
        print(("  OK   " if ok else "  FALLO ") + txt)
        fallos += 0 if ok else 1

    faltan = [h for h in list(HOJAS) + [FIGURA] if h not in wb.sheetnames]
    check(not faltan, f"nombres de hoja (faltan: {faltan or 'ninguno'})")
    extra = [h for h in wb.sheetnames
             if h not in HOJAS and h != FIGURA and not h.startswith(("M5_", "M4_"))]
    print(f"  info  hojas escritas por otros scripts (M5 corre en paralelo): "
          f"{len(extra)} — {extra or 'ninguna'}")

    for hoja, csvs in HOJAS.items():
        if hoja not in wb.sheetnames:
            check(False, f"{hoja}: no existe")
            continue
        ws = wb[hoja]
        esp = sum(len(pd.read_csv(c)) for c in csvs)
        n = filas_datos(ws)
        check(n == esp, f"{hoja}: {n} filas de datos = {esp} de los CSV "
                        f"({' + '.join(c.name for c in csvs)})")
        fila = resumen(wb["Resumen"], hoja) if "Resumen" in wb.sheetnames else None
        check(fila is not None and fila[0] == n,
              f"{hoja}: fila en Resumen con puntos = {n} (fila: {fila})")
        cab = [str(ws.cell(1, c).value) for c in range(1, ws.max_column + 1)
               if ws.cell(1, c).value is not None]
        check(all(h.endswith(")") for h in cab),
              f"{hoja}: {len(cab)} cabeceras con unidad")
        check(bool(ws.auto_filter.ref), f"{hoja}: auto_filtro en {ws.auto_filter.ref}")
        desc = any(str(ws.cell(r, 1).value or "").startswith("Descripción")
                   for r in range(n + 2, ws.max_row + 1))
        check(desc, f"{hoja}: descripción al pie (2 filas vacías + texto)")
        miss = [u for u in UNIDADES.get(hoja, []) if u not in cab]
        if UNIDADES.get(hoja):
            check(not miss, f"{hoja}: unidades nuevas (faltan: {miss or 'ninguna'})")

    if "M4_lhs500" in wb.sheetnames:
        ws = wb["M4_lhs500"]
        check(str(ws.cell(1, 1).value).startswith("lote"),
              f"M4_lhs500: la columna A de la cabecera es «{ws.cell(1, 1).value}»")
        lotes = [ws.cell(r, 1).value for r in range(2, filas_datos(ws) + 2)]
        check(lotes.count("M4") == 300 and lotes.count("M4b") == 200,
              f"M4_lhs500: lote M4={lotes.count('M4')} y M4b={lotes.count('M4b')} "
              f"(esperado 300 y 200)")

    if FIGURA in wb.sheetnames:
        wsf = wb[FIGURA]
        rutas = [wsf.cell(r, 2).value for r in range(2, 2 + len(FIGS))]
        esperadas = [(M4B / n).relative_to(_RAIZ).as_posix() for n in FIGS]
        check(rutas == esperadas, f"{FIGURA}: rutas = {rutas}")
        print(f"  info  {FIGURA}: imagen(es) insertada(s) = {len(wsf._images)} "
              f"(0 = solo las rutas, por falta de Pillow)")
    else:
        check(False, f"{FIGURA}: no existe")

    print(f"huella del libro: {huella(wb)}")
    print("OK" if not fallos else f"FALLO ({fallos} comprobaciones)")
    return 0 if not fallos else 1


if __name__ == "__main__":
    sys.exit(main())
