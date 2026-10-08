"""M6 — hojas finales en `datos/Excel_final_kalina.xlsx`.

Con el mismo patrón que `excel_m5c_m7_m8.py`: abre el libro ya escrito por los
seis scripts anteriores (`excel_final.py`, `excel_m2.py`, `excel_m3_m2r.py`,
`excel_m2s_m4.py`, `excel_m5.py` y `excel_m5c_m7_m8.py`), crea o reemplaza las
hojas `M6_verificacion`, `M6_comparacion`, `M6_piso_30_4C`,
`M6_subenfriamiento` y `M6_subenf_evals` con los CSV de
`resultados/2026-10-08_final_M6` vía `excel_final.escribir_tabla` (mismo
formato, unidades, colores de clase, filtro y descripción al pie) y añade o
actualiza sus filas en `Resumen`. Es idempotente y se corre SIEMPRE el último.
"""
from __future__ import annotations

import sys
from pathlib import Path

_DIR = Path(__file__).resolve().parent
_RAIZ = _DIR.parents[1]
for _p in (str(_RAIZ), str(_DIR), str(_DIR.parent)):   # parent = scripts/
    if _p not in sys.path:
        sys.path.insert(0, _p)

import pandas as pd  # noqa: E402
from openpyxl import load_workbook  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

import excel_final as ef  # noqa: E402

M6 = _RAIZ / "resultados" / "2026-10-08_final_M6"
XLSX = ef.XLSX
PREFIJOS = ("M6_",)
# hoja -> (csv, campaña de sus clases (None = la del propio CSV), descripción)
HOJAS = (
    ("M6_verificacion", "m6_verificacion.csv", None,
     "M6-A: los 12 puntos evaluados con teqp y con el motor real "
     "(AmmoniaWaterAdapter), 24 filas; `motor` dice cuál."),
    ("M6_comparacion", "m6_comparacion.csv", "M6",
     "M6-A: comparación por punto teqp vs. motor real; `dentro_banda` = "
     "|d_eta_pp| ≤ 0.7 pp (banda de η, G4-01); en blanco si no hay η_real."),
    ("M6_piso_30_4C", "m6_piso_30_4C.csv", None,
     "M6-B1: los 12 puntos con el piso por defecto de O2 "
     "(t_amb_diseno = 303.55 K = 30.4 °C), solo teqp."),
    ("M6_subenfriamiento", "m6_subenfriamiento.csv", "M6",
     "M6-B2: P_baja mínima por bisección para exigir ΔT_sub = 2.0 y 5.0 K; "
     "`nota` avisa huecos NO_CONVERGIO del predicado."),
    ("M6_subenf_evals", "m6_subenfriamiento_evals.csv", None,
     "M6-B2: todas las evaluaciones de las 16 bisecciones "
     "(`paso` = lo, hi, it1…); `cumple` = KALINA y m_O2 ≥ ΔT_sub."),
)
# Unidades de columnas que las reglas del formateador no cubren.
UNIDADES = {"id": "\u2212", "origen": "texto", "paso": "\u2212",
            "ΔT_sub": "K", "alcanzable": "V/F", "cumple": "V/F",
            "P_baja_orig": "kPa", "P_baja_req": "kPa", "ΔP_baja": "kPa",
            "rP_req": "\u2212", "m_O1_req": "\u2212", "iteraciones": "n",
            "coincide_clase": "V/F", "dentro_banda": "V/F",
            "clase_teqp": "categoría", "clase_real": "categoría",
            "eta_teqp": "\u2212", "eta_real": "\u2212",
            "m_O1_teqp": "\u2212", "m_O1_real": "\u2212",
            "m_O2_teqp": "K", "m_O2_real": "K",
            "eta_orig": "\u2212", "eta_req": "\u2212", "pid": "n"}


def _campanas() -> dict:
    """Base de las hojas derivadas: la mitad teqp de la verificación."""
    ver = pd.read_csv(M6 / "m6_verificacion.csv")
    return {"M6": ver[ver["motor"] == "teqp"]}


def _fila_resumen(hoja: str, df, camp) -> list:
    """[hoja, filas, KALINA, CORREGIBLE, INVIABLE, NO_CONVERGIO, errores]."""
    if "clasificacion" in df.columns:
        n = df["clasificacion"].value_counts()
        err = int(df["error_inesperado"].notna().sum()) if "error_inesperado" in df.columns else 0
    else:
        n = camp["clasificacion"].value_counts()
        err = int(camp["error_inesperado"].notna().sum())
    return [hoja, len(df)] + [int(n.get(c, 0)) for c in ef.CLASES] + [err]


def _ultima_fila(ws) -> int:
    """Última fila de datos de `Resumen` (deja fuera las 2 vacías y la descripción)."""
    ult = 1
    for r in range(2, ws.max_row + 1):
        v = ws.cell(r, 1).value
        if v in (None, "") or str(v).startswith("Descripción"):
            break
        ult = r
    return ult


def _reescribir_resumen(ws, filas: list) -> None:
    """Quita las filas viejas de estas hojas y reinserta las nuevas."""
    for r in range(ws.max_row, 1, -1):
        if str(ws.cell(r, 1).value or "").startswith(PREFIJOS):
            ws.delete_rows(r)
    pos = _ultima_fila(ws) + 1
    ws.insert_rows(pos, len(filas))
    for k, fila in enumerate(filas):
        for j, v in enumerate(fila, 1):
            ws.cell(pos + k, j, v)
    j = 1
    while ws.cell(1, j + 1).value is not None:
        j += 1
    ws.auto_filter.ref = f"A1:{get_column_letter(j)}{pos + len(filas) - 1}"


def main() -> None:
    faltan = [str(M6 / csv) for _h, csv, _k, _d in HOJAS if not (M6 / csv).exists()]
    if not XLSX.exists():
        raise SystemExit(f"No existe {XLSX}: correr antes scripts/barrido_final/excel_final.py")
    if faltan:
        raise SystemExit("Faltan: " + ", ".join(faltan))
    ef.UNIDAD.update(UNIDADES)
    wb = load_workbook(XLSX)
    if "Resumen" not in wb.sheetnames:
        raise SystemExit("El libro no tiene la hoja Resumen: correr antes excel_final.py")
    for h in [h for h, _c, _k, _d in HOJAS if h in wb.sheetnames]:
        del wb[h]
    camp = _campanas()
    filas = []
    for hoja, csv, clave, desc in HOJAS:
        df = pd.read_csv(M6 / csv)
        ef.escribir_tabla(wb, hoja, list(df.columns), df.values.tolist(), extra=[desc])
        filas.append(_fila_resumen(hoja, df, camp.get(clave)))
    _reescribir_resumen(wb["Resumen"], filas)
    wb.save(XLSX)
    print(f"escrito {XLSX}")
    print("hojas:", len(wb.sheetnames), wb.sheetnames)


if __name__ == "__main__":
    main()
