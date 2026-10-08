"""M5c/M4c/M7/M8 — hojas finales en `datos/Excel_final_kalina.xlsx`.

Abre el libro que ya escribieron `excel_final.py`, `excel_m2.py`,
`excel_m3_m2r.py`, `excel_m2s_m4.py` y `excel_m5.py`, crea o reemplaza las
hojas M5c_*, M4c_*, M7_*, M8_ANECDOTICO_oat y la de figuras
`M5c_M4c_M7_M8_figuras` con los CSV de `resultados/2026-10-07_final_*` vía
`excel_final.escribir_tabla` (mismo formato, unidades, colores de clase,
filtro y descripción al pie), y añade o actualiza sus filas en `Resumen`. Es
idempotente y se corre SIEMPRE el último.
"""
from __future__ import annotations

import sys
from pathlib import Path

_DIR = Path(__file__).resolve().parent
_RAIZ = _DIR.parents[1]
for _p in (str(_RAIZ), str(_DIR), str(_DIR.parent)):   # parent = scripts/
    if _p not in sys.path:
        sys.path.insert(0, _p)

import pandas as pd  # noqa: E402
from openpyxl import load_workbook  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

import excel_final as ef  # noqa: E402
from formatear_excel_barridos import estilo_titulo  # noqa: E402

M5C = _RAIZ / "resultados" / "2026-10-07_final_M5c"
M4C = _RAIZ / "resultados" / "2026-10-07_final_M4c"
M7 = _RAIZ / "resultados" / "2026-10-07_final_M7"
M8 = _RAIZ / "resultados" / "2026-10-07_final_M8_anecdotico"
XLSX = ef.XLSX
FIGURA = "M5c_M4c_M7_M8_figuras"
FIGS = (M5C / "m5c_mapa.png", M4C / "m4c_techo.png",
        M7 / "m7_convergencia.png", M8 / "m8_curvas.png")
PREFIJOS = ("M5c_", "M4c_", "M7_", "M8_")
ANECDOTO = "ANECDÓTICO: valores fuera de [0.40, 0.95] no son realistas"
# hoja -> (carpeta, csv, campaña de sus clases (None = la propia), descripción)
HOJAS = (
    ("M5c_malla", M5C, "m5c_malla.csv", None,
     "M5c: la malla de 165 puntos (T_fuente × eps_cond × T_sumidero con las "
     "presiones del profesor); `parte`/`id`/`tag` identifican cada punto."),
    ("M5c_frontera", M5C, "m5c_frontera.csv", "M5c",
     "M5c: una fila por (T_sumidero, eps_cond) con las celdas KALINA y el "
     "rango de T_fuente (mín–máx) donde aparece."),
    ("M4c_lhs", M4C, "m4c_lhs.csv", None,
     "M4c: el LHS de 150 puntos con P_baja ∈ [1500, 2500] kPa; `omitido` "
     "marca los que no se evaluaron por P_baja ≥ P_alta − 300."),
    ("M4c_techo", M4C, "m4c_techo.csv", "M4c",
     "M4c: el conteo de KALINA por franjas de 100 kPa de P_baja (con M4 + "
     "M4b + M4c) y su P_baja máxima."),
    ("M4c_techo_Palta", M4C, "m4c_techo_por_palta.csv", "M4c",
     "M4c: el mismo conteo por 6 franjas de P_alta de 2000 a 5000 kPa."),
    ("M7_optimos", M7, "m7_optimos.csv", None,
     "M7: el último punto aceptado de cada arranque A1…A5 (el óptimo de la "
     "búsqueda por patrones)."),
    ("M7_trayectorias", M7, "m7_trayectorias.csv", None,
     "M7: los puntos aceptados de cada arranque, en orden de iteración."),
    ("M7_evaluaciones", M7, "m7_evaluaciones.csv", None,
     "M7: todas las evaluaciones del optimizador, aceptadas o no."),
    ("M8_ANECDOTICO_oat", M8, "m8_oat.csv", None, ANECDOTO),
)
# Unidades de columnas nuevas de esta tarea.
UNIDADES = {"parte": "\u2212", "id": "\u2212", "tag": "\u2212",
            "omitido": "V/F", "pid": "n", "codigos_falla": "texto",
            "repliegues_lo": "n", "repliegues_hi": "n",
            "arranque": "\u2212", "iteracion": "n", "paso": "\u2212",
            "aceptado": "V/F", "evaluaciones": "n", "iteraciones": "n",
            "base": "\u2212", "variable": "\u2212", "valor": "\u2212",
            "realista": "V/F", "fuente": "\u2212", "franja": "\u2212",
            "lo": "kPa", "hi": "kPa", "n": "n", "n_kalina": "\u2212",
            "pct_kalina": "%", "P_baja_max_kalina": "kPa"}


def _campanas() -> dict:
    """Punto base de cada campaña para las hojas derivadas (sin clases propias)."""
    return {"M5c": pd.read_csv(M5C / "m5c_malla.csv"),
            "M4c": pd.read_csv(M4C / "m4c_lhs.csv")}


def _fila_resumen(hoja: str, df, camp) -> list:
    """[hoja, filas, KALINA, CORREGIBLE, INVIABLE, NO_CONVERGIO, errores].

    La clase sale del propio df si tiene `clasificacion`; si no, de la campaña
    de la que deriva la tabla.
    """
    if "clasificacion" in df.columns:
        n = df["clasificacion"].value_counts()
        err = int(df["error_inesperado"].notna().sum()) if "error_inesperado" in df.columns else 0
    else:
        n = camp["clasificacion"].value_counts()
        err = int(camp["error_inesperado"].notna().sum())
    return [hoja, len(df)] + [int(n.get(c, 0)) for c in ef.CLASES] + [err]


def _ultima_fila(ws) -> int:
    """Última fila de datos de `Resumen` (deja fuera las 2 vacías y la descripción)."""
    ult = 1
    for r in range(2, ws.max_row + 1):
        v = ws.cell(r, 1).value
        if v in (None, "") or str(v).startswith("Descripción"):
            break
        ult = r
    return ult


def _reescribir_resumen(ws, filas: list) -> None:
    """Quita las filas viejas de estas hojas y reinserta las nuevas."""
    for r in range(ws.max_row, 1, -1):
        if str(ws.cell(r, 1).value or "").startswith(PREFIJOS):
            ws.delete_rows(r)
    pos = _ultima_fila(ws) + 1
    ws.insert_rows(pos, len(filas))
    for k, fila in enumerate(filas):
        for j, v in enumerate(fila, 1):
            ws.cell(pos + k, j, v)
    j = 1
    while ws.cell(1, j + 1).value is not None:
        j += 1
    ws.auto_filter.ref = f"A1:{get_column_letter(j)}{pos + len(filas) - 1}"


def _figuras(wb) -> str:
    """Hoja con las cuatro PNG; sin Pillow o sin PNG, solo las rutas."""
    ws = wb.create_sheet(FIGURA)
    for j in (1, 2):
        estilo_titulo(ws.cell(1, j), "texto")
    for k, p in enumerate(FIGS, 2):
        ws.cell(k, 1, p.name)
        ws.cell(k, 2, p.relative_to(_RAIZ).as_posix())
    c = len(FIGS) + 3
    ws.cell(c, 1, "Figuras: (a) mapa de M5c; (b) techo por P_baja de M4c; "
                  "(c) convergencia de M7; (d) curvas de M8 (ANECDÓTICO).")
    ws.freeze_panes = "A2"
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 64
    faltan = [p.name for p in FIGS if not p.exists()]
    try:
        from openpyxl.drawing.image import Image as XLImage
    except ImportError:
        ws.cell(c + 2, 1, "Pillow no disponible: no se insertaron las imágenes.")
        return f"{FIGURA}: sin Pillow, solo las rutas"
    if faltan:
        ws.cell(c + 2, 1, "Faltan: " + ", ".join(faltan))
        return f"{FIGURA}: faltan {', '.join(faltan)}"
    fila = c + 2
    for p in FIGS:
        img = XLImage(str(p))
        ws.add_image(img, f"A{fila}")
        fila += -(-img.height // 20) + 1   # filas que ocupa la imagen (20 píxeles) + 1
    return f"{len(FIGS)} imágenes insertadas en {FIGURA}"


def main() -> None:
    faltan = [str(carp / csv) for _h, carp, csv, _k, _d in HOJAS
              if not (carp / csv).exists()]
    faltan += [str(p) for p in FIGS if not p.exists()]
    if not XLSX.exists():
        raise SystemExit(f"No existe {XLSX}: correr antes scripts/barrido_final/excel_final.py")
    if faltan:
        raise SystemExit("Faltan: " + ", ".join(faltan))
    ef.UNIDAD.update(UNIDADES)
    wb = load_workbook(XLSX)
    if "Resumen" not in wb.sheetnames:
        raise SystemExit("El libro no tiene la hoja Resumen: correr antes excel_final.py")
    viejas = [h for h, _p, _c, _k, _d in HOJAS if h in wb.sheetnames]
    for h in viejas + ([FIGURA] if FIGURA in wb.sheetnames else []):
        del wb[h]
    camp = _campanas()
    filas = []
    for hoja, carpeta, csv, clave, desc in HOJAS:
        df = pd.read_csv(carpeta / csv)
        ef.escribir_tabla(wb, hoja, list(df.columns), df.values.tolist(), extra=[desc])
        filas.append(_fila_resumen(hoja, df, camp.get(clave)))
    msg = _figuras(wb)
    _reescribir_resumen(wb["Resumen"], filas)
    wb.save(XLSX)
    print(f"escrito {XLSX}")
    print("hojas:", wb.sheetnames)
    print(msg)


if __name__ == "__main__":
    main()