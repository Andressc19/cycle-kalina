"""M10 — verifica el Excel final tras `excel_m10.py`.

Comprueba que las hojas M10 existen con el número de filas de datos de su CSV,
que todas las hojas anteriores (incluidas las de M9) siguen presentes y que
`Resumen` tiene las filas M10. Imprime OK/FALLA y sale con código 1 si algo falla.
"""
from __future__ import annotations

import sys  # noqa: E402
from pathlib import Path

_DIR = Path(__file__).resolve().parent
_RAIZ = _DIR.parents[1]
sys.path[:0] = [str(_RAIZ), str(_DIR)]

import csv  # noqa: E402

from openpyxl import load_workbook  # noqa: E402

import excel_final as ef  # noqa: E402
import excel_m10 as e10  # noqa: E402
import excel_m9 as e9  # noqa: E402

HOJAS_ANTES = {
    "M0_protocolo", "M1_profesor_P_alta", "M1_profesor_x_b", "M1_profesor_T_amb",
    "M1_profesor_base", "M1_comparacion", "Resumen", "M2_N1_malla",
    "M2_N1_biseccion", "M2_N1_fronteras", "M2_N1_resumen", "M2_N2_malla",
    "M2_N2_biseccion", "M2_N2_fronteras", "M2_N2_kalina", "M2_N2_resumen",
    "M3_ofat", "M3_resumen", "M3_cruces", "M2R_n1", "M2R_n2a", "M2R_n2b",
    "M2R_n3", "M2R_fronteras", "M2R_kalina", "M2R_resumen", "M3_figuras",
    "M2S_malla", "M2S_biseccion", "M2S_resumen", "M2S_kalina", "M4_lhs500",
    "M4_rsa", "M4_prcc", "M4_pawn", "M4_presiones_rangos", "M4_ventana_Palta",
    "M4_ventana_Tfuente", "M4_ventana_Tsumidero", "M4_bloqueo_rP",
    "M4_top_kalina", "M4_convergencia", "M4_figuras", "M5_lhs", "M5_resumen",
    "M5_rangos", "M5_cercano_profesor", "M5_escenarios", "M5_figuras",
    "M6_verificacion", "M6_comparacion", "M6_piso_30_4C", "M6_subenfriamiento",
    "M6_subenf_evals", "M5c_malla", "M5c_frontera", "M4c_lhs", "M4c_techo",
    "M4c_techo_Palta", "M7_optimos", "M7_trayectorias", "M7_evaluaciones",
    "M8_ANECDOTICO_oat", "M5c_M4c_M7_M8_figuras",
    *(h for h, _c, _k, _d in e9.HOJAS), e9.FIGURA,
}
_falla = 0


def chk(nombre, ok, detalle="") -> None:
    global _falla
    if not ok:
        _falla += 1
    print(f'{"OK  " if ok else "FALLA"} {nombre}{" — " + detalle if detalle else ""}',
          flush=True)


def _n_csv(p: Path) -> int:
    with open(p, encoding="utf-8", newline="") as fh:
        return sum(1 for _ in csv.reader(fh)) - 1


def _n_datos(ws) -> int:
    """Filas de datos: desde la 2 hasta la primera celda A vacía o la descripción."""
    n = 0
    for r in range(2, ws.max_row + 1):
        v = ws.cell(r, 1).value
        if v in (None, "") or str(v).startswith("Descripción"):
            break
        n += 1
    return n


def main() -> None:
    if not ef.XLSX.exists():
        raise SystemExit(f"No existe {ef.XLSX}")
    wb = load_workbook(ef.XLSX, read_only=False)
    for h in sorted(HOJAS_ANTES):
        chk(f"hoja previa {h}", h in wb.sheetnames)
    for hoja, csvf, _d in e10.HOJAS:
        if hoja not in wb.sheetnames:
            chk(f"hoja {hoja} existe", False)
            continue
        n = _n_datos(wb[hoja])
        ncsv = _n_csv(e10.OUT / csvf)
        chk(f"{hoja}: filas", n == ncsv, f"{n} vs {ncsv}")
    chk(f"hoja {e10.FIGURA} existe", e10.FIGURA in wb.sheetnames)
    chk("imágenes M10 insertadas",
        e10.FIGURA in wb.sheetnames
        and len(getattr(wb[e10.FIGURA], "_images", [])) > 0)
    if "Resumen" in wb.sheetnames:
        ws = wb["Resumen"]
        m10 = [r for r in range(2, ws.max_row + 1)
               if str(ws.cell(r, 1).value or "").startswith("M10_")]
        chk("Resumen tiene filas M10", len(m10) > 0, f"{len(m10)} filas")
    print(f"\n{'TODO OK' if _falla == 0 else str(_falla) + ' FALLAS'}", flush=True)
    sys.exit(1 if _falla else 0)


if __name__ == "__main__":
    main()
