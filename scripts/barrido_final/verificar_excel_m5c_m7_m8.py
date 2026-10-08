"""Verifica las hojas M5c/M4c/M7/M8 de `datos/Excel_final_kalina.xlsx`.

Comprobaciones:
- nombre y filas de datos de cada hoja frente a su CSV de origen
- `M8_ANECDOTICO_oat` lleva la descripción ANECDÓTICO
- `M5c_M4c_M7_M8_figuras` lista las cuatro rutas m5c_mapa/m4c_techo/
  m7_convergencia/m8_curvas
- `Resumen` tiene las filas nuevas (las 9 hojas de datos, sin la de figuras)
- todas las hojas anteriores (M0…M5) siguen presentes
- huella del libro (idempotencia)

Imprime cada comprobación y termina con OK si todas pasan.
"""
from __future__ import annotations

import hashlib
import sys
from pathlib import Path

_DIR = Path(__file__).resolve().parent
_RAIZ = _DIR.parents[1]
sys.path.insert(0, str(_RAIZ))
sys.path.insert(0, str(_DIR))
sys.path.insert(0, str(_DIR.parent))

from openpyxl import load_workbook  # noqa: E402
import pandas as pd  # noqa: E402

XLSX = _RAIZ / "datos" / "Excel_final_kalina.xlsx"
M5C = _RAIZ / "resultados" / "2026-10-07_final_M5c"
M4C = _RAIZ / "resultados" / "2026-10-07_final_M4c"
M7 = _RAIZ / "resultados" / "2026-10-07_final_M7"
M8 = _RAIZ / "resultados" / "2026-10-07_final_M8_anecdotico"
# hoja -> (csv de origen, filas esperadas)
NUEVAS = {
    "M5c_malla": (M5C / "m5c_malla.csv", 165),
    "M5c_frontera": (M5C / "m5c_frontera.csv", 15),
    "M4c_lhs": (M4C / "m4c_lhs.csv", 150),
    "M4c_techo": (M4C / "m4c_techo.csv", 12),
    "M4c_techo_Palta": (M4C / "m4c_techo_por_palta.csv", 6),
    "M7_optimos": (M7 / "m7_optimos.csv", 5),
    "M7_trayectorias": (M7 / "m7_trayectorias.csv", 41),
    "M7_evaluaciones": (M7 / "m7_evaluaciones.csv", 560),
    "M8_ANECDOTICO_oat": (M8 / "m8_oat.csv", 160),
}
FIGURA = "M5c_M4c_M7_M8_figuras"
FIGS = ("m5c_mapa.png", "m4c_techo.png", "m7_convergencia.png", "m8_curvas.png")
PRE = {"M0_protocolo", "M1_profesor_P_alta", "M1_profesor_x_b",
       "M1_profesor_T_amb", "M1_profesor_base", "M1_comparacion",
       "M2_N1_malla", "M2_N1_biseccion", "M2_N1_fronteras", "M2_N1_resumen",
       "M2_N2_malla", "M2_N2_biseccion", "M2_N2_fronteras", "M2_N2_kalina",
       "M2_N2_resumen", "M3_ofat", "M3_resumen", "M3_cruces", "M2R_n1",
       "M2R_n2a", "M2R_n2b", "M2R_n3", "M2R_fronteras", "M2R_kalina",
       "M2R_resumen", "M3_figuras", "M2S_malla", "M2S_biseccion",
       "M2S_resumen", "M2S_kalina", "M4_lhs500", "M4_rsa", "M4_prcc",
       "M4_pawn", "M4_presiones_rangos", "M4_ventana_Palta",
       "M4_ventana_Tfuente", "M4_ventana_Tsumidero", "M4_bloqueo_rP",
       "M4_top_kalina", "M4_convergencia", "M4_figuras", "M5_lhs",
       "M5_resumen", "M5_rangos", "M5_cercano_profesor", "M5_escenarios",
       "M5_figuras"}


def filas_datos(ws) -> int:
    """Filas de datos: desde la 2 hasta la fila vacía o la descripción."""
    n = 0
    for r in range(2, ws.max_row + 1):
        v = ws.cell(r, 1).value
        if v in (None, "") or str(v).startswith("Descripción"):
            break
        n += 1
    return n


def descripcion(ws) -> str:
    """Todo el bloque de descripción al pie (desde la fila «Descripción»)."""
    for r in range(2, ws.max_row + 1):
        if str(ws.cell(r, 1).value or "").startswith("Descripción"):
            return "\n".join(str(ws.cell(k, 1).value or "")
                             for k in range(r, ws.max_row + 1))
    return ""


def main() -> int:
    fallos = 0

    def check(ok, txt):
        nonlocal fallos
        print(("OK   " if ok else "FALLO") + " " + txt)
        if not ok:
            fallos += 1

    if not XLSX.exists():
        print("FALLO no existe " + str(XLSX))
        return 1
    wb = load_workbook(XLSX)
    nombres = set(wb.sheetnames)
    faltan_pre = sorted(PRE - nombres)
    check(not faltan_pre, f"hojas previas presentes (faltan: {faltan_pre or 'ninguna'})")
    faltan_new = sorted((set(NUEVAS) | {FIGURA}) - nombres)
    check(not faltan_new, f"hojas nuevas presentes (faltan: {faltan_new or 'ninguna'})")
    for hoja, (csv, esp) in NUEVAS.items():
        ws = wb[hoja]
        n = filas_datos(ws)
        real = len(pd.read_csv(csv))
        check(n == esp and real == esp, f"{hoja}: {n} filas (CSV {real}, esperado {esp})")
    ws = wb["M8_ANECDOTICO_oat"]
    check("ANECDÓTICO: valores fuera de [0.40, 0.95] no son realistas" in descripcion(ws),
          "M8 descripción con la frase ANECDÓTICO")
    ws = wb[FIGURA]
    rutas = [ws.cell(r, 1).value for r in range(2, 2 + len(FIGS))]
    check(rutas == list(FIGS), f"{FIGURA}: rutas {rutas}")
    ws = wb["Resumen"]
    resumen = [ws.cell(r, 1).value for r in range(2, ws.max_row + 1)]
    resumen = [v for v in resumen if v not in (None, "")]
    faltan_res = [h for h in NUEVAS if h not in resumen]
    check(not faltan_res, f"Resumen con filas nuevas (faltan: {faltan_res or 'ninguna'})")
    check(FIGURA not in resumen, "Resumen sin la hoja de figuras")
    print("huella del libro: " + hashlib.md5(XLSX.read_bytes()).hexdigest())
    print("OK" if not fallos else f"FALLO ({fallos} comprobaciones)")
    return 0 if not fallos else 1


if __name__ == "__main__":
    sys.exit(main())