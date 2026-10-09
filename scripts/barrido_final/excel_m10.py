"""M10 — hojas finales en `datos/Excel_final_kalina.xlsx`.

Se corre SIEMPRE el último (después de `excel_final.py` … `excel_m9.py`): abre el
libro existente, crea o reemplaza las hojas `M10_catalogo`, `M10_recorrida_F3`,
`M10_ofat`, `M10_ofat_rangos`, `M10_xb_pbaja` y `M10_figuras` con los CSV de
`resultados/2026-10-09_final_M10_husavik/` vía `excel_final.escribir_tabla`
(mismo formato, unidades, colores de clase, filtro y descripción al pie), y
añade o actualiza sus filas en `Resumen`. Es idempotente.
"""
from __future__ import annotations

import sys  # noqa: E402
from pathlib import Path

_DIR = Path(__file__).resolve().parent
_RAIZ = _DIR.parents[1]
for _p in (str(_RAIZ), str(_DIR), str(_DIR.parent)):
    if _p not in sys.path:
        sys.path.insert(0, _p)

import pandas as pd  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

import excel_final as ef  # noqa: E402
from m10_datos import OUT  # noqa: E402

XLSX = ef.XLSX
FIGURA = "M10_figuras"
FIGS = (OUT / "m10_xb_pbaja.png", OUT / "m10_ofat.png")
PREFIJOS = ("M10_",)
# hoja -> (csv, descripción)
HOJAS = (
    ("M10_catalogo", "m10_catalogo.csv",
     "M10: una fila por familia; `KALINA sept` vs `KALINA ahora`, `coincide` y "
     "`d_eta_pp` (puntos porcentuales de η) frente a septiembre."),
    ("M10_recorrida_F3", "m10_recorrida_fase3.csv",
     "M10 paso A: cada punto KALINA de septiembre re-evaluado con teqp; "
     "`coincide_clase` dice si mantiene su clase."),
    ("M10_ofat", "m10_ofat.csv",
     "M10 paso B1: 15 valores por variable alrededor del ancla de Húsavík "
     "(OFAT); el resto de variables queda en el ancla."),
    ("M10_ofat_rangos", "m10_ofat_rangos.csv",
     "M10 paso B1: rango de valor y de η donde cada variable mantiene el punto "
     "como KALINA, con el criterio que lo rompe a cada lado."),
    ("M10_xb_pbaja", "m10_malla_xb_pbaja.csv",
     "M10 paso B2: malla x_b (0.70–0.91) × P_baja (500–1200 kPa) con el ancla "
     "fija en todo lo demás."),
)
UNIDADES = {
    "familia": "\u2212", "origen": "\u2212", "variables_movidas": "\u2212",
    "piso_O2": "K", "KALINA_sept": "n", "KALINA_ahora": "n",
    "confiables_ahora": "n", "coincide": "n", "coincide_clase": "V/F",
    "d_eta_pp_medio_abs": "pp", "d_eta_pp_max_abs": "pp", "d_eta_pp": "pp",
    "archivo_origen": "\u2212", "fila_origen": "n", "clasificacion": "\u2212",
    "eta_sept": "\u2212", "codigos_falla": "texto", "fuera_rango": "V/F",
    "B_verificado": "V/F", "A_verificacion_incompleta": "V/F",
    "A_recurrencias": "n", "repliegues_lo": "K", "repliegues_hi": "K",
    "variable": "\u2212", "valor": "\u2212", "min_kalina": "\u2212",
    "max_kalina": "\u2212", "criterio_bajo": "texto", "criterio_alto": "texto",
    "eta_min": "\u2212", "eta_max": "\u2212", "n_kalina": "n",
}


def _fila_resumen(hoja, df):
    """[hoja, filas, KALINA, CORREGIBLE, INVIABLE, NO_CONVERGIO, errores]."""
    if "clasificacion" not in df.columns:
        return [hoja, len(df)] + [None] * (len(ef.CLASES) + 1)
    n = df["clasificacion"].value_counts()
    err = int(df["error_inesperado"].notna().sum()) if "error_inesperado" in df.columns else 0
    return [hoja, len(df)] + [int(n.get(c, 0)) for c in ef.CLASES] + [err]


def _reescribir_resumen(ws, filas) -> None:
    for r in range(ws.max_row, 1, -1):
        if str(ws.cell(r, 1).value or "").startswith(PREFIJOS):
            ws.delete_rows(r)
    pos = 2
    while ws.cell(pos, 1).value not in (None, "") and not str(
            ws.cell(pos, 1).value).startswith("Descripción"):
        pos += 1
    ws.insert_rows(pos, len(filas))
    for k, fila in enumerate(filas):
        for j, v in enumerate(fila, 1):
            ws.cell(pos + k, j, v)


def _figuras(wb) -> str:
    """Hoja con las 2 PNG de M10; sin Pillow o sin PNG, solo las rutas."""
    ws = wb.create_sheet(FIGURA)
    for k, p in enumerate(FIGS, 2):
        ws.cell(k, 1, p.name)
        ws.cell(k, 2, p.relative_to(_RAIZ).as_posix())
    c = len(FIGS) + 3
    ws.cell(c, 1, "Figuras: (a) clase en la malla x_b × P_baja con isolíneas de η; "
                  "(b) η contra cada variable (ancla = línea roja).")
    ws.freeze_panes = "A2"
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 64
    faltan = [p.name for p in FIGS if not p.exists()]
    try:
        from openpyxl.drawing.image import Image as XLImage
    except ImportError:
        ws.cell(c + 2, 1, "Pillow no disponible: no se insertaron las imágenes.")
        return f"{FIGURA}: sin Pillow, solo las rutas"
    if faltan:
        ws.cell(c + 2, 1, "Faltan: " + ", ".join(faltan))
        return f"{FIGURA}: faltan {', '.join(faltan)}"
    fila = c + 2
    for p in FIGS:
        img = XLImage(str(p))
        ws.add_image(img, f"A{fila}")
        fila += -(-img.height // 20) + 1
    return f"{len(FIGS)} imágenes insertadas en {FIGURA}"


def main() -> None:
    faltan = [str(OUT / csv) for _h, csv, _d in HOJAS if not (OUT / csv).exists()]
    if not XLSX.exists():
        raise SystemExit(f"No existe {XLSX}: correr antes excel_final.py")
    if faltan:
        raise SystemExit("Faltan: " + ", ".join(faltan))
    ef.UNIDAD.update(UNIDADES)
    wb = load_workbook(XLSX)
    if "Resumen" not in wb.sheetnames:
        raise SystemExit("El libro no tiene la hoja Resumen: correr antes excel_final.py")
    viejas = [h for h, _c, _d in HOJAS if h in wb.sheetnames]
    for h in viejas + ([FIGURA] if FIGURA in wb.sheetnames else []):
        del wb[h]
    filas = []
    for hoja, csv, desc in HOJAS:
        df = pd.read_csv(OUT / csv)
        ef.escribir_tabla(wb, hoja, list(df.columns), df.values.tolist(), extra=[desc])
        filas.append(_fila_resumen(hoja, df))
    msg = _figuras(wb)
    _reescribir_resumen(wb["Resumen"], filas)
    wb.save(XLSX)
    print(f"escrito {XLSX}")
    print("hojas:", wb.sheetnames)
    print(msg)


if __name__ == "__main__":
    main()
