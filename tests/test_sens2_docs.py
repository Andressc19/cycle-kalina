# -*- coding: utf-8 -*-
"""Sens2 - pruebas de coherencia entre documentos y datos (solo lectura).

Las cifras globales de los documentos escritos a mano van marcadas como `<!--g:clave-->valor<!--/g-->`; estas pruebas fallan si alguna no
coincide con `p5_globales.py` (que las calcula de los archivos) o si desaparece un marcador. Ejecuta `scripts/sens2/docs_sync.py` para corregirlas.
"""
import os
import re
import sys

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(RAIZ, "scripts", "sens2"))
sys.path.insert(0, os.path.join(RAIZ, "scripts", "sens1"))
import docs_sync as DS  # noqa: E402
import p5_globales as G  # noqa: E402

MINIMO_MARCADORES = 9


def test_cifras_marcadas_al_dia():
    malas, total = DS.revisar(escribir=False)
    assert not malas, "documentos desactualizados (ejecuta docs_sync.py): %s" % malas
    assert total >= MINIMO_MARCADORES, "faltan marcadores de cifras globales (%d)" % total


def test_leeme_cita_los_puntos_del_motor():
    from openpyxl import load_workbook
    ws = load_workbook(G.LIBRO, read_only=True)["LEEME"]
    texto = " ".join(str(c) for fila in ws.iter_rows(values_only=True) for c in fila if c is not None)
    u1, b = G.puntos_motor()
    assert "anade %d puntos nuevos" % u1 in texto
    assert "%d puntos de la etapa B" % sum(b.values()) in texto


def test_informe_cita_hojas_y_graficos_del_libro():
    g = G.globales()
    with open(os.path.join(RAIZ, "docs", "INFORME_SENS2.md"), encoding="utf-8") as fh:
        t = fh.read()
    assert "(%d hojas, %d gráficos)" % (g["n_hojas"][0], g["n_graficos"][0]) in t


def test_marcadores_bien_formados():
    for ruta in DS.DOCS:
        with open(ruta, encoding="utf-8") as fh:
            s = fh.read()
        assert s.count("<!--g:") == s.count("<!--/g-->"), os.path.basename(ruta)
        assert len(re.findall(r"<!--g:", s)) == len(DS.PATRON.findall(s)), os.path.basename(ruta)


def _texto_hoja(nombre):
    from openpyxl import load_workbook
    ws = load_workbook(G.LIBRO, read_only=True)[nombre]
    return " ".join(str(c) for fila in ws.iter_rows(values_only=True) for c in fila if c is not None)


def test_mapa_de_preguntas_incluye_la_etapa_b():
    """Las conclusiones y hojas de la etapa B figuran en las tablas de preguntas de la consolidada y del LEEME."""
    cons = _texto_hoja("Conclusion consolidada")
    for ref in ("D8; D9", "U1 a U5", "M6; M7", "E9; E10", "F4 Optimo vs T_fuente", "F4 Caudales y bomba", "F3 Pisos 25 kPa y reintento"):
        assert ref in cons, "falta %r en la tabla de preguntas de la consolidada" % ref
    leeme = _texto_hoja("LEEME")
    for hoja in ("F3 Pisos 25 kPa y reintento", "F4 Caudales y bomba", "F4 Optimo vs T_fuente", "F5 Incertidumbre simetrica"):
        assert hoja in leeme, "falta %r en la guia del LEEME" % hoja


def test_pdf_incluye_las_conclusiones_de_la_etapa_b():
    ruta = os.path.join(RAIZ, "docs", "INFORME_TECNICO_SENS2.html")
    with open(ruta, encoding="utf-8") as fh:
        t = fh.read()
    for cid in ("D9", "L10", "E10", "U5", "M7"):
        assert ("<b>%s" % cid in t) or (cid in t), "el PDF no cita %s" % cid
    assert "etapa B" in t


def test_guia_del_leeme_sigue_el_orden_real_de_las_hojas():
    from openpyxl import load_workbook
    import s2_leeme
    real = load_workbook(G.LIBRO, read_only=True).sheetnames
    assert [n for n, _ in s2_leeme.GUIA] == real


def test_tabla_de_preguntas_del_leeme_y_de_la_consolidada_incluyen_la_etapa_b():
    """Cada pregunta del profesor lista las hojas de la etapa B que la responden (la tabla de preguntas, no solo la guia de hojas)."""
    import s2_conclusion
    import s2_leeme
    exigidas = {0: ["F3 Pisos 25 kPa y reintento"], 1: ["F5 Incertidumbre simetrica"], 2: ["F4 Optimo vs T_fuente"], 3: ["F4 Caudales y bomba"]}
    for nombre, tabla in (("LEEME", s2_leeme.PREGUNTAS), ("consolidada", s2_conclusion.PREGUNTAS)):
        for i, hojas in exigidas.items():
            for h in hojas:
                assert h in tabla[i][1], "%s: la pregunta %d no lista %r" % (nombre, i + 1, h)


def test_mapa_asigna_cada_conclusion_y_cada_hoja():
    """`s2_mapa` es la unica definicion del mapa de preguntas: ninguna conclusion ni hoja puede quedar sin asignar."""
    import pandas as pd
    import s2_mapa
    from openpyxl import load_workbook
    ids = list(pd.read_csv(os.path.join(RAIZ, "resultados", "2026-10-17_sens2_P5", "p5_conclusiones.csv"))["id"].astype(str))
    sin = [i for i in ids if i not in set(s2_mapa.conclusiones_asignadas())]
    assert not sin, "conclusiones sin asignar en s2_mapa: %s" % sin
    desconocidas = [i for i in set(s2_mapa.conclusiones_asignadas()) if i not in ids]
    assert not desconocidas, "s2_mapa cita conclusiones que no existen: %s" % desconocidas
    hojas = load_workbook(G.LIBRO, read_only=True).sheetnames
    sin_h = [h for h in hojas if h not in set(s2_mapa.hojas_asignadas())]
    assert not sin_h, "hojas sin asignar en s2_mapa: %s" % sin_h
    fantasmas = [h for h in set(s2_mapa.hojas_asignadas()) if h not in hojas]
    assert not fantasmas, "s2_mapa cita hojas que no existen: %s" % fantasmas


def test_pdf_incluye_cada_conclusion_una_vez():
    """Cada una de las conclusiones del CSV aparece en el PDF (como bloque propio), sin repetirse."""
    import pandas as pd
    ids = list(pd.read_csv(os.path.join(RAIZ, "resultados", "2026-10-17_sens2_P5", "p5_conclusiones.csv"))["id"].astype(str))
    with open(os.path.join(RAIZ, "docs", "INFORME_TECNICO_SENS2.html"), encoding="utf-8") as fh:
        t = fh.read()
    for cid in ids:
        n = t.count('<div class="c"><b>%s</b> [' % cid)
        assert n == 1, "la conclusion %s aparece %d veces en el PDF" % (cid, n)


def test_limites_citan_conclusiones_que_existen_y_el_leeme_los_lista():
    import pandas as pd
    import s2_limites
    ids = set(pd.read_csv(os.path.join(RAIZ, "resultados", "2026-10-17_sens2_P5", "p5_conclusiones.csv"))["id"].astype(str))
    for cid in s2_limites.conclusiones_citadas():
        assert cid in ids, "s2_limites cita %s, que no existe" % cid
    texto = _texto_hoja("LEEME")
    for limite, *_ in s2_limites.LIMITES:
        assert limite in texto, "el LEEME no lista el limite %r" % limite[:60]
