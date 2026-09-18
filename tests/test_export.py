"""Tests del módulo de exportación a Excel del ciclo Kalina KSC-11.

Usa datos de ejemplo construidos a mano (`tests/_datos_ejemplo.py`, misma
forma que `resolver_ciclo`/`calcular_exergia`) — nunca el backend real ni el
solver (una corrida real tarda ~8 min). Verifica que el `.xlsx` generado se
vuelve a abrir con `openpyxl.load_workbook` y tiene las 5 hojas esperadas con
el contenido correcto. `generar_excel` no debe escribir a disco.
"""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import load_workbook

from src.export_excel import VERSION_PROYECTO, generar_excel
from tests._datos_ejemplo import (BACKEND_EJEMPLO, PARAMETROS_EJEMPLO,
                                  resultado_ciclo_ejemplo,
                                  resultado_exergia_ejemplo)

_RAIZ = Path(__file__).resolve().parents[1]


@pytest.fixture
def libro():
    datos = generar_excel(resultado_ciclo_ejemplo(),
                          resultado_exergia_ejemplo(),
                          PARAMETROS_EJEMPLO, BACKEND_EJEMPLO)
    assert isinstance(datos, bytes) and len(datos) > 0
    return load_workbook(BytesIO(datos))


def test_devuelve_bytes_reabribles_con_las_5_hojas(libro):
    assert libro.sheetnames == ["Entradas", "Estados", "Componentes",
                                "Exergía", "Resumen"]


def test_hoja_entradas_una_fila_por_parametro_mas_backend(libro):
    ws = libro["Entradas"]
    assert [c.value for c in ws[1]] == ["Variable", "Valor", "Unidad"]
    assert ws.max_row == len(PARAMETROS_EJEMPLO) + 2    # header + parámetros + backend
    filas = {ws.cell(row=r, column=1).value: (ws.cell(row=r, column=2).value,
                                              ws.cell(row=r, column=3).value)
             for r in range(2, ws.max_row + 1)}
    for nombre, (valor, unidad) in PARAMETROS_EJEMPLO.items():
        assert filas[nombre] == (valor, unidad), nombre
    assert filas["backend"] == (BACKEND_EJEMPLO, "—")


def test_hoja_estados_10_filas_con_valores_y_unidades(libro):
    ws = libro["Estados"]
    assert [c.value for c in ws[1]] == ["Estado", "T [K]", "P [kPa]",
                                        "h [kJ/kg]", "s [kJ/kg·K]", "x [-]",
                                        "m [kg/s]", "q [-]", "Fase",
                                        "ex_fisica [kJ/kg]"]
    assert ws.max_row == 11                             # header + 10 estados
    # e1: T=350, P=3000, h=358, s=4.0, x=0.5, m=1.0, q=0.0, 'líquido',
    # ex_fisica = −97.1 (T0=300 K, P0=100 kPa).
    fila1 = [ws.cell(row=2, column=c).value for c in range(1, 11)]
    assert fila1[0] == "1" and fila1[1] == 350.0 and fila1[2] == 3000.0
    assert fila1[3] == pytest.approx(358.0)
    assert fila1[4] == pytest.approx(4.0)
    assert fila1[5] == 0.5 and fila1[6] == 1.0 and fila1[7] == 0.0
    assert fila1[8] == "líquido" and fila1[9] == pytest.approx(-97.1)
    # Última fila (e10).
    fila10 = [ws.cell(row=11, column=c).value for c in range(1, 11)]
    assert fila10[0] == "10" and fila10[9] == pytest.approx(-77.1)


def test_hoja_componentes_energia_relevante_y_convencion(libro):
    ws = libro["Componentes"]
    assert [c.value for c in ws[1]] == ["Componente", "Energía [kW]",
                                        "Tipo de energía"]
    assert ws.max_row == 9                              # header + 8 componentes
    filas = {ws.cell(row=r, column=1).value: ws.cell(row=r, column=2).value
             for r in range(2, ws.max_row + 1)}
    assert filas["HRVG"] == 50.0
    assert filas["Turbina"] == 11.3
    assert filas["Regenerador"] == 35.0
    assert filas["Condensador"] == 44.0
    assert filas["Bomba"] == 12.6
    # Convención documentada: sin energía externa -> celda "—".
    assert filas["Separador"] == "—"
    assert filas["Válvula"] == "—"
    assert filas["Absorbedor"] == "—"


def test_hoja_exergia_sgen_ed_y_residual_del_balance(libro):
    ws = libro["Exergía"]
    assert [c.value for c in ws[1]] == ["Componente", "Sgen [kW/K]", "Ed [kW]"]
    assert ws.max_row == 13                             # header + 8 + 4 resumen
    filas = {ws.cell(row=r, column=1).value: (ws.cell(row=r, column=2).value,
                                              ws.cell(row=r, column=3).value)
             for r in range(2, ws.max_row + 1)}
    assert filas["Turbina"] == (pytest.approx(-0.1), pytest.approx(-30.0))
    assert filas["Bomba"] == (pytest.approx(0.1), pytest.approx(30.0))
    assert filas["Ed_total (suma de componentes)"][1] == pytest.approx(6.5)
    assert filas["Ex_Qi (exergía de la fuente)"][1] == pytest.approx(12.5)
    assert filas["Ex_Qout (exergía del calor rechazado)"][1] == \
        pytest.approx(0.0)
    # Residual trazable: Ex_Qi − Wnet − Ex_Qout − Ed_total = 12.5 + 1.3 − 6.5.
    residual = filas["Residual: Ex_Qi - Wnet - Ex_Qout - Ed_total"][1]
    assert residual == pytest.approx(12.5 - (-1.3) - 0.0 - 6.5)


def test_hoja_resumen_energetico_exergetico_y_metadatos(libro):
    ws = libro["Resumen"]
    filas = {ws.cell(row=r, column=1).value: (ws.cell(row=r, column=2).value,
                                              ws.cell(row=r, column=3).value)
             for r in range(2, ws.max_row + 1)}
    assert filas["Qi"] == (50.0, "kW") and filas["Qout"] == (44.0, "kW")
    assert filas["Wt"] == (11.3, "kW") and filas["Wp"] == (12.6, "kW")
    assert filas["Wnet"] == (-1.3, "kW")
    assert filas["eta"][0] == pytest.approx(-0.026) and filas["eta"][1] == "-"
    assert filas["eta_exergetico (Wnet/Ex_Qi)"][0] == \
        pytest.approx(-1.3 / 12.5)
    assert filas["Versión del proyecto"] == (VERSION_PROYECTO, "—")
    assert filas["Backend"] == (BACKEND_EJEMPLO, "—")
    assert isinstance(filas["Fecha de generación"][0], str)
    assert filas["Fecha de generación"][0]


def test_generar_excel_no_escribe_a_disco(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    generar_excel(resultado_ciclo_ejemplo(), resultado_exergia_ejemplo(),
                  PARAMETROS_EJEMPLO, BACKEND_EJEMPLO)
    assert list(tmp_path.iterdir()) == []


_MODULOS_NUEVOS = ["src/export_excel.py", "src/plots.py",
                   "tests/test_export.py", "tests/test_plots.py",
                   "tests/_datos_ejemplo.py"]


def test_nuevos_archivos_dentro_del_limite_de_200_lineas():
    for ruta in _MODULOS_NUEVOS:
        assert len((_RAIZ / ruta).read_text(encoding="utf-8").splitlines()) <= 200


def test_export_y_plots_no_importan_backends_concretos():
    prohibidos = ("iapws", "AmmoniaWaterAdapter", "IAPWSAdapter",
                  "PyfluidsAdapter", "pyfluids", "CoolProp")
    for ruta in ("src/export_excel.py", "src/plots.py"):
        texto = (_RAIZ / ruta).read_text(encoding="utf-8")
        assert not any(p in texto for p in prohibidos), ruta