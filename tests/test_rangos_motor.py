"""Rango confiable de los motores (IAPWS G4-01 §6), modo A: extrapolar y marcar."""

import io
import warnings

import openpyxl
import pytest

from src import sensitivity as sens
from src._rango_barrido import anotar_rango
from src.excel_rango import GRIS_RANGO, marcar_estados, marcar_tabla
from src.properties.adapter import PropertyRangeError
from src.rangos_motor import (RangoMotorWarning, T_triple, evaluar_rango,
                              evaluar_rango_entradas)
from src.state import EstadoTermo

pytestmark = pytest.mark.filterwarnings("ignore::src.rangos_motor.RangoMotorWarning")


def _res(**cambios):
    """Ciclo sintético de 10 estados dentro de rango; `cambios` = {"e4": dict(T=...)}."""
    est = {}
    for i in range(1, 11):
        d = dict(T=350.0, P=3000.0, h=0.0, s=0.0, x=0.5, fase="liquido")
        d.update(cambios.get(f"e{i}", {}))
        est[f"e{i}"] = EstadoTermo(etiqueta=str(i), **d)
    return {"estados": est, "eta": 0.1, "Wnet": 50.0}


@pytest.mark.parametrize("w,T_ref", [(0.05, 267), (0.20, 236), (0.35, 175),
                                     (0.55, 190), (0.65, 194), (0.75, 188),
                                     (0.95, 192)])
def test_T_triple_reproduce_valores_de_control(w, T_ref):
    assert T_triple(w) == pytest.approx(T_ref, abs=1.0)


def test_dentro_de_rango_sin_violaciones_ni_avisos():
    with warnings.catch_warnings():
        warnings.simplefilter("error")
        assert evaluar_rango("teqp", _res()) == []


@pytest.mark.parametrize("cambio,param,tipo,col,limite", [
    (dict(P=45000.0), "P", "modelo", "P4", 40000.0),
    (dict(T=180.0, x=0.95), "T", "modelo", "T4", T_triple(0.95)),
    (dict(T=430.0, fase="liquido"), "T", "respaldo_experimental", "T4", 420.0),
    (dict(P=12000.0, fase="vapor", T=500.0), "P", "respaldo_experimental", "P4", 10000.0),
])
def test_cada_limite_genera_su_violacion(cambio, param, tipo, col, limite):
    v = evaluar_rango("real", _res(e4=cambio))
    assert len(v) == 1
    v = v[0]
    assert (v.parametro, v.tipo, v.columna, v.motor) == (param, tipo, col, "real")
    assert v.limite == pytest.approx(limite)
    assert col in v.texto() and "[motor: real]" in v.texto()


def test_bifasico_aplica_ambos_limites_experimentales():
    v = evaluar_rango("teqp", _res(e2=dict(T=430.0, P=12000.0, fase="bifasico")))
    assert {(x.parametro, x.tipo) for x in v} == {
        ("T", "respaldo_experimental"), ("P", "respaldo_experimental")}


def test_aviso_explicito():
    with pytest.warns(RangoMotorWarning, match="T4=430"):
        evaluar_rango("teqp", _res(e4=dict(T=430.0)))


class _Backend:
    def __init__(self, fase=None, falla=False):
        self.fase, self.falla, self.llamadas = fase, falla, 0

    def fase_de(self, P, T, x):
        self.llamadas += 1
        if self.falla:
            raise PropertyRangeError("sin campana")
        return self.fase, 0.0


def test_fase_none_se_calcula_con_fase_de():
    be = _Backend(fase="liquido")
    res = _res(**{f"e{i}": dict(fase=None) for i in range(1, 11)},)
    res["estados"]["e5"].T = 430.0
    v = evaluar_rango("real", res, backend=be)
    assert be.llamadas == 10 and [x.columna for x in v] == ["T5"]


def test_fase_no_determinable_se_registra_sin_adivinar():
    res = _res(e3=dict(fase=None, T=430.0))
    v = evaluar_rango("real", res, backend=_Backend(falla=True))
    assert [x.parametro for x in v] == ["fase"]


def test_modo_bloquear_reservado_y_motor_invalido():
    with pytest.raises(NotImplementedError):
        evaluar_rango("real", _res(), modo="bloquear")
    with pytest.raises(ValueError):
        evaluar_rango("otro", _res())


def test_entradas_no_convergio():
    combo = dict(P_alta=50000.0, P_baja=400.0, T_sumidero=300.0, x_b=0.5)
    v = evaluar_rango_entradas("teqp", combo)
    assert [x.columna for x in v] == ["P_alta"]


def test_anotar_rango_none_no_toca_la_fila():
    fila = {"eta": 0.1}
    assert anotar_rango(dict(fila), {}, _res(), None, None, "extrapolar_marcar") == fila


def test_kalina_y_fuera_rango_conviven_sin_reclasificar():
    fila = dict(clasificacion="KALINA", eta=0.1, Wnet=50.0, P_alta=3000.0)
    out = anotar_rango(fila, {}, _res(e4=dict(T=430.0)), "teqp", None,
                       "extrapolar_marcar")
    assert out["clasificacion"] == "KALINA" and out["fuera_rango"] is True
    assert out["columnas_rango"] == "eta;Wnet" and "T4=430" in out["detalle_rango"]


def test_no_convergio_marca_la_entrada():
    combo = dict(P_alta=50000.0, P_baja=400.0, T_sumidero=300.0, x_b=0.5)
    fila = dict(combo, clasificacion="NO_CONVERGIO", eta=None, Wnet=None)
    out = anotar_rango(fila, combo, None, "real", None, "extrapolar_marcar")
    assert out["clasificacion"] == "NO_CONVERGIO"
    assert out["fuera_rango"] is True and out["columnas_rango"] == "P_alta"


def test_ejecutar_barrido_integra_columnas(monkeypatch):
    class _Val:
        clasificacion = sens.Clasificacion.KALINA
        def mensaje_reporte(self):
            return "ok"
    monkeypatch.setattr(sens, "resolver_ciclo", lambda b, **k: _res(e4=dict(T=430.0)))
    monkeypatch.setattr(sens, "evaluar_ciclo", lambda *a, **k: _Val())
    var = {k: sens.Fijo(1.0) for k in sens.VARIABLES_BARRIBLES}
    base = sens.tabla_barrido(sens.ejecutar_barrido(None, var))
    assert "fuera_rango" not in base.columns
    t = sens.tabla_barrido(sens.ejecutar_barrido(None, var, motor_rango="teqp"))
    assert list(t.columns[-3:]) == ["fuera_rango", "detalle_rango", "columnas_rango"]
    assert t.loc[0, "clasificacion"] == "KALINA" and bool(t.loc[0, "fuera_rango"])


def _color(c):
    return c.fill.fgColor.rgb[-6:]


def test_excel_tabla_gris_y_comentario():
    wb = openpyxl.Workbook()
    ws = wb.active
    cab = ["eta", "Wnet", "clasificacion", "fuera_rango", "detalle_rango", "columnas_rango"]
    ws.append(cab)
    ws.append([0.1, 50.0, "KALINA", True, "T4=430 K > 420 K", "eta;Wnet"])
    ws.append([0.1, 50.0, "KALINA", False, None, None])
    assert marcar_tabla(ws, cab) == 2
    assert _color(ws["A2"]) == GRIS_RANGO and "T4=430" in ws["A2"].comment.text
    assert ws["A3"].comment is None and _color(ws["C2"]) != GRIS_RANGO


def test_export_excel_marca_estados():
    from src.export_excel import generar_excel
    res = _res(e4=dict(T=430.0))
    res.update(Qi=1.0, Qout=1.0, Wt=1.0, Wp=1.0, Qreg=1.0)
    comps = ("hrvg", "separador", "turbina", "regenerador", "valvula",
             "absorbedor", "condensador", "bomba")
    exe = dict(estados={}, sgen=dict.fromkeys(comps, 0.0), ed=dict.fromkeys(comps, 0.0),
               ed_total=0.0, ex_qi=1.0, ex_qout=0.0)
    v = evaluar_rango("teqp", res)
    wb = openpyxl.load_workbook(io.BytesIO(generar_excel(res, exe, {}, "teqp", v)))
    ws = wb["Estados"]
    fila = next(r for r in range(2, 12) if ws.cell(r, 1).value == "4")
    assert _color(ws.cell(fila, 2)) == GRIS_RANGO
    assert marcar_estados(ws, []) == 0
